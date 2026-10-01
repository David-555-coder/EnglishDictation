"""Run the real close/backup/replace/relaunch lifecycle against temporary data.

Only file dialogs, message boxes, and the final external launch are simulated.
The real main event loop, single-instance lock, window and recovery run normally.
"""
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from PySide6.QtCore import QLockFile, QProcess, QTimer
from PySide6.QtWidgets import QFileDialog, QMessageBox

import app as application
from src.database.repository import Repository
from src.services.library import ArticleLibrary, backup_library
from src.services.practice_service import new_full_session
from src.ui import main_window


def main():
    root = Path(tempfile.mkdtemp(prefix="EnglishDictationRecovery_"))
    old = Repository(root / "source" / "data" / "english_dictation.db")
    document = Document()
    paragraph = document.add_paragraph("A ")
    paragraph.add_run("rainbow").bold = True
    paragraph.add_run("（彩虹） appeared.")
    path = root / "文章.docx"
    document.save(path)
    article = ArticleLibrary(old).import_file(path)
    session = new_full_session(article)
    session.answers[session.target_ids[0]] = "saved before export"
    old.save_draft(session)
    backup = backup_library(old, root / "学习备份.zip")
    current_root = root / "target"
    current = Repository(current_root / "data" / "english_dictation.db")
    (current_root / "logs").mkdir()
    current.set_setting("current_only", True)
    application.resolve_storage_root = lambda: current_root
    QFileDialog.getOpenFileName = lambda *_: (str(backup), "")
    QMessageBox.question = lambda *_: QMessageBox.Yes
    messages = []
    QMessageBox.information = lambda *args: messages.append(args[1]) or QMessageBox.Ok
    QMessageBox.critical = lambda *args: (_ for _ in ()).throw(AssertionError(args))
    original_window = main_window.MainWindow

    def window_factory(repository):
        window = original_window(repository)
        QTimer.singleShot(220, window.restore_data)
        return window

    main_window.MainWindow = window_factory
    launches = []

    def launch(program, arguments, directory):
        lock = QLockFile(str(current.database_path.parent / "application.lock"))
        assert lock.tryLock(100), "Restart attempted before the data lock was released"
        lock.unlock()
        launches.append((program, arguments, directory))
        return True, 12345

    QProcess.startDetached = launch
    assert application.main() == 0
    assert launches and "恢复完成" in messages
    restored = Repository(current.database_path)
    assert restored.get_setting("current_only") is None
    assert restored.load_drafts()[0].answers == session.answers
    assert len(list((current_root / "backups" / "recovery").glob("*.zip"))) == 1
    assert not list(current_root.glob(".ed_restore_*"))
    print(f"{root}\nReal event loop, lock, close/save, safety backup, replace and restart dispatch verified", flush=True)


if __name__ == "__main__":
    main()
