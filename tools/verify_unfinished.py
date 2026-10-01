"""Visual checks for the new lightweight entries and draft selector, in isolation."""
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from PySide6.QtCore import QPoint
from PySide6.QtGui import QFont, QFontDatabase
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from src.database.repository import Repository
from src.services.practice_service import new_full_session
from src.services.library import ArticleLibrary
from src.ui.main_window import MainWindow
from src.ui.unfinished_dialog import UnfinishedPracticeDialog


def main():
    root = Path(tempfile.mkdtemp(prefix="EnglishDictationUnfinished_"))
    app = QApplication([])
    app.setStyle("Fusion")
    for filename in ("msyh.ttc", "msyhbd.ttc"):
        QFontDatabase.addApplicationFont(str(Path("C:/Windows/Fonts") / filename))
    app.setFont(QFont("Microsoft YaHei UI", 10))
    repo = Repository(root / "data" / "english_dictation.db")
    for index in range(10):
        document = Document()
        paragraph = document.add_paragraph("A ")
        paragraph.add_run("rainbow").bold = True
        paragraph.add_run("（彩虹） appeared.")
        path = root / f"文章{index + 1} - A Long Article Name.docx"
        document.save(path)
        article = ArticleLibrary(repo).import_file(path)
        session = new_full_session(article)
        if index % 2:
            session.answers[session.target_ids[0]] = "rainbow"
        repo.save_draft(session)
    QMessageBox.question = lambda *_: QMessageBox.Yes
    window = MainWindow(repo)
    window.show()
    QTest.qWait(220)
    for name, menu in (("practice", window.practice_menu), ("import", window.import_button.menu()),
                       ("export", window.export_button.menu())):
        menu.popup(window.menuBar().mapToGlobal(QPoint(120, 30)))
        QTest.qWait(40)
        menu.grab().save(str(root / f"{name}.png"))
        menu.close()
    window.show_preview()
    window.grab().save(str(root / "preview.png"))
    dialog = UnfinishedPracticeDialog(repo.unfinished_articles(), window)
    dialog.show()
    QTest.qWait(80)
    assert dialog.list.count() == 10
    dialog.grab().save(str(root / "unfinished.png"))
    dialog.resize(420, 330)
    dialog.setStyleSheet("QWidget { font-size:20px; }")
    QTest.qWait(80)
    dialog.grab().save(str(root / "unfinished-large-font.png"))
    dialog.close()
    window.close()
    app.processEvents()
    print(root, flush=True)


if __name__ == "__main__":
    main()
