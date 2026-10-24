"""Create isolated visual smoke-test artifacts; never open the real learning database."""
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from PySide6.QtWidgets import QApplication, QMessageBox
from PySide6.QtGui import QFontDatabase, QFont
from PySide6.QtTest import QTest
from src.database.repository import Repository
from src.ui.main_window import MainWindow
from src.ui.dialogs import WrongWordsDialog


def main():
    root = Path(tempfile.mkdtemp(prefix="EnglishDictationVisual_"))
    path = root / "The Big Family Get-Together.docx"
    doc = Document(); doc.add_paragraph("The Big Family Get-Together", "Title")
    paragraphs = [
        [("Lily's ", False), ("folks", True), ("（父母） planned a big family ", False), ("get-together", True), ("（聚会） at a small diner on the outskirts of town.", False)],
        [("The ", False), ("extensive", True), ("（广泛的） research gained significant ", False), ("momentum", True), ("（势头，冲劲）. This sentence provides more context for the first round.", False)],
        [("Her brother Max was ", False), ("reluctant", True), ("（勉强的）. He sat in a shady corner with his teddy bear.\nA second line stays in the same paragraph.", False)],
        [("We can ", False), ("take advantage of", True), ("（利用） the opportunity to learn something new every day.", False)],
    ]
    for parts in paragraphs:
        p = doc.add_paragraph()
        for text, bold in parts: p.add_run(text).bold = bold
    doc.save(path)
    app = QApplication([]); app.setStyle("Fusion")
    # Windows offscreen plugin does not enumerate installed fonts automatically.
    for name in ("msyh.ttc", "msyhbd.ttc"):
        font_path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / name
        if font_path.exists(): QFontDatabase.addApplicationFont(str(font_path))
    app.setFont(QFont("Microsoft YaHei UI", 10))
    QMessageBox.question = lambda *args: QMessageBox.Yes
    QMessageBox.information = lambda *args: QMessageBox.Ok
    window = MainWindow(Repository(root / "data" / "english_dictation.db"))
    window.show(); QTest.qWait(180)
    window._import_paths([str(path)]); QTest.qWait(250)
    window.grab().save(str(root / "01-preview.png"))
    window.start_full_practice(); QTest.qWait(120)
    for index, edit in enumerate(window.answer_edits.values()):
        if index < 3: edit.setText("my answer")
    window.grab().save(str(root / "02-practice.png"))
    window.actions['显示文章列表'].trigger(); QTest.qWait(80)
    window.grab().save(str(root / "02-collapsed.png"))
    window.actions['显示文章列表'].trigger(); QTest.qWait(80)
    window.set_zoom(150); QTest.qWait(100)
    window.grab().save(str(root / "03-zoom.png"))
    window.set_zoom(100); window.submit(); QTest.qWait(100)
    window.grab().save(str(root / "04-results.png"))
    window.retry_current_wrong(); QTest.qWait(100)
    window.grab().save(str(root / "05-wrong-preview.png"))
    window.start_full_practice(); QTest.qWait(100)
    window.grab().save(str(root / "06-wrong-dictation.png"))
    dialog = WrongWordsDialog(window.repository, window); dialog.show(); QTest.qWait(100)
    dialog.grab().save(str(root / "07-wordbook.png")); dialog.close()
    window.close(); app.processEvents()
    print(root, flush=True)


if __name__ == "__main__": main()
