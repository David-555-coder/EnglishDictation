"""Create an isolated native-EXE verification folder and a two-answer fixture."""
from pathlib import Path
import shutil
import sys
import tempfile

from docx import Document

root = Path(tempfile.mkdtemp(prefix="EnglishDictationSmoke_")) / "中文 测试"
root.mkdir()
shutil.copy2(Path(sys.argv[1]), root / "EnglishDictation.exe")
doc = Document()
doc.add_paragraph("Everyday Practice", "Title")
p = doc.add_paragraph("I have ")
p.add_run("ambition").bold = True
p.add_run("（抱负）. We can ")
p.add_run("take advantage of").bold = True
p.add_run("（利用） this opportunity.")
doc.save(root / "启动测试.docx")
print(root, flush=True)
