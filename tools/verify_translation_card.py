"""Synthetic, offline integrated-card snapshots. Not native desktop UX proof."""
import argparse
import json
import os
from pathlib import Path
import tempfile
import time


def main():
    parser=argparse.ArgumentParser();parser.add_argument('output',type=Path);args=parser.parse_args()
    from docx import Document
    from PySide6.QtCore import QPoint
    from PySide6.QtGui import QFontDatabase
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication
    from src.database.repository import Repository
    from src.services.reading_translation import Query
    from src.ui.main_window import MainWindow
    from src.ui.widgets.text_menus import TextLabel,plain_label_text,text_menu
    from src.ui.widgets.reading_selection import units
    app=QApplication([]);app.setStyle('Fusion')
    for name in ('msyh.ttc','msyhbd.ttc','segoeui.ttf'):
        path=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/name
        if path.exists():QFontDatabase.addApplicationFont(str(path))
    args.output.mkdir(parents=True,exist_ok=True)
    class CardWindow(MainWindow):
        def _restore_article_library(self):pass
    with tempfile.TemporaryDirectory(prefix='TranslationCardVisual_') as folder:
        root=Path(folder);document=Document()
        document.add_paragraph('Anna rinsed her glass in the sink. I was utterly exhausted.')
        document.add_paragraph('There was a charge for the service. As far as I am concerned, it was fair.')
        path=root/'sample.docx';document.save(path)
        window=CardWindow(Repository(root/'test.db'));window.resize(1100,850)
        window.show();window._article_switch_delay=0;window._import_paths([str(path)]);app.processEvents()
        rows=[]
        for word in ('utterly','rinsed','charge','As far as I am concerned'):
            window.reading_selection.clear()
            label=next(w for w in window.content.findChildren(TextLabel)
                       if w.isVisibleTo(window.content) and word in plain_label_text(w))
            full=plain_label_text(label);start=full.index(word)
            label.setSelection(units(full[:start]),units(word));menu=text_menu(label)
            began=time.perf_counter()
            next(a for a in menu.actions() if a.text()=='翻译').trigger();menu.deleteLater()
            popup=window._reading_translation.popup;deadline=time.monotonic()+3
            while not popup.lookup_done and time.monotonic()<deadline:QTest.qWait(5)
            assert popup.lookup_done;app.processEvents()
            elapsed=round((time.perf_counter()-began)*1000,3)
            popup.grab().save(str(args.output/(word.replace(' ','-')+'.png')))
            rows.append({'case':word,'width':popup.width(),'height':popup.height(),
                         'lookup_ms':elapsed,
                         'inside_screen':popup.screen().availableGeometry().contains(popup.frameGeometry())})
            if word=='rinsed':
                popup.show_online(Query(word,'Anna rinsed her glass in the sink.',True),'done','安娜在水槽里冲洗她的杯子。')
                app.processEvents();popup.grab().save(str(args.output/'rinsed-online.png'))
                rows.append({'case':'rinsed-online-synthetic-display-only','width':popup.width(),'height':popup.height(),
                             'inside_screen':popup.screen().availableGeometry().contains(popup.frameGeometry())})
        report={'scope':'Qt offscreen; online text is a synthetic appearance sample, not a service-quality test',
                'scale':os.environ.get('QT_SCALE_FACTOR','1'),'normal_learning_database_opened':False,
                'external_network_used':False,'network_client_created':window._reading_translation.network is not None,
                'rows':rows}
        window.close();app.processEvents()
        assert not popup._outside_filter_installed
    (args.output/'visual-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))
    assert all(row['inside_screen'] for row in rows)


if __name__=='__main__':main()
