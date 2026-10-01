"""Render the focus/continuation flow using generated articles and a temporary DB."""
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from PySide6.QtCore import QPoint, Qt
from PySide6.QtGui import QFont, QFontDatabase
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from src.database.repository import Repository
from src.ui.main_window import MainWindow


def main():
    root = Path(tempfile.mkdtemp(prefix="EnglishDictationFocus_"))
    path = root / "Long Article.docx"
    doc = Document()
    doc.add_paragraph("A Small Opportunity", "Title")
    words = [("rainbow", "彩虹"), ("condolence", "哀悼"),
             ("take advantage of", "利用"), ("get-together", "聚会")]
    for index in range(100):
        word, hint = words[index % len(words)]
        p = doc.add_paragraph("After the rain, Mia saw a ")
        p.add_run(word).bold = True
        p.add_run(f"（{hint}）. She shared this opportunity with her family and friends.")
    doc.save(path)
    app = QApplication.instance() or QApplication([])
    app.setStyle("Fusion")
    for name in ("msyh.ttc", "msyhbd.ttc"):
        font_path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / name
        if font_path.exists():
            QFontDatabase.addApplicationFont(str(font_path))
    app.setFont(QFont("Microsoft YaHei UI", 10))
    QMessageBox.question = lambda *args: QMessageBox.Yes
    QMessageBox.information = lambda *args: QMessageBox.Ok
    window = MainWindow(Repository(root / "data" / "english_dictation.db"))
    window.show()
    window.activateWindow()
    QTest.qWait(180)
    window._import_paths([str(path)])
    window.set_zoom(130)
    QTest.qWait(300)  # Let the original page finish showing before clicking Start.
    QTest.mouseClick(window.primary_action_button, Qt.LeftButton,
                     pos=window.primary_action_button.rect().center())
    window.activateWindow()  # The offscreen platform has no desktop window manager.
    # Offscreen Qt posts additional show/layout passes for a long new document.
    QTest.qWait(160)
    assert len(window.answer_edits) == 100
    QTest.qWait(500)
    assert window.scroll.verticalScrollBar().maximum() > 0
    assert all(edit.isVisible() and edit.height() > 0 for edit in window.answer_edits.values())
    assert QApplication.focusWidget() is window.answer_edits[window.session.target_ids[0]]
    window.grab().save(str(root / "01-normal.png"))
    window.view_menu.popup(window.menuBar().mapToGlobal(QPoint(210, 26)))
    QTest.qWait(30)
    window.view_menu.grab().save(str(root / "02-view-menu.png"))
    window.view_menu.close()
    window.activateWindow()

    ids = window.session.target_ids
    for item_id in ids[:32]:
        window.answer_edits[item_id].setText(window.session.article.item_map[item_id].answer)
    edit = window.answer_edits[ids[32]]
    edit.setFocus()
    QTest.qWait(30)
    before = edit.mapTo(window.content, QPoint()).y() - window.scroll.verticalScrollBar().value()
    QTest.keyClick(edit, Qt.Key_F8)
    QTest.qWait(80)
    assert window.actions["专注模式"].isChecked()
    after = edit.mapTo(window.content, QPoint()).y() - window.scroll.verticalScrollBar().value()
    assert abs(after - before) <= 3
    window.grab().save(str(root / "03-focus.png"))
    window.show_answers()
    QTest.qWait(80)
    window.grab().save(str(root / "04-answers.png"))
    window.show_answers()
    QTest.qWait(30)
    key = window.session.session_key
    window.show_preview()
    QTest.qWait(80)
    assert window.primary_action_button.text() == "开始默写"
    window.grab().save(str(root / "05-original-with-draft.png"))
    from src.ui.unfinished_dialog import UnfinishedPracticeDialog
    from PySide6.QtWidgets import QDialog
    original_exec = UnfinishedPracticeDialog.exec
    UnfinishedPracticeDialog.exec = lambda self: QDialog.Accepted
    try:
        window.resume_draft()
    finally:
        UnfinishedPracticeDialog.exec = original_exec
    QTest.qWait(160)
    assert window.session.session_key == key and window.session.filled == 32
    assert QApplication.focusWidget() is window.answer_edits[ids[32]]
    window.grab().save(str(root / "06-resumed.png"))

    window.resize(1000, 650)
    window.set_zoom(180)
    QTest.qWait(100)
    assert window.scroll.horizontalScrollBar().maximum() == 0
    window.grab().save(str(root / "07-narrow-180.png"))
    window.resize(2500, 1000)
    window.set_zoom(130)
    QTest.qWait(100)
    window.grab().save(str(root / "08-wide.png"))
    for index in range(32, 75):
        window._move_from(ids[index], 1)
    QTest.qWait(80)
    assert QApplication.focusWidget() is window.answer_edits[ids[75]]
    window.grab().save(str(root / "09-rapid-follow.png"))
    window.close()
    app.processEvents()
    print(f"{root}\n100 targets; focus, answers, resume, narrow/wide and rapid follow verified", flush=True)


if __name__ == "__main__":
    main()
