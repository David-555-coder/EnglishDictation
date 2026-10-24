"""Explicit, isolated packaged-release check. Never accesses normal storage."""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path
import tempfile
import time
import traceback


def run(output: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    result = {"passed": False, "normal_learning_database_opened": False}
    output.parent.mkdir(parents=True, exist_ok=True)
    ui_errors = []
    original_hook = sys.excepthook
    sys.excepthook = lambda *error: ui_errors.append("".join(traceback.format_exception(*error)))
    try:
        from datetime import datetime
        from docx import Document
        from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QTimer
        from PySide6.QtGui import QFontDatabase
        from PySide6.QtWidgets import QApplication, QMessageBox
        from src.database.repository import Repository
        from src.services.library import ArticleLibrary
        from src.services.fsrs7 import next_state, interval, MODEL_ID
        from src.ui.main_window import MainWindow
        from src.ui.widgets.flow_layout import FlowLayout

        started = time.perf_counter()
        state = next_state(None, 3)
        assert abs(state.stability - 3.9221) < 1e-5
        assert abs(state.difficulty - 3.5307243) < 1e-4
        assert abs(interval(state, .8) - 37.62522) < .1
        result["model"] = MODEL_ID
        app = QApplication.instance() or QApplication([])
        app.setStyle("Fusion")
        # Windows' offscreen plugin has no system font discovery; load installed
        # fonts for meaningful QA images only. Normal startup is unchanged.
        if not QFontDatabase.families() and os.name == "nt":
            fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
            for name in ("msyh.ttc", "segoeui.ttf"):
                QFontDatabase.addApplicationFont(str(fonts / name))
        result["qa_fonts_available"] = bool(QFontDatabase.families())
        QMessageBox.information = lambda *args: QMessageBox.Ok
        with tempfile.TemporaryDirectory(prefix="EnglishDictationRelease_") as directory:
            root = Path(directory)
            repo = Repository(root / "data" / "english_dictation.db")
            doc = Document(); doc.add_paragraph("Release Verification", "Title")
            p = doc.add_paragraph("This was the most "); p.add_run("thrilling").bold = True
            p.add_run("（令人激动的、惊险的） find ever. ")
            p.add_run("As far as I'm concerned").bold = True
            p.add_run("（就我而言）, it was excellent."); path = root / "测试.docx"; doc.save(path)
            article = ArticleLibrary(repo).import_file(path)
            with repo.connection() as conn:
                for item in article.word_items:
                    row = conn.execute("""INSERT INTO wrong_words(normalized_answer,answer,chinese,error_count,
                        first_error_at,last_error_at,review_due_date) VALUES (?,?,?,1,?,?,?)""",
                                       (item.normalized_answer,item.answer,item.chinese,datetime.now().isoformat(),
                                        datetime.now().isoformat(),repo.today().isoformat()))
                    conn.execute("""INSERT INTO word_examples(wrong_word_id,source_article,source_path,sentence)
                        VALUES (?,?,?,?)""",(row.lastrowid,item.filename,item.source_path,item.sentence))
            window = MainWindow(repo); window._article_switch_delay = 0
            # These synthetic imported errors are legacy words, admitted explicitly.
            repo.configure_review_capacity(40, 2)
            window.show()
            startup = QEventLoop(); QTimer.singleShot(180, startup.quit); startup.exec()
            app.processEvents()
            window.start_daily_review(); app.processEvents()
            assert window.session.total == 2 and not window._answers_visible

            def check_layout():
                app.processEvents()
                assert window.scroll.horizontalScrollBar().maximum() == 0
                scale = window.zoom_percent / 100
                for edit in window.answer_edits.values():
                    cap = 170 if len(edit.property("answerText").split()) == 1 else 320
                    assert edit.width() <= round(cap * scale)
                for flow in window.content.findChildren(FlowLayout):
                    left, _, right, bottom = flow.getContentsMargins()
                    parent = flow.parentWidget()
                    available = parent.width() - left - right
                    visible = [flow.itemAt(i).widget() for i in range(flow.count())
                               if not flow.itemAt(i).widget().isHidden()]
                    for widget in visible:
                        rect = widget.geometry()
                        assert rect.left() >= left and rect.right() < parent.width() - right
                        assert rect.bottom() < parent.height() - bottom
                        if widget.hasHeightForWidth():
                            assert widget.height() >= widget.heightForWidth(widget.width())
                        if widget.property("flowText"):
                            natural = widget.fontMetrics().horizontalAdvance(widget.text()) + 2
                            if widget.text().startswith("（") and natural <= available:
                                assert widget.height() <= widget.fontMetrics().height() + 2
                        if widget.property("flowAnswerHint"):
                            assert widget.width() <= round(280 * scale)
                            margins = widget.contentsMargins()
                            natural = (widget.fontMetrics().horizontalAdvance(widget.text())
                                       + margins.left() + margins.right() + 2)
                            if natural <= widget.width():
                                assert widget.height() <= (widget.fontMetrics().height()
                                                           + margins.top() + margins.bottom() + 2)
                    for index, widget in enumerate(visible):
                        assert not any(widget.geometry().intersects(other.geometry())
                                       for other in visible[index+1:])

            result["layout_matrix"] = []
            for width in (980, 1480, 2400):
                for zoom in (70, 100, 130, 180):
                    window.resize(width, 1000); window.set_zoom(zoom)
                    check_layout()
                    result["layout_matrix"].append({"width": width, "zoom": zoom})
            window.resize(1480, 1000); window.set_zoom(130); check_layout()
            window.grab().save(str(output.with_name(output.stem + "-practice.png")))
            first, second = window.session.target_ids
            window.answer_edits[first].setText(window.session.article.item_map[first].answer)
            window.show_answers(); window.answer_edits[second].setText("incorrect")
            check_layout()
            window.grab().save(str(output.with_name(output.stem + "-answers.png")))
            assert window._save_active_draft()
            draft = Repository(repo.database_path).load_drafts()[0]
            assert draft.view_state["memory"]["exposed"]
            window._commit_submission(); app.processEvents()
            assert window.session.correct_count == 1 and window.session.wrong_count == 1
            key = window.session.session_key
            repo.save_submitted_session(window.session)
            with repo.connection() as conn:
                rows = conn.execute("SELECT kind,rating FROM memory_events WHERE session_key=?",(key,)).fetchall()
                assert len(rows) == 2
                assert {tuple(r) for r in rows} == {("first",3),("initial_learning",1)}
            assert repo.review_status()["completed"] == 2
            window.grab().save(str(output.with_suffix(".png")))
            window.retry_current_wrong(); window.start_full_practice()
            edit = next(iter(window.answer_edits.values()))
            edit.setText(window.session.article.item_map[edit.item_id].answer)
            window._commit_submission(); app.processEvents()
            assert window.session is not None and "纠错后结果" in window.score_text.text()
            correction = window.session
            window._restore_article_library(); app.processEvents()
            assert window.session is correction and "纠错后结果" in window.score_text.text()
            assert repo.review_status()["completed"] == 2
            with repo.connection() as conn:
                assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
                assert conn.execute("SELECT count(*) FROM memory_events WHERE kind='relearning' AND rating=3").fetchone()[0] == 1
            from src.services.review_dialog_check import verify
            dialog_checks = verify(repo, window, app, output)
            from src.services.review_stream_check import verify as verify_streams
            stream_checks = verify_streams(root, window, app, output)
            from src.services.learning_flow_check import verify as verify_learning_flow
            learning_flow_checks = verify_learning_flow(root, window, app)
            from src.services.scroll_position_check import verify as verify_scroll_position
            scroll_position_checks = verify_scroll_position(root, app, output)
            from src.services.continuity_check import verify as verify_continuity
            continuity_checks = verify_continuity(root, window, app, output)
            from src.services.startup_check import verify as verify_startup
            startup_checks = verify_startup(root, app, output)
            from src.services.ui_polish_check import verify as verify_ui_polish
            ui_polish_checks = verify_ui_polish(root, app, output)
            from src.services.reading_surface_check import verify as verify_reading_surface
            reading_surface_checks = verify_reading_surface(root, app, output)
            from src.services.learning_content_check import verify as verify_learning_content
            learning_content_checks = verify_learning_content(root, app, output)
            from src.services.text_menu_check import verify as verify_text_menus
            text_menu_checks = verify_text_menus(root, app, output)
            from src.services.sentence_selection_check import verify as verify_sentence_selection
            sentence_selection_checks = verify_sentence_selection(root, app, output)
            from src.services.ui_performance_check import verify as verify_ui_performance
            ui_performance_checks = verify_ui_performance(root, app, output)
            from src.services.shortcut_help_check import verify as verify_shortcut_help
            shortcut_help_checks = verify_shortcut_help(root, app, output)
            from src.services.quick_write_check import verify as verify_quick_write
            quick_write_checks = verify_quick_write(root, app, output)
            from src.services.reading_translation_check import verify as verify_translation
            translation_checks = verify_translation(root, app, output)
            window.close(); window.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete); app.processEvents()
            assert not ui_errors, ui_errors
            result["checks"] = ["official_default_vector", "cold_daily_start", "preserved_pre_reveal_answer",
                                 "exposure_draft_restore", "separate_relearning", "no_extra_daily_slot",
                                 "idempotent_save", "database_integrity", "word_document_import",
                                 "late_startup_restore_preserves_results", "no_qt_callback_errors",
                                 "compact_field_bounds", "chinese_annotation_layout", "bounded_answer_hints"] + dialog_checks + stream_checks + learning_flow_checks + scroll_position_checks + continuity_checks + startup_checks + ui_polish_checks + reading_surface_checks + learning_content_checks + text_menu_checks + sentence_selection_checks + ui_performance_checks + shortcut_help_checks + quick_write_checks + translation_checks
        result["passed"] = True
        result["elapsed_seconds"] = round(time.perf_counter()-started,3)
    except Exception:
        result["error"] = traceback.format_exc()
    finally:
        sys.excepthook = original_hook
    output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return 0 if result["passed"] else 1
