"""Check help entry points on synthetic data, including the packaged UI."""


def verify(root, app, output):
    from copy import deepcopy
    import traceback
    from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, Qt, QTimer
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QLabel, QPushButton, QScrollArea
    from src.database.repository import Repository
    from src.services.library import ArticleLibrary
    from src.services.practice_service import new_full_session
    from src.ui.main_window import MainWindow
    from src.ui.shortcut_help import ShortcutHelpDialog
    from docx import Document
    from shiboken6 import isValid

    folder = root / 'shortcut-help'; folder.mkdir()
    doc = Document(); p = doc.add_paragraph('Keep a steady ')
    p.add_run('pace').bold = True; p.add_run('（节奏）.')
    path = folder / 'Help.docx'; doc.save(path)
    repo = Repository(folder / 'data' / 'copy.db')
    article = ArticleLibrary(repo).import_file(path)
    window = MainWindow(repo); window.show()
    loop = QEventLoop(); QTimer.singleShot(180, loop.quit); loop.exec()
    window._activate_session(new_full_session(article))
    edit = next(iter(window.answer_edits.values())); edit.setText('my answer'); edit.setFocus()
    app.processEvents()
    original_session = window.session
    memory = deepcopy(window.session.view_state.get('memory'))
    hint = window.shortcut_hint.text()
    assert [a.text() for a in window.menuBar().actions()][-1] == '帮助'
    assert window.actions['快捷键说明'].shortcut().toString() == 'F1'
    opened = []; errors = []

    def inspect():
        dialog = app.activeModalWidget()
        try:
            assert isinstance(dialog, ShortcutHelpDialog)
            opened.append(dialog)
            assert len(dialog.findChildren(QLabel, 'shortcutSection')) == 4
            scroll = dialog.findChild(QScrollArea)
            assert scroll.horizontalScrollBar().maximum() == 0
            if dialog.height() >= 660:
                assert scroll.verticalScrollBar().maximum() == 0
            # Save both the default view and the minimum-size scrolling fallback.
            if len(opened) == 1:
                dialog.grab().save(str(output.with_name(output.stem + '-shortcut-help.png')))
                dialog.resize(560, 420); app.processEvents()
                assert scroll.horizontalScrollBar().maximum() == 0
                assert scroll.verticalScrollBar().maximum() > 0
                scroll.verticalScrollBar().setValue(scroll.verticalScrollBar().maximum())
                app.processEvents()
                dialog.grab().save(str(output.with_name(output.stem + '-shortcut-help-small.png')))
                QTest.keyClick(dialog, Qt.Key_Escape)
            else:
                dialog.findChild(QPushButton).click()
            assert not dialog.isVisible()
        except Exception:
            errors.append(traceback.format_exc())
        finally:
            if dialog is not None and isValid(dialog):
                dialog.reject()

    for trigger in ('menu', 'f1', 'focus_f1'):
        if trigger == 'focus_f1':
            window._set_focus_mode(True)
        window.activateWindow(); edit.setFocus(); app.processEvents()
        timer = QTimer(window); timer.setSingleShot(True)
        timer.timeout.connect(inspect); timer.start(50)
        if trigger == 'menu':
            window.actions['快捷键说明'].trigger()
        else:
            QTest.keyClick(edit, Qt.Key_F1)
        timer.stop(); timer.deleteLater()
        assert not errors, errors
        assert len(opened) == ('menu', 'f1', 'focus_f1').index(trigger) + 1
        assert window.session is original_session and edit.text() == 'my answer'
        assert window.session.view_state.get('memory') == memory
        assert window.shortcut_hint.text() == hint
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete); app.processEvents()
        assert not isValid(opened[-1])
    window.close(); window.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete); app.processEvents()
    return ['help_menu_entry', 'help_f1_from_answer', 'help_f1_in_focus_mode',
            'help_escape_and_close_button', 'help_small_window_scroll',
            'help_preserves_answers_and_memory', 'help_dialog_released']
