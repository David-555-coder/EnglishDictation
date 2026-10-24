"""Real drag/copy checks inside the packaged app, temporary learning data only."""


def verify(root, app, output):
    from pathlib import Path
    from docx import Document
    from PySide6.QtCore import QCoreApplication, QEvent, QPoint, Qt
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication, QFrame
    from src.database.repository import Repository
    from src.services.practice_service import new_full_session, new_long_term_session, group_items_by_sentence
    from src.services.daily_review import new_daily_session
    from src.services.serialization import session_to_dict
    from src.ui.main_window import MainWindow
    from src.ui.widgets.text_menus import TextLabel, text_menu

    folder=Path(root)/'sentence-selection';folder.mkdir()
    path=folder/'Selection Story.docx';doc=Document()
    first='On the whole（总体而言）, it was a stress-free（无压力的） night.'
    second='Her goal was to outdo（超过） her prior self & keep reading.'
    p=doc.add_paragraph('On the whole（总体而言）, it was a ');p.add_run('stress-free').bold=True;p.add_run('（无压力的） night.')
    p=doc.add_paragraph('Her goal was to ');p.add_run('outdo').bold=True;p.add_run('（超过） her prior self & keep reading.')
    doc.save(path);repo=Repository(folder/'data'/'db');window=MainWindow(repo)
    window._article_switch_delay=0;window._import_paths([str(path)]);window.show();app.processEvents()
    checks=[]

    def words():
        return [w for w in window.content.findChildren(TextLabel) if w.isVisibleTo(window.content)
                and w.property('flowText') and not w.property('flowAnswerHint') and w.objectName()!='learningNote']

    def drag_copy(a,b,reverse=False):
        if reverse:a,b=b,a
        start=QPoint(a.width()-1 if reverse else 1,1)
        end=a.mapFromGlobal(b.mapToGlobal(QPoint(1 if reverse else b.width()-1,b.height()-1)))
        QTest.mousePress(a,Qt.LeftButton,pos=start);QTest.mouseMove(a,end,10)
        QTest.mouseRelease(a,Qt.LeftButton,pos=end);QTest.keyClick(a,Qt.Key_C,Qt.ControlModifier)
        assert not window.reading_selection._scroll_timer.isActive()
        return QApplication.clipboard().text()

    def assert_page(name, expected=first+'\n\n'+second):
        app.processEvents();labels=words()
        for reverse in (False,True):
            copied = drag_copy(labels[0],labels[-1],reverse).strip()
            assert copied == expected, (name, repr(copied), repr(expected))
        checks.append('sentence_drag_'+name)

    try:
        window.set_zoom(130);app.processEvents()
        original=window.content.findChild(QFrame,'originalPage').findChildren(TextLabel)
        assert drag_copy(original[0],original[-1]).strip()==first+'\n\n'+second
        checks.append('sentence_drag_original_rich')
        window._show_wrong_preview(new_full_session(window.current_article))
        for width,zoom in ((980,70),(980,180),(1480,100),(2400,130)):
            window.resize(width,1000);window.set_zoom(zoom);assert_page(f'preview_{width}_{zoom}')
            assert window.scroll.horizontalScrollBar().maximum()==0
        window.resize(1480,1000);window.set_zoom(130);assert_page('preview')
        window.grab().save(str(output.with_name(output.stem+'-sentence-preview.png')))
        label=words()[0];menu=text_menu(label);menu.ensurePolished()
        assert [a.text() for a in menu.actions() if not a.isSeparator()]==['复制','全选','翻译']
        next(a for a in menu.actions() if a.text()=='复制').trigger()
        assert QApplication.clipboard().text().strip()==first+'\n\n'+second
        assert menu.font().pixelSize()==14;menu.deleteLater();checks.append('sentence_context_copy')
        window.start_full_practice();app.processEvents()
        session=window.session;before=session_to_dict(session)
        copied=drag_copy(words()[0],words()[-1])
        assert '____' in copied and 'stress-free' not in copied and 'outdo' not in copied
        assert session_to_dict(session)==before;checks.append('sentence_no_hidden_answer_leak')
        edits=list(window.answer_edits.values());edits[0].setText('typed answer');edits[0].setFocus();edits[0].selectAll()
        QTest.keyClick(edits[0],Qt.Key_C,Qt.ControlModifier);assert QApplication.clipboard().text()=='typed answer'
        checks.append('sentence_native_input_copy')
        window._commit_submission();app.processEvents()
        label=words()[0];QTest.keyClick(label,Qt.Key_A,Qt.ControlModifier)
        QTest.keyClick(label,Qt.Key_C,Qt.ControlModifier)
        assert '正确答案：stress-free' in QApplication.clipboard().text()
        checks.append('sentence_visible_results_copy')
        window.read_current_wrong();app.processEvents()
        before=session_to_dict(window.session)
        with repo.connection() as db: before_events=db.execute('SELECT count(*) FROM memory_events').fetchone()[0]
        assert_page('wrong_reading');window.grab().save(str(output.with_name(output.stem+'-sentence-wrong-reading.png')))
        assert before==session_to_dict(window.session)
        with repo.connection() as db: assert before_events==db.execute('SELECT count(*) FROM memory_events').fetchone()[0]
        checks.append('sentence_selection_no_model_updates')
        with repo.connection() as db: ids=[r[0] for r in db.execute('SELECT id FROM wrong_words')]
        records=repo.get_review_examples(ids)
        for name,session in (('daily',new_daily_session(records,repo.today().isoformat())),
                             ('wordbook',new_long_term_session(records))):
            # Daily sessions intentionally shuffle their examples. Validate the
            # model's actual sentence order, not one lucky random arrangement.
            expected = '\n\n'.join(group[0].paragraph_text[group[0].sentence_start:group[0].sentence_end]
                                     for group in group_items_by_sentence(session))
            window._show_wrong_preview(session);assert_page(name+'_preview', expected)
            window.start_full_practice();app.processEvents()
            text=drag_copy(words()[0],words()[-1]);assert 'stress-free' not in text and 'outdo' not in text
            checks.append('sentence_'+name+'_native_practice')
        window.start_saved_wrong_practice(window.repository.list_history()[0]['id']);assert_page('history_preview')
        assert not window.reading_selection._scroll_timer.isActive()
        checks.append('sentence_no_idle_selection_timer')
    finally:
        window.close();window.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete);app.processEvents()
    return checks
