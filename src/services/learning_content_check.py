"""Synthetic, isolated content-editor checks included in the packaged executable."""
import hashlib


def verify(root, app, output):
    from docx import Document
    from PySide6.QtWidgets import QLabel
    from src.database.repository import Repository
    from src.models.rules import RecognitionRules
    from src.services.library import ArticleLibrary
    from src.services.practice_service import new_full_session, new_wrong_session, new_saved_wrong_session
    from src.services.grading import grade_session
    from src.services.learning_content import prepare_article, edit_hint
    from src.services.serialization import session_to_dict
    from src.ui.editing_dialogs import TargetEditorDialog
    from src.ui.main_window import MainWindow

    folder = root / 'learning-content'; folder.mkdir()
    source = folder / 'Learning with Context.docx'
    doc = Document(); doc.add_paragraph('Learning with Context', 'Title')
    p = doc.add_paragraph('A good story can '); p.add_run('hold').bold = True
    p.add_run("（容纳） everyone's attention. A clear example makes a difficult word easier to remember.")
    p = doc.add_paragraph('Mia checked the '); p.add_run('internal').bold = True
    p.add_run('（内部的） rules before the meeting. She wanted to understand why the team had made this choice.')
    p = doc.add_paragraph('Later, she asked Leo to '); p.add_run('hold').bold = True
    p.add_run('（握住） the rope. The same word can have a different meaning in another sentence.')
    doc.save(source); digest = hashlib.sha256(source.read_bytes()).hexdigest()
    repo = Repository(folder/'data'/'db'); article = ArticleLibrary(repo).import_file(source)
    first = new_full_session(article); grade_session(first); repo.save_submitted_session(first)
    draft = new_wrong_session(first); repo.save_draft(draft); frozen = session_to_dict(draft)
    editor = TargetEditorDialog(article)
    editor.show(); app.processEvents()
    editor.chinese_edit.setText('保持大家的注意力')
    editor.note_edit.setPlainText('hold attention：保持注意力；hold a rope：握住绳子。')
    editor._apply_hint(); app.processEvents()
    editor.grab().save(str(output.with_name(output.stem+'-content-editor.png')))
    assert editor.article.word_items[-1].chinese == '握住'
    assert repo.save_learning_content(editor.original_article,editor.article) == 1
    editor.accept(); editor.deleteLater(); app.processEvents()
    assert session_to_dict(repo.load_drafts()[0]) == frozen
    assert repo.get_wrong_answers(first.database_id)[0]['chinese'] == '容纳'
    renewed = repo.prepare_learning_session(new_wrong_session(first))
    assert renewed.article.word_items[0].chinese == '保持大家的注意力'
    rows = repo.get_wrong_answers(first.database_id)
    history = repo.prepare_learning_session(new_saved_wrong_session({'id':first.database_id},rows))
    assert history.article.word_items[0].chinese == '保持大家的注意力'
    current = repo.load_article(article.source_path); candidate = prepare_article(current)
    edit_hint(candidate,candidate.word_items[1].id,'里面的','internal / inherent')
    repo.save_learning_content(current,candidate)
    with repo.connection() as db:
        word_id = db.execute("SELECT id FROM wrong_words WHERE answer='internal'").fetchone()[0]
    assert repo.get_review_examples([word_id])[0]['chinese'] == '里面的'
    reparsed = ArticleLibrary(repo).reparse(repo.load_article(article.source_path), RecognitionRules())
    assert reparsed.word_items[0].chinese == '保持大家的注意力'
    window = MainWindow(repo); window._article_switch_delay=0; window._restore_article_library(); window.show_preview()
    window.show(); app.processEvents()
    window.grab().save(str(output.with_name(output.stem+'-content-reading.png')))
    window.session = new_full_session(repo.load_article(article.source_path))
    window._render_practice(full_article=True); app.processEvents()
    notes = window.content.findChildren(QLabel,'learningNote')
    assert notes and all(n.isHidden() for n in notes)
    window.show_answers(); assert all(not n.isHidden() for n in notes)
    window.show_answers(); assert all(n.isHidden() for n in notes)
    window.close(); window.deleteLater(); app.processEvents()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    with repo.connection() as db: assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    return ['content_sentence_specific_hint','content_inline_chinese_sync','content_preserves_old_draft',
            'content_preserves_historical_grade','content_fresh_correction','content_fresh_history_practice',
            'content_daily_example_revision','content_reparse_preserves_revision','content_note_visibility',
            'content_original_word_unchanged','content_database_integrity']
