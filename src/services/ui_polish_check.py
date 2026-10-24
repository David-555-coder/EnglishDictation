"""Generated-data UI checks. Never opens the normal learning database."""
from datetime import timedelta


def verify(root, app, output):
    from docx import Document
    from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QPoint, QTimer
    from PySide6.QtWidgets import QFrame, QInputDialog
    from src.database.repository import Repository
    from src.services.library import ArticleLibrary
    from src.services.serialization import session_to_dict
    from src.ui.main_window import MainWindow

    def settle(ms=60):
        loop = QEventLoop(); QTimer.singleShot(ms, loop.quit); loop.exec(); app.processEvents()

    def dispose(window):
        window.close(); window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete); app.processEvents()

    def rows(repo):
        with repo.connection() as db:
            return {t: [tuple(r) for r in db.execute('SELECT * FROM '+t)] for t in
                    ('practice_sessions', 'practice_answers', 'practice_drafts', 'wrong_words', 'daily_review_items', 'memory_events')}

    folder = root / 'ui-polish'; folder.mkdir()
    repo = Repository(folder / 'data' / 'copy.db')
    words = 'workforce forthcoming ambition change river landscape condition sensitive sound guilt lag offline curtain dormitory sustainability indispensable conversation ethical nature map town sunlight nightmare wake'.split()
    for name in ('河谷故事', '阅读练习'):
        doc = Document(); doc.add_paragraph('A Small Change', 'Title')
        for i in range(24):
            p = doc.add_paragraph('They used '); p.add_run(words[i]).bold = True
            p.add_run('（示例词） to describe the change in their town.')
        path = folder / (name+'.docx'); doc.save(path); ArticleLibrary(repo).import_file(path)
    repo.set_setting('last_article_path', str(folder / '阅读练习.docx'))
    window = MainWindow(repo)
    try:
        window.show(); settle(350); window.start_full_practice(); settle()
        key = window.session.session_key; first = window.session.target_ids[0]
        window.answer_edits[first].setText('saved answer'); assert window._save_active_draft()
    finally:
        dispose(window)
    before = rows(repo)
    window = MainWindow(repo)
    try:
        window.show(); window.set_zoom(160); settle(350)
        assert window.session is None and not window.answer_edits
        assert window.current_article.filename == '阅读练习.docx' and window.mode_title.text().startswith('原文模式')
        assert window.progress.isHidden() and window.score_text.isHidden() and rows(repo) == before
        window.grab().save(str(output.with_name(output.stem+'-original-startup.png')))
        window._activate_session(next(s for s, _ in repo.unfinished_practices() if s.session_key == key)); settle()
        assert window.answer_edits[first].text() == 'saved answer'
        window._commit_submission(); settle()
        original = session_to_dict(window.session); before = rows(repo)
        bar = window.scroll.verticalScrollBar(); bar.setValue(bar.maximum() // 2)
        offset = bar.value(); assert offset > 0
        window.read_current_wrong(); settle()
        assert window._wrong_reading and not window.answer_edits and window.pending_review_session is None
        assert app.activeModalWidget() is None and session_to_dict(window.session) == original and rows(repo) == before
        assert 'workforce' in '\n'.join(c.property('paragraphText') for c in window.content.findChildren(QFrame) if c.property('paragraphText'))
        for width in (980, 1480, 2400):
            for zoom in (70, 100, 160, 180):
                window.resize(width, 900); window.set_zoom(zoom); settle(20)
                assert window.scroll.horizontalScrollBar().maximum() == 0
                assert abs(window.zoom_panel.mapTo(window, QPoint()).x()-window.status_panel.mapTo(window, QPoint()).x()) <= 2
                assert window.zoom_panel.width() == window.status_panel.width()
                if bar.isVisible():
                    gap = window.status_panel.mapTo(window, QPoint()).x()-bar.mapTo(window, QPoint()).x()-bar.width()
                    assert gap <= 11
        window.resize(1480, 900); window.set_zoom(160); settle()
        window.grab().save(str(output.with_name(output.stem+'-main-wrong-reading.png')))
        window.actions['专注模式'].trigger(); settle()
        assert window._wrong_reading and window.status_panel.isVisible() and window.reading_back_button.isVisible()
        window.grab().save(str(output.with_name(output.stem+'-focused-reading.png')))
        window.actions['专注模式'].trigger(); settle()
        window.reading_back_button.click(); settle()
        assert not window._wrong_reading and session_to_dict(window.session) == original
        assert abs(bar.value()-offset) <= 3 and rows(repo) == before
        window.read_current_wrong(); window.primary_action_button.click(); settle()
        assert not window._wrong_reading and not window.session.submitted and window.pending_review_session is None
        window._commit_submission(); settle(); window.read_current_wrong(); settle()
        correction_key = window.session.session_key; saved_get_int = QInputDialog.getInt
        try:
            QInputDialog.getInt = lambda *args: (3, False)
            window.primary_action_button.click()
            assert window._wrong_reading and window.session.session_key == correction_key
            QInputDialog.getInt = lambda *args: (3, True)
            window.primary_action_button.click(); settle()
            assert window.session.total == 3 and not window.session.submitted
            assert window.session.view_state['consolidation_origin']['session_key'] == correction_key
        finally:
            QInputDialog.getInt = saved_get_int
    finally:
        dispose(window)

    quantity = Repository(folder / 'quantity' / 'data' / 'copy.db')
    day = quantity.today().isoformat(); future = (quantity.today()+timedelta(days=2)).isoformat()
    quantity.enable_review_continuity()
    with quantity.connection() as db:
        for i in range(100):
            word = f'word{i}'
            wid = db.execute('''INSERT INTO wrong_words(answer,normalized_answer,chinese,error_count,
                first_error_at,last_error_at,review_due_date,review_active) VALUES (?,?,?,1,?,?,?,?)''',
                (word, word, '示例词', day, day, future if 71 <= i < 77 else day, int(i < 77))).lastrowid
            db.execute("INSERT INTO word_examples(wrong_word_id,source_article,source_path,sentence) VALUES (?,'generated','generated',?)",
                       (wid, f'They used {word}（示例词） to describe the change.'))
    quantity.configure_review_capacity(80, 20)
    window = MainWindow(quantity)
    try:
        window.show(); settle(200); window.start_daily_review(); settle()
        assert window.session.total == 80 and window.mode_detail.text() == '今日安排 80 词 · 每日最多 80 词'
        assert '到期复习：71' in window.progress_text.text() and '加入旧词：9' in window.progress_text.text()
        before = rows(quantity); window._refresh_daily_review(); assert rows(quantity) == before
        window.set_zoom(160); settle()
        window.grab().save(str(output.with_name(output.stem+'-capacity-explanation.png')))
    finally:
        dispose(window)
    return ['startup_original_last_article', 'manual_draft_answers_preserved', 'preview_sidebar_context',
            'main_pane_wrong_reader', 'reading_no_automatic_next_round', 'reading_result_and_scroll_return',
            'reader_zoom_and_f8', 'zoom_sidebar_alignment', 'compact_article_gutter',
            'explicit_correction_from_reader', 'optional_bounded_group_from_reader', 'actual_daily_capacity_explanation']
