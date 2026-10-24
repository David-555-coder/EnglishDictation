"""Synthetic release checks for the scoped UI/performance refinement."""
from copy import deepcopy


def verify(root, app, output):
    from pathlib import Path
    from docx import Document
    from PySide6.QtCore import QCoreApplication, QEvent
    from PySide6.QtGui import QPixmap
    from PySide6.QtWidgets import QFrame, QLabel
    from src.database.repository import Repository
    from src.services.library import ArticleLibrary
    from src.services.learning_content import prepare_article, edit_hint
    from src.services.practice_service import new_full_session
    from src.services.grading import grade_session
    from src.ui.main_window import MainWindow
    from src.ui.dialogs import WrongWordsDialog, WrongPracticeHistoryDialog

    folder = root/'ui-performance-polish'; folder.mkdir()
    path = folder/'Context and Recall.docx'
    doc = Document(); doc.add_paragraph('Small Steps, Lasting Progress', 'Title')
    p = doc.add_paragraph('A clear '); p.add_run('context').bold = True
    p.add_run('（语境） makes a difficult word easier to remember.')
    p = doc.add_paragraph('She '); p.add_run('persisted').bold = True
    p.add_run('（坚持） with her daily practice, one small step at a time.')
    doc.save(path)
    repo = Repository(folder/'data'/'copy.db')
    article = ArticleLibrary(repo).import_file(path)
    window = MainWindow(repo); window.show()
    arrow = Path(__file__).resolve().parents[2] / 'assets' / 'chevron-down.svg'
    assert not QPixmap(str(arrow)).isNull() and arrow.as_posix() in window.styleSheet()
    book = WrongWordsDialog(repo, window)
    history = WrongPracticeHistoryDialog(repo, window)
    assert '暂无' in book.result_count.text() and '暂无' in history.result_count.text()
    book.close(); book.deleteLater(); history.close(); history.deleteLater()
    first = new_full_session(article); grade_session(first); repo.save_submitted_session(first)
    before = prepare_article(repo.load_article(article.source_path)); after = deepcopy(before)
    edit_hint(after, after.word_items[0].id, '语境', 'context 是具体语境；content 是内容。')
    repo.save_learning_content(before, after)
    with repo.connection() as db:
        db.execute('UPDATE wrong_words SET review_due_date=?', (repo.today().isoformat(),))
    import src.services.learning_content as content
    original = content.ContentResolver.__init__
    calls = []
    def counted(self, connection):
        calls.append(1); original(self, connection)
    try:
        content.ContentResolver.__init__ = counted
        repo.configure_review_capacity(20, 20)
    finally:
        content.ContentResolver.__init__ = original
    assert len(calls) == 1
    book = WrongWordsDialog(repo, window); book.show(); app.processEvents()
    book.search.setText('not-a-matching-word')
    assert '没有匹配' in book.result_count.text()
    assert not book.start_button.isEnabled() and not book.delete_button.isEnabled()
    book.grab().save(str(output.with_name(output.stem+'-empty-search.png')))
    book.close(); book.deleteLater()
    window._activate_session(new_full_session(repo.load_article(article.source_path)))
    for zoom in (70, 100, 130, 180):
        window.set_zoom(zoom); app.processEvents()
        notes = window.content.findChildren(QLabel, 'learningNote')
        assert notes and all(n.isHidden() and n.font().pixelSize() == round(15*zoom/100) for n in notes)
    window.set_zoom(130)
    window._commit_submission(); app.processEvents()
    assert '本轮答题正确' not in window.score_text.text()
    assert '原词回忆结果' in window.score_text.text()
    assert all(not n.isHidden() for n in window.content.findChildren(QLabel, 'learningNote'))
    window.grab().save(str(output.with_name(output.stem+'-compact-result.png')))
    session = window.session
    result_root, result_edits = window.content, dict(window.answer_edits)
    window.read_current_wrong(); app.processEvents()
    assert window._wrong_reading and window.session is session
    assert window.content is not result_root and not result_root.isVisible()
    assert window._result_page_cache is not None and not window.answer_edits
    assert all(n.font().pixelSize() == round(15*1.3) for n in window.content.findChildren(QLabel, 'learningNote'))
    window.grab().save(str(output.with_name(output.stem+'-secondary-note.png')))
    window.return_from_wrong_reading(); app.processEvents()
    assert window.session is session and session.submitted
    assert window.content is result_root and window._result_page_cache is None
    assert all(window.answer_edits[key] is edit for key, edit in result_edits.items())
    assert len(window.content.findChildren(QFrame, 'wrongSummary')) == 1
    for _ in range(3):
        window.read_current_wrong(); app.processEvents()
        window.return_from_wrong_reading(); app.processEvents()
        assert window.content is result_root and window._result_page_cache is None
        assert len(window.content.findChildren(QFrame, 'wrongSummary')) == 1
        assert window.reading_selection.root is result_root
    window.close(); window.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete); app.processEvents()
    return ['packaged_sort_dropdown_indicator', 'review_transaction_resolver_reuse', 'review_empty_state', 'review_search_no_match',
            'compact_result_preserves_first_round', 'secondary_note_zoom_and_visibility',
            'polished_reader_returns_same_results','result_page_reuses_exact_widgets',
            'cached_page_hidden_from_wrong_reader','bounded_single_result_cache',
            'repeated_return_has_one_wrong_summary','selection_rebinds_to_active_page']
