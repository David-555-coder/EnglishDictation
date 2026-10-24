"""Packaged checks of the reading surface and navigation affordances."""
from copy import deepcopy


def verify(root, app, output):
    from docx import Document
    from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QPoint, QPointF, Qt, QTimer
    from PySide6.QtGui import QKeyEvent, QMouseEvent
    from PySide6.QtWidgets import QFrame, QLabel, QPushButton
    from src.database.repository import Repository
    from src.services.library import ArticleLibrary
    from src.ui.main_window import MainWindow

    def settle(ms=60):
        loop = QEventLoop(); QTimer.singleShot(ms, loop.quit); loop.exec(); app.processEvents()

    def dispose(window):
        window.close(); window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete); app.processEvents()

    folder = root / 'reading-surface'; folder.mkdir()
    repo = Repository(folder / 'data' / 'copy.db')
    doc = Document(); doc.add_paragraph().add_run('A Quiet Morning').bold = True
    for i in range(24):
        p = doc.add_paragraph('Mia looked at the '); p.add_run('landscape').bold = True
        p.add_run('（风景） and thought about the changes in her town.')
    path = folder / '阅读示例.docx'; doc.save(path); ArticleLibrary(repo).import_file(path)
    window = MainWindow(repo)
    try:
        window.show(); window.activateWindow(); settle(250)
        assert window.session is None and not window.answer_edits
        article = deepcopy(window.current_article)
        assert not window.import_button.property('keyboardFocus')
        assert window.import_button.focusPolicy() & Qt.TabFocus
        page, = window.content.findChildren(QFrame, 'originalPage')
        paragraphs = page.findChildren(QFrame, 'readingParagraph')
        assert len(paragraphs) == 25 and not window.content.findChildren(QFrame, 'paragraphCard')
        assert [p.property('paragraphText') for p in paragraphs] == [p.text for p in article.paragraphs if p.text.strip()]
        assert window.scroll.verticalScrollBar().value() == 0
        window.set_zoom(160); settle()
        window.grab().save(str(output.with_name(output.stem+'-continuous-page.png')))
        window.import_button.setFocus()
        app.sendEvent(window.import_button, QKeyEvent(QEvent.KeyPress, Qt.Key_Tab, Qt.NoModifier))
        app.sendEvent(window.import_button, QKeyEvent(QEvent.KeyRelease, Qt.Key_Tab, Qt.NoModifier)); settle()
        assert app.focusWidget() is window.export_button and window.export_button.property('keyboardFocus')
        window.grab().save(str(output.with_name(output.stem+'-keyboard-focus.png')))
        pos = QPointF(2, 2); global_pos = QPointF(window.content.mapToGlobal(QPoint(2, 2)))
        app.sendEvent(window.content, QMouseEvent(QEvent.MouseButtonPress, pos, global_pos,
                                               Qt.LeftButton, Qt.LeftButton, Qt.NoModifier))
        app.sendEvent(window.content, QMouseEvent(QEvent.MouseButtonRelease, pos, global_pos,
                                               Qt.LeftButton, Qt.NoButton, Qt.NoModifier)); settle()
        assert not window.import_button.property('keyboardFocus') and not window.export_button.property('keyboardFocus')
        for width in (980, 1480, 2400):
            window.resize(width, 900)
            for zoom in (70, 100, 160, 180):
                window.set_zoom(zoom); settle(15)
                assert window.scroll.horizontalScrollBar().maximum() == 0
                assert page.findChildren(QFrame, 'readingParagraph') == paragraphs
                for p in paragraphs:
                    label = p.findChild(QLabel)
                    assert label.height() >= label.heightForWidth(label.width())
        window.resize(1480, 900); window.set_zoom(160); settle()
        assert not any('收起文章列表' in b.text() for b in window.findChildren(QPushButton))
        window.actions['显示文章列表'].trigger(); settle()
        assert window.library_shell.isHidden() and window.splitter.sizes()[0] == 0 and window.main_toolbar.isVisible()
        window.actions['专注模式'].trigger(); settle()
        window.actions['专注模式'].trigger(); settle()
        assert window.library_shell.isHidden() and window.splitter.sizes()[0] == 0
        window.actions['显示文章列表'].trigger(); settle()
        assert window.library_shell.isVisible() and window.splitter.sizes()[0] >= 200
        assert window.current_article == article and repo.load_drafts() == [] and repo.list_history() == []
        window.actions['显示文章列表'].trigger(); settle()
    finally:
        dispose(window)
    window = MainWindow(repo)
    try:
        window.show(); settle(250)
        assert window.library_collapsed and window.library_shell.isHidden() and window.splitter.sizes()[0] == 0
        window.actions['显示文章列表'].trigger(); settle()
        assert window.library_shell.isVisible() and not window.library_collapsed
    finally:
        dispose(window)
    return ['continuous_original_page', 'original_paragraphs_and_text_preserved', 'page_zoom_and_focus_layout',
            'no_full_width_collapse_button', 'menu_and_saved_library_visibility', 'quiet_initial_toolbar_focus',
            'keyboard_toolbar_navigation', 'mouse_clears_toolbar_outline']
