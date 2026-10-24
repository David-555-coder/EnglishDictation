"""Isolated packaged checks. Synthetic text, public lexicon, loopback network."""
from copy import deepcopy
import json
from pathlib import Path
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlsplit


def verify(root, app, output):
    from docx import Document
    from PySide6.QtCore import QPoint, Qt
    from PySide6.QtGui import QFontDatabase
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QPushButton
    from src.database.repository import Repository
    from src.services.serialization import session_to_dict
    from src.ui.main_window import MainWindow
    from src.ui.widgets.text_menus import TextLabel, plain_label_text, text_menu
    from src.ui.widgets.reading_selection import units

    folder = root / 'translation'; folder.mkdir()
    path = folder / 'reading.docx'
    doc = Document()
    p = doc.add_paragraph('I was '); p.add_run('utterly').bold = True
    p.add_run('（完全地） exhausted. Please do not spill the beans.')
    doc.save(path)
    window = MainWindow(Repository(folder / 'test.db'))
    window._article_switch_delay = 0
    window._import_paths([str(path)])
    window.resize(1100, 850); window.show(); app.processEvents()
    checks = []

    def wait(condition):
        limit = time.monotonic() + 5
        while not condition() and time.monotonic() < limit:
            QTest.qWait(10)
        assert condition()

    def menu_for(text):
        app.processEvents(); window.reading_selection.clear()
        label = next(w for w in window.content.findChildren(TextLabel)
                     if w.isVisibleTo(window.content) and text in plain_label_text(w))
        full = plain_label_text(label); start = full.index(text)
        label.setSelection(units(full[:start]), units(text))
        return text_menu(label)

    def lookup(text):
        menu = menu_for(text)
        try: next(a for a in menu.actions() if a.text() == '翻译').trigger()
        finally: menu.deleteLater()
        wait(lambda: window._reading_translation.popup.lookup_done)
        return window._reading_translation

    try:
        assert window._reading_translation is None
        checks.append('translation_no_clients_at_startup')
        controller = lookup('utterly')
        assert '完全' in controller.popup.result.toPlainText()
        assert controller.network is None
        checks += ['translation_public_dictionary_sidecar', 'translation_default_zero_network']
        assert controller.popup.width() <= 360 and controller.popup.height() < 250
        assert not controller.popup.isModal()
        checks += ['translation_compact_popup', 'translation_non_modal']
        controller.popup.grab().save(str(output.with_name(output.stem + '-translation.png')))
        window.show_preview(); assert not controller.popup.isVisible()
        checks.append('translation_switch_cancels')
        controller = lookup('spill the beans')
        assert '说漏嘴' in controller.popup.result.toPlainText()
        checks.append('translation_whole_idiom')

        requests = []
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args): pass
            def do_GET(self):
                text = parse_qs(urlsplit(self.path).query)['q'][0]; requests.append(text)
                body = json.dumps({'responseStatus': 200, 'responseData': {'translatedText': '回环测试译文'}}).encode()
                self.send_response(200); self.send_header('Content-Length', str(len(body)))
                self.end_headers(); self.wfile.write(body)
        server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        server.daemon_threads = True
        thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
        try:
            controller.online_action()
            assert controller.network is None and controller.popup.net_state == 'consent'
            checks.append('translation_explicit_inline_consent')
            from src.services.reading_translation import MyMemory, TranslationClient
            provider = MyMemory(); provider.endpoint = f'http://127.0.0.1:{server.server_port}/'
            controller.network = TranslationClient(controller, provider=provider)
            controller.network.completed.connect(controller._network_completed)
            controller.online_action(); wait(lambda: controller.popup.net_state == 'done')
            assert requests == ['Please do not spill the beans.']
            assert '回环测试译文' in controller.popup.net_text.toPlainText()
            checks.append('translation_packaged_async_network_loopback')
            controller.popup.copy.click();assert '说漏嘴' in app.clipboard().text()
            controller.popup.net_copy.click();assert app.clipboard().text()=='回环测试译文'
            checks.append('translation_explicit_copy_targets')
            controller.popup.grab().save(str(output.with_name(output.stem+'-translation-online.png')))
            controller.online_action(); assert not controller.popup.network.isVisible()
            assert '说漏嘴' in controller.popup._copy_text
            checks.append('translation_online_does_not_replace_dictionary')
        finally:
            server.shutdown(); server.server_close(); thread.join(timeout=2)

        assert controller.popup._outside_filter_installed
        button=QPushButton('outside',window);button.setGeometry(10,10,80,30);button.show()
        clicks=[];button.clicked.connect(lambda:clicks.append(True))
        QTest.mouseClick(button,Qt.LeftButton)
        assert clicks==[True] and not controller.popup.isVisible()
        assert not controller.popup._outside_filter_installed
        checks+=['translation_outside_first_click_preserved','translation_hidden_listener_removed']
        controller=lookup('utterly')
        QTest.mouseClick(controller.popup.copy,Qt.LeftButton)
        assert controller.popup.isVisible() and '完全' in app.clipboard().text()
        checks.append('translation_inside_click_keeps_popup')

        window.start_full_practice(); app.processEvents()
        assert not controller.popup.isVisible()
        before = deepcopy(session_to_dict(window.session))
        menu = menu_for('exhausted')
        assert '翻译' not in [a.text() for a in menu.actions()]; menu.deleteLater()
        window._translate_reading_selection(None, QPoint(20, 20))
        assert not controller.popup.isVisible()
        assert session_to_dict(window.session) == before
        checks += ['translation_independent_recall_disabled', 'translation_no_learning_mutation']
        for edit in window.answer_edits.values():
            edit.setText('incorrect')  # Exercise normal submit without a modal blank-answer prompt.
        window._confirm_finished_submission = lambda: True
        window.submit(); app.processEvents()
        before = deepcopy(session_to_dict(window.session))
        controller = lookup('exhausted')
        assert controller.popup.lookup_done and session_to_dict(window.session) == before
        checks.append('translation_submitted_results_reading')
        window.read_current_wrong(); app.processEvents(); lookup('exhausted')
        checks.append('translation_wrong_sentence_reading')
    finally:
        window.close(); window.deleteLater(); app.processEvents()
    return checks
