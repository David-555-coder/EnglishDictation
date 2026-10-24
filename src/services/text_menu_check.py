"""Exercise real Qt menus with synthetic data, including packaged releases."""
from pathlib import Path


def verify(root, app, output):
    from docx import Document
    from PySide6.QtCore import Qt, QEvent, QPoint, QTimer, QCoreApplication
    from PySide6.QtGui import QContextMenuEvent, QPalette
    from PySide6.QtWidgets import QApplication, QLabel, QLineEdit, QMenu, QInputDialog
    from src.database.repository import Repository
    from src.services.practice_service import new_wrong_session
    from src.ui.main_window import MainWindow
    from src.ui.editing_dialogs import TargetEditorDialog
    from src.ui.quick_write_dialogs import QuickWriteEditor, QuickWritePanel
    from src.ui.dialogs import WrongWordsDialog
    from src.ui.daily_review_dialog import ReviewQuantityDialog
    from src.ui.widgets.text_menus import text_menu, TextLabel, bind_text_menus

    folder = root/'text-menus'; folder.mkdir()
    path = folder/'A Calm Reading Page.docx'
    doc = Document(); doc.add_paragraph('A Calm Reading Page', 'Title')
    p = doc.add_paragraph('Mia '); p.add_run('woke up').bold = True
    p.add_run('（醒来） from a '); p.add_run('nightmare').bold = True
    p.add_run('（噩梦）. A quiet room helped her relax.'); doc.save(path)
    repo = Repository(folder/'data'/'db')
    window = MainWindow(repo); window._article_switch_delay=0
    window._import_paths([str(path)]); window.show(); app.processEvents()
    checks=[]

    def check(widget, name, screenshot=False):
        widget.ensurePolished()
        if isinstance(widget,QLabel):
            widget.setSelection(0, 3)
        elif isinstance(widget,QLineEdit):
            widget.selectAll()
        else:
            widget.selectAll()
        menu=text_menu(widget); menu.ensurePolished()
        assert menu.font().pixelSize()==14 and not menu.font().bold(), name
        assert menu.palette().color(QPalette.Window).name()=='#ffffff', name
        assert all(a.font().pixelSize()==14 for a in menu.actions()), name
        assert not any('修正' in a.text() or 'Copy' in a.text() or 'Select All' in a.text() for a in menu.actions()), name
        copy=next(a for a in menu.actions() if a.text()=='复制')
        if screenshot:
            menu.popup(widget.mapToGlobal(QPoint(12,12)));menu.setActiveAction(copy);app.processEvents()
            menu.grab().save(str(output.with_name(output.stem+'-'+name+'.png')));menu.close()
        menu.deleteLater();checks.append(name)

    for zoom in (70,100,160,180):
        window.set_zoom(zoom);app.processEvents()
        label=next(l for l in window.content.findChildren(TextLabel) if 'nightmare' in l.text())
        check(label, 'menu_original_'+str(zoom), screenshot=zoom==160)
    window.start_full_practice();app.processEvents()
    check(next(iter(window.answer_edits.values())), 'menu_active_answer')
    window.show_answers();app.processEvents()
    check(next(iter(window.answer_hints.values())), 'menu_blue_answer_hint', True)
    window.show_answers()
    for n, edit in enumerate(window.answer_edits.values()):
        edit.setText(window.session.article.item_map[edit.item_id].answer if n==0 else 'wrong')
    window._commit_submission();app.processEvents()
    for n, edit in enumerate(window.answer_edits.values()):
        check(edit, 'menu_submitted_'+('correct' if n==0 else 'wrong'))
        menu=text_menu(edit)
        assert all(not a.isEnabled() for a in menu.actions() if a.text() in ('剪切','粘贴','删除'))
        menu.deleteLater()
    visible_hint=next(h for h in window.answer_hints.values() if not h.isHidden())
    check(visible_hint,'menu_red_result_hint',True)
    window.read_current_wrong();app.processEvents()
    blue=next(l for l in window.content.findChildren(TextLabel) if 'nightmare' in l.text())
    check(blue,'menu_wrong_reading',True)
    window._show_wrong_preview(new_wrong_session(window.session));app.processEvents()
    red=next(l for l in window.content.findChildren(TextLabel) if 'nightmare' in l.text())
    check(red,'menu_wrong_preview')
    from src.services.daily_review import new_daily_session
    from src.services.practice_service import new_long_term_session
    with repo.connection() as db:
        ids=[r[0] for r in db.execute('SELECT id FROM wrong_words')]
    records=repo.get_review_examples(ids)
    window._activate_session(new_daily_session(records,repo.today().isoformat()));app.processEvents()
    check(next(iter(window.answer_edits.values())),'menu_daily_review')
    window._activate_session(new_long_term_session(records));app.processEvents()
    check(next(iter(window.answer_edits.values())),'menu_long_term_practice')

    dialogs=[TargetEditorDialog(repo.load_article(str(path.resolve())),window),
             QuickWriteEditor(repo,window),QuickWritePanel(repo,window),WrongWordsDialog(repo,window),
             ReviewQuantityDialog(40,10,repo.continuity_status(),window)]
    for dialog in dialogs:bind_text_menus(dialog)
    editor,quick,quick_panel,words,quantity=dialogs
    check(editor.text,'menu_editor_preview')
    check(editor.chinese_edit,'menu_editor_hint')
    editor.note_edit.setPlainText('internal / inherent');check(editor.note_edit,'menu_editor_note')
    editor._edit_paragraph();check(editor.paragraph_edit,'menu_editor_body')
    editor._leave_paragraph()
    check(quick.name_edit,'menu_quick_name')
    check(quick.prompt_edit,'menu_quick_prompt')
    check(quick_panel.filename,'menu_quick_filename')
    check(words.search,'menu_wordbook_search')
    check(window.article_search,'menu_article_search')
    check(quantity.findChild(QLineEdit),'menu_quantity_native_field')

    # A context event must really reach the common menu, not just its factory.
    captured=[]
    def inspect():
        menu=QApplication.activePopupWidget()
        if isinstance(menu,QMenu):
            captured.append(menu.font().pixelSize())
            next(a for a in menu.actions() if a.text()=='复制').trigger();menu.close()
    editor.show();editor.activateWindow();editor.chinese_edit.setFocus();app.processEvents()
    editor.chinese_edit.setText('保持注意力');editor.chinese_edit.selectAll()
    QTimer.singleShot(0,inspect)
    QApplication.sendEvent(editor.chinese_edit,QContextMenuEvent(QContextMenuEvent.Mouse,QPoint(5,5),editor.chinese_edit.mapToGlobal(QPoint(5,5))))
    assert captured==[14] and QApplication.clipboard().text()=='保持注意力', (captured,QApplication.clipboard().text(),editor.chinese_edit.selectedText())
    checks.append('menu_real_context_copy')
    captured.clear();quick_panel.show();quick_panel.activateWindow();quick_panel.filename.setFocus();app.processEvents()
    quick_panel.filename.setText('sample');quick_panel.filename.selectAll()
    QTimer.singleShot(0,inspect)
    QApplication.sendEvent(quick_panel.filename,QContextMenuEvent(QContextMenuEvent.Mouse,QPoint(5,5),quick_panel.filename.mapToGlobal(QPoint(5,5))))
    assert captured==[14] and QApplication.clipboard().text()=='sample'
    checks.append('menu_popup_context_copy')
    for dialog in dialogs:
        dialog.hide();dialog.deleteLater()
    window.close();window.deleteLater()
    QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete);app.processEvents()
    return checks
