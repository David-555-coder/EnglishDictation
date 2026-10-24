"""Packaged UI regression checks using only generated articles and a temporary DB."""
from pathlib import Path


def verify(root, app, output):
    from docx import Document
    from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QPoint, Qt, QTimer
    from PySide6.QtGui import QKeyEvent
    from src.database.repository import Repository
    from src.ui.main_window import MainWindow

    def settle(milliseconds=60):
        loop = QEventLoop()
        QTimer.singleShot(milliseconds, loop.quit)
        loop.exec()
        app.processEvents()

    def dispose(window):
        window.close()
        window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        app.processEvents()

    def offset(window, item_id):
        return (window.answer_edits[item_id].mapTo(window.content, QPoint()).y()
                - window.scroll.verticalScrollBar().value())

    directory = Path(root) / 'scroll-position'
    directory.mkdir()
    path = directory / 'generated-long-article.docx'
    doc = Document()
    for index in range(60):
        paragraph = doc.add_paragraph(f'Sentence {index}: ')
        paragraph.add_run('take advantage of').bold = True
        paragraph.add_run('（利用） this opportunity with enough context to read.')
    doc.save(path)
    repo = Repository(directory / 'data' / 'copy.db')
    window = MainWindow(repo)
    try:
        window.show(); window.activateWindow(); settle(180)
        window._article_switch_delay = 0
        window._import_paths([str(path)])
        window.start_full_practice(); settle()
        # Offscreen Qt has no window manager to reactivate the main window
        # after the import progress window closes.
        window.activateWindow(); settle()
        item_id = window.session.target_ids[28]
        edit = window.answer_edits[item_id]
        edit.setText('saved draft')
        edit.setFocus(); settle()
        before = offset(window, item_id)
        assert window.scroll.verticalScrollBar().value() > 0, {
            'ready': window._reading_layout_ready(),
            'content_height': window.content.height(),
            'flow_height': window.content_layout.totalHeightForWidth(window.content.width()),
            'viewport_height': window.scroll.viewport().height(),
            'range': window.scroll.verticalScrollBar().maximum(),
            'pending_restore': window._pending_view_restore,
            'pending_follow': window._pending_follow,
            'focused': app.focusWidget() is edit,
            'offset': before,
        }
        for _ in range(6):
            window.actions['专注模式'].trigger(); settle()
            assert abs(offset(window, item_id) - before) <= 3
            assert app.focusWidget() is edit
        window.grab().save(str(output.with_name(output.stem + '-focus-scroll.png')))
        for zoom in (130, 100):
            window.set_zoom(zoom); settle()
            assert abs(offset(window, item_id) - before) <= 3
        for _ in range(2):
            window.show_answers(); settle()
            assert abs(offset(window, item_id) - before) <= 3
        key = window.session.session_key
        window.show_preview(); settle()
        draft = next(s for s in Repository(repo.database_path).load_drafts() if s.session_key == key)
        window._activate_session(draft); settle(150)
        assert window.session.session_key == key
        assert window.answer_edits[item_id].text() == 'saved draft'
        assert app.focusWidget() is window.answer_edits[item_id]
        assert abs(offset(window, item_id) - before) <= 3
        window.grab().save(str(output.with_name(output.stem + '-resumed-scroll.png')))
    finally:
        dispose(window)

    reopened = MainWindow(Repository(repo.database_path))
    try:
        reopened.show(); reopened.activateWindow(); settle(240)
        assert reopened.session is None and reopened.mode_title.text().startswith('原文模式')
        draft = next(s for s, _ in reopened.repository.unfinished_practices() if s.session_key == key)
        reopened._activate_session(draft); settle(150)
        assert reopened.session.session_key == key
        assert reopened.answer_edits[item_id].text() == 'saved draft'
        assert app.focusWidget() is reopened.answer_edits[item_id]
        assert abs(offset(reopened, item_id) - before) <= 3, {
            'before': before, 'after': offset(reopened, item_id),
            'saved': reopened.session.view_state,
            'bar': reopened.scroll.verticalScrollBar().value(),
            'ready': reopened._reading_layout_ready(),
            'pending_restore': reopened._pending_view_restore,
            'pending_follow': reopened._pending_follow,
            'timers': (reopened._layout_timer.isActive(), reopened._follow_timer.isActive()),
            'viewport': (reopened.scroll.viewport().width(), reopened.scroll.viewport().height()),
        }
        generation, revision = reopened._content_generation, reopened._view_restore_revision
        bar = reopened.scroll.verticalScrollBar()
        QCoreApplication.sendEvent(bar, QKeyEvent(QEvent.KeyPress, Qt.Key_Home, Qt.NoModifier))
        bar.setValue(0)
        reopened._restore_view(key, generation, revision); settle()
        assert bar.value() == 0
        assert reopened._pending_view_restore is None
        probes = []
        ready = reopened._reading_layout_ready
        reopened._reading_layout_ready = lambda: (probes.append(1), ready())[1]
        settle(180)
        assert not probes
        assert not reopened._layout_timer.isActive() and not reopened._follow_timer.isActive()
    finally:
        dispose(reopened)
    return ['focus_mode_scroll_anchor', 'zoom_answer_scroll_anchor',
            'draft_scroll_restore', 'reopen_draft_scroll_restore',
            'manual_scroll_overrides_late_restore', 'no_idle_scroll_polling']
