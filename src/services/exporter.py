from __future__ import annotations

import os
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run

from src.models.article import Article, RunStyle
from src.parser.word_structure import iter_paragraphs, iter_runs, is_heading


def _apply_style(run, style: RunStyle) -> None:
    if style.font_name:
        run.font.name = style.font_name
    if style.font_size_pt:
        run.font.size = Pt(style.font_size_pt)
    run.bold, run.italic, run.underline = style.bold, style.italic, style.underline
    if style.color_rgb:
        try:
            run.font.color.rgb = RGBColor.from_string(style.color_rgb[-6:])
        except ValueError:
            pass


def protect_output(output_path, protected_paths):
    output = Path(output_path).resolve()
    for source in protected_paths:
        if not source: continue
        path = Path(source).resolve()
        if output == path or (output.exists() and path.exists() and output.samefile(path)):
            raise ValueError("不能覆盖原始 Word 或本地文章副本，请换一个文件名或导出目录。")
    if output.suffix.lower() != ".docx":
        raise ValueError("请将文件保存为 .docx 格式。")
    return output


def _replace_span(paragraph, start, end, replacement, answer_mode):
    cursor = 0
    overlaps = []
    for run in iter_runs(paragraph):
        text = run.text
        left, right = max(start, cursor), min(end, cursor + len(text))
        if left < right:
            allowed = {qn("w:" + name) for name in ("rPr", "t", "tab", "br", "cr", "lastRenderedPageBreak")}
            if any(child.tag not in allowed for child in run._r):
                raise ValueError("目标词与图片或特殊 Word 对象共用一个文本块，请将它们分开后再导出。")
            overlaps.append((run, left - cursor, right - cursor))
        cursor += len(text)
    for number, (run, left, right) in enumerate(overlaps):
        text = run.text
        middle = text[left:right] if answer_mode else (replacement if number == 0 else "")
        for part, marked in ((text[:left], False), (middle, True), (text[right:], False)):
            if not part: continue
            element = OxmlElement("w:r")
            if run._r.rPr is not None: element.append(deepcopy(run._r.rPr))
            fragment = Run(element, paragraph)
            fragment.text = part
            if marked:
                fragment.bold = bool(answer_mode)
                fragment.font.color.rgb = RGBColor(0, 0, 0)
            run._r.addprevious(element)
        run._r.getparent().remove(run._r)


def _plain_empty_paragraph(paragraph):
    """Only remove whitespace-only text paragraphs, never objects or section breaks."""
    allowed = {qn("w:" + name) for name in ("p", "pPr", "r", "rPr", "t", "tab")}
    return not paragraph.text.strip() and all(
        node.tag in allowed or any(parent.tag in {qn("w:pPr"), qn("w:rPr")}
                                   for parent in node.iterancestors())
        for node in paragraph._p.iter()
    ) and paragraph._p.find('.//' + qn('w:sectPr')) is None


def _format_print_copy(document):
    """Normalize only the export copy, after all offset-based replacements."""
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(22)
        section.left_margin = section.right_margin = Mm(24)
    for paragraph in list(iter_paragraphs(document)):
        if _plain_empty_paragraph(paragraph):
            # Table cells must end in a paragraph. Keep their structural blanks tiny.
            if paragraph._p.getparent().tag == qn("w:body"):
                paragraph._p.getparent().remove(paragraph._p)
                continue
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            for run in iter_runs(paragraph): run.font.size = Pt(1)
            continue
        heading = is_heading(paragraph)
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(6 if heading else 0)
        fmt.space_after = Pt(14)
        fmt.line_spacing = 1.3 if heading else 1.6
        fmt.keep_with_next = heading
        # Keep ordinary study paragraphs intact so blanks and their Chinese
        # hints don't land on opposite pages; very long paragraphs can still split.
        fmt.keep_together = heading or len(paragraph.text) <= 600
        fmt.widow_control = True
        fmt.page_break_before = False
        fmt.left_indent = fmt.right_indent = fmt.first_line_indent = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ppr = paragraph._p.get_or_add_pPr()
        indent = ppr.find(qn('w:ind'))
        for key in ('firstLineChars', 'hangingChars', 'leftChars', 'rightChars', 'hanging'):
            indent.attrib.pop(qn('w:' + key), None)
        # A large source paragraph-mark font can enlarge every line even when
        # all visible runs have been normalized.
        mark_style = ppr.find(qn('w:rPr'))
        if mark_style is not None:
            for name in ('sz', 'szCs'):
                element = mark_style.find(qn('w:' + name))
                if element is not None: element.set(qn('w:val'), str(34 if heading else 26))
        # Direct sizes take precedence over arbitrary source styles and document grids.
        grid = paragraph._p.get_or_add_pPr().find(qn("w:snapToGrid"))
        if grid is None:
            grid = OxmlElement("w:snapToGrid"); paragraph._p.get_or_add_pPr().append(grid)
        grid.set(qn("w:val"), "0")
        for run in iter_runs(paragraph):
            run.font.name = "Calibri"
            run.font.size = Pt(17 if heading else 13)
            rpr = run._r.get_or_add_rPr()
            cs_size = rpr.find(qn('w:szCs'))
            if cs_size is not None: cs_size.set(qn('w:val'), str(34 if heading else 26))
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + key), None)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if heading: run.bold = True
    body = [p for p in document.paragraphs if p.text.strip()]
    if len(body) >= 2 and len(body[-1].text) <= 80 and len(body[-2].text) <= 600:
        # Avoid a final one-sentence page without squeezing the note-taking space.
        body[-2].paragraph_format.keep_with_next = True


def _export(article, output_path, answer_mode=False, protected_paths=()):
    source = article.metadata.get("document_path") or article.source_path
    output = protect_output(output_path, [source, article.source_path, *protected_paths])
    if not source or not Path(source).is_file():
        raise ValueError("找不到原 Word 或本地副本，请重新导入文章后导出。")
    document = Document(source)
    paragraphs = list(iter_paragraphs(document))
    for paragraph in article.paragraphs:
        if paragraph.index >= len(paragraphs) or paragraphs[paragraph.index].text != paragraph.text:
            raise ValueError("Word 内容已发生变化，请重新导入，避免导出位置错误。")
    for item in sorted(article.word_items, key=lambda i: (i.paragraph_index, i.start), reverse=True):
        if paragraphs[item.paragraph_index].text[item.start:item.end] != item.answer:
            raise ValueError("目标词位置与 Word 内容不一致，无法安全导出。")
        _replace_span(paragraphs[item.paragraph_index], item.start, item.end,
                      "_" * max(8, min(40, round(len(item.answer) * 1.6))), answer_mode)
    _format_print_copy(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(suffix=".docx", dir=output.parent)
    os.close(handle)
    try:
        document.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary): os.unlink(temporary)
    return output


def export_worksheet(article: Article, output_path: str | Path, protected_paths=()) -> Path:
    return _export(article, output_path, False, protected_paths)


def export_answer(article: Article, output_path: str | Path, protected_paths=()) -> Path:
    return _export(article, output_path, True, protected_paths)
