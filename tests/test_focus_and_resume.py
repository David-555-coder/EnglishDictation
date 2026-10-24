"""Interaction regression tests use temporary libraries, never the user's data."""
import pytest
from docx import Document
from PySide6.QtCore import QPoint, Qt, QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox, QDialog

from src.database.repository import Repository
from src.ui.widgets.flow_layout import FlowLayout
from src.ui.unfinished_dialog import UnfinishedPracticeDialog
from test_library_ui import ui, sample_word

REAL_SINGLE_SHOT = QTimer.singleShot


def practice(tmp_path, ui, count=1):
    path = sample_word(tmp_path)
    if count > 1:
        document = Document()
        for index in range(count):
            paragraph = document.add_paragraph(f"Sentence {index}: ")
            paragraph.add_run("take advantage of").bold = True
            paragraph.add_run("（利用） this opportunity with enough context to read.")
        document.save(path)
    window = ui(Repository(tmp_path / "learning.db"))
    window._import_paths([str(path)])
    window.start_full_practice()
    window.show()
    window.activateWindow()
    QApplication.processEvents()
    return window


@pytest.mark.parametrize("collapsed,toolbar_visible", [(False, True), (True, False), (False, False)])
def test_focus_mode_restores_layout_and_keeps_menu_and_input(tmp_path, ui, collapsed, toolbar_visible):
    window = practice(tmp_path, ui)
    edit = next(iter(window.answer_edits.values()))
    edit.setText("keep this answer")
    window._set_library_collapsed(collapsed)
    window.main_toolbar.setVisible(toolbar_visible)
    QApplication.processEvents()
    sizes = window.splitter.sizes()
    edit.setFocus()
    edit.setSelection(2, 4)
    key = window.session.session_key
    window.actions["专注模式"].trigger()
    QApplication.processEvents()
    assert window.menuBar().isVisible()
    assert window.view_menu.menuAction() in window.menuBar().actions()
    assert window.main_toolbar.isHidden() and window.library_shell.isHidden()
    assert window.status_panel.isVisible()
    assert not window.actions["显示主工具栏"].isEnabled()
    assert window.actions["未完成练习…"] in window.practice_menu.actions()
    window.actions["专注模式"].trigger()
    QApplication.processEvents()
    assert window.library_collapsed == collapsed
    assert window.main_toolbar.isVisible() == toolbar_visible
    assert window.library_shell.isVisible() == (not collapsed)
    assert abs(window.splitter.sizes()[0] - sizes[0]) <= 5
    assert window.session.session_key == key
    assert next(iter(window.answer_edits.values())) is edit
    assert edit.text() == "keep this answer" and edit.selectedText() == "ep t"
    assert QApplication.focusWidget() is edit


def test_focus_shortcut_and_saved_visibility_are_independent(tmp_path, ui):
    window = practice(tmp_path, ui)
    window.actions["显示主工具栏"].trigger()
    window._set_library_collapsed(True)
    QApplication.processEvents()
    assert window.main_toolbar.isHidden()
    QTest.keyClick(window, Qt.Key_F8)
    QApplication.processEvents()
    assert window.actions["专注模式"].isChecked()
    window.close()
    reopened = ui(Repository(tmp_path / "learning.db"))
    reopened.show()
    QApplication.processEvents()
    assert not reopened.actions["专注模式"].isChecked()
    assert reopened.library_collapsed and reopened.main_toolbar.isHidden()


def test_preview_returns_to_same_round_and_restart_is_separate(tmp_path, ui, monkeypatch):
    window = practice(tmp_path, ui, 24)
    item_id = window.session.target_ids[18]
    edit = window.answer_edits[item_id]
    edit.setText("unfinished long answer")
    edit.setFocus()
    window._center_answer(item_id)
    old_key = window.session.session_key
    old_scroll = window.scroll.verticalScrollBar().value()
    window.show_preview()
    assert window.session is None
    assert window.primary_action_button.text() == "开始默写"
    assert not hasattr(window, "restart_action_button")
    monkeypatch.setattr(UnfinishedPracticeDialog, "exec", lambda _: QDialog.Accepted)
    window.resume_draft()
    QApplication.processEvents()  # The production restore runs after the layout has settled.
    window._restore_view(old_key)
    QApplication.processEvents()
    assert window.session.session_key == old_key
    assert window.answer_edits[item_id].text() == "unfinished long answer"
    assert QApplication.focusWidget() is window.answer_edits[item_id]
    assert abs(window.scroll.verticalScrollBar().value() - old_scroll) <= 3
    window.show_preview()
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.No)
    window.primary_action_button.click()
    assert window.session is None
    assert window._drafts[old_key].answers[item_id] == "unfinished long answer"
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.Yes)
    window.primary_action_button.click()
    assert window.session.session_key != old_key
    assert window.session.filled == 0
    drafts = {session.session_key: session for session in window.repository.load_drafts()}
    assert drafts[old_key].answers[item_id] == "unfinished long answer"
    assert not window.repository.list_history()


def test_failed_save_does_not_leave_practice_for_preview(tmp_path, ui, monkeypatch):
    window = practice(tmp_path, ui)
    old = window.session
    next(iter(window.answer_edits.values())).setText("not saved yet")
    monkeypatch.setattr(window.repository, "save_draft", lambda *args: (_ for _ in ()).throw(OSError("unavailable")))
    window.show_preview()
    assert window.session is old
    assert window.primary_action_button.text() == "提交批改"
    monkeypatch.undo()


def test_capped_field_keeps_restored_text_without_relayout_on_each_keystroke(tmp_path, ui, monkeypatch):
    window = practice(tmp_path, ui)
    edit = next(iter(window.answer_edits.values()))
    assert edit.placeholderText() == ""
    restored = "a longer draft"
    edit.setText(restored)
    window.show_preview()
    monkeypatch.setattr(UnfinishedPracticeDialog, "exec", lambda _: QDialog.Accepted)
    window.resume_draft()
    window.set_zoom(130)
    QApplication.processEvents()
    edit = next(iter(window.answer_edits.values()))
    # Long drafts stay intact and scroll inside a compact field; restoring one
    # must not undo the single-word width cap.
    assert edit.text() == restored and window.session.answers[edit.item_id] == restored
    assert round(120 * 1.3) <= edit.width() <= round(170 * 1.3)
    width = edit.property("preferredFlowWidth")
    edit.setText("x" * 600)
    QApplication.processEvents()
    assert edit.property("preferredFlowWidth") == width
    assert edit.text() == "x" * 600 and window.session.answers[edit.item_id] == "x" * 600
    assert edit.maximumWidth() == round(170 * 1.3)


def test_focus_follow_coalesces_and_manual_scroll_cancels_pending_follow(tmp_path, ui):
    window = practice(tmp_path, ui, 70)
    window.resize(1000, 650)
    QApplication.processEvents()
    ids = window.session.target_ids
    window.answer_edits[ids[1]].setFocus()
    for index in range(1, 30):
        window._move_from(ids[index], 1)
    QApplication.processEvents()
    last = window.answer_edits[ids[30]]
    assert QApplication.focusWidget() is last
    relative = last.mapTo(window.content, QPoint(0, 0)).y() - window.scroll.verticalScrollBar().value()
    assert 0 <= relative <= window.scroll.viewport().height() - last.height()
    window.answer_edits[ids[-1]].setFocus()
    window._cancel_pending_follow()
    window.scroll.verticalScrollBar().setValue(0)
    QApplication.processEvents()
    assert window.scroll.verticalScrollBar().value() == 0
    window._move_from(ids[-1], -1)
    QApplication.processEvents()
    assert window.scroll.verticalScrollBar().value() > 0


def test_focus_layout_caps_wide_reading_and_reflows_small_window(tmp_path, ui):
    window = practice(tmp_path, ui, 3)
    window.actions["专注模式"].trigger()
    window.resize(2500, 1000)
    QApplication.processEvents()
    margins = window.content_layout.contentsMargins()
    assert margins.left() > 4 and margins.right() > 12
    width = window.content.width() - margins.left() - margins.right()
    assert width <= round(1000 * window.zoom_percent / 100) + 2
    window.resize(1000, 650)
    window.set_zoom(180)
    QApplication.processEvents()
    assert window.scroll.horizontalScrollBar().maximum() == 0
    for flow in window.content.findChildren(FlowLayout):
        for index in range(flow.count()):
            widget = flow.itemAt(index).widget()
            if widget and not widget.isHidden():
                assert widget.geometry().right() <= flow.parentWidget().width()


def test_history_wrong_action_label(tmp_path, ui):
    window = ui(Repository(tmp_path / "learning.db"))
    assert window.actions["历史错词练习"].shortcut().toString() == "Ctrl+R"
    assert "重新练习错词" not in [action.text() for action in window.main_toolbar.actions()]


@pytest.mark.parametrize("change", ["answers", "zoom", "focus", "toolbar"])
def test_layout_changes_keep_the_current_word_at_the_same_height(tmp_path, ui, change):
    window = practice(tmp_path, ui, 40)
    window.resize(1000, 650)
    QApplication.processEvents()
    item_id = window.session.target_ids[24]
    edit = window.answer_edits[item_id]
    edit.setFocus()
    QTest.qWait(20)
    bar = window.scroll.verticalScrollBar()
    before = edit.mapTo(window.content, QPoint()).y() - bar.value()
    if change == "answers":
        window.show_answers()
    elif change == "zoom":
        window.set_zoom(130)
    elif change == "focus":
        window.actions["专注模式"].trigger()
    else:
        window.actions["显示主工具栏"].trigger()
    QTest.qWait(30)
    after = edit.mapTo(window.content, QPoint()).y() - bar.value()
    assert abs(after - before) <= 3
    assert QApplication.focusWidget() is edit and bar.value() > 0
    if change == "answers":
        window.show_answers()
    elif change == "zoom":
        window.set_zoom(100)
    elif change == "focus":
        window.actions["专注模式"].trigger()
    else:
        window.actions["显示主工具栏"].trigger()
    QTest.qWait(30)
    assert abs(edit.mapTo(window.content, QPoint()).y() - bar.value() - before) <= 3


def test_manual_scrollbar_input_cancels_queued_follow(tmp_path, ui):
    window = practice(tmp_path, ui, 30)
    window.answer_edits[window.session.target_ids[-1]].setFocus()
    bar = window.scroll.verticalScrollBar()
    # A real interaction event must cancel the pending request, not just the helper.
    QTest.keyClick(bar, Qt.Key_Home)
    bar.setValue(0)
    QApplication.processEvents()
    assert bar.value() == 0 and window._pending_follow is None


def test_old_delayed_restore_cannot_move_a_new_render_of_the_same_round(tmp_path, ui, monkeypatch):
    window = practice(tmp_path, ui, 30)
    callbacks = []
    monkeypatch.setattr(QTimer, "singleShot", lambda delay, callback: callbacks.append(callback))
    window._render_practice(True)
    old_restore = callbacks[-1]
    window.show_preview()
    monkeypatch.setattr(UnfinishedPracticeDialog, "exec", lambda _: QDialog.Accepted)
    window.resume_draft()
    QApplication.processEvents()
    edit = window.answer_edits[window.session.target_ids[-1]]
    edit.setFocus()
    QApplication.processEvents()
    before = window.scroll.verticalScrollBar().value()
    old_restore()
    QApplication.processEvents()
    assert QApplication.focusWidget() is edit
    assert window.scroll.verticalScrollBar().value() == before


def test_real_delayed_restore_and_restart_survive_reopening(tmp_path, ui, monkeypatch):
    window = practice(tmp_path, ui, 30)
    monkeypatch.setattr(QTimer, "singleShot", REAL_SINGLE_SHOT)
    item_id = window.session.target_ids[23]
    edit = window.answer_edits[item_id]
    edit.setText("saved draft")
    edit.setFocus()
    QTest.qWait(30)
    old_key = window.session.session_key
    window.show_preview()
    monkeypatch.setattr(UnfinishedPracticeDialog, "exec", lambda _: QDialog.Accepted)
    window.resume_draft()
    QTest.qWait(150)
    assert window.session.session_key == old_key
    assert QApplication.focusWidget() is window.answer_edits[item_id]
    assert window.scroll.verticalScrollBar().value() > 0
    window.start_full_practice()
    new_key = window.session.session_key
    assert new_key != old_key
    QTest.qWait(150)
    window.close()
    reopened = ui(Repository(tmp_path / "learning.db"))
    reopened.show()
    QTest.qWait(240)
    assert reopened.session is None and reopened.mode_title.text().startswith('原文模式')
    reopened.resume_draft()
    assert reopened.session.session_key == new_key and reopened.session.filled == 0
    old = next(d for d in reopened.repository.load_drafts() if d.session_key == old_key)
    assert old.answers[item_id] == "saved draft"
