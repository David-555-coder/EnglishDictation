"""All backup, corruption and rollback cases use isolated temporary libraries."""
import hashlib
import json
import os
import sqlite3
import zipfile
from contextlib import closing
from pathlib import Path

import pytest
from docx import Document

from src.database.repository import Repository
from src.services import backup_restore
from src.services.backup_restore import prepare_library_restore, restore_library, recover_interrupted_restore
from src.services.grading import grade_session
from src.services.library import ArticleLibrary, backup_library
from src.services.practice_service import new_full_session


def seed_library(root, name="文章", submitted=False):
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"{name}.docx"
    document = Document()
    paragraph = document.add_paragraph("A ")
    paragraph.add_run("rainbow").bold = True
    paragraph.add_run("（彩虹） appeared.")
    document.save(path)
    repository = Repository(root / "data" / "english_dictation.db")
    article = ArticleLibrary(repository).import_file(path)
    session = new_full_session(article)
    session.answers[session.target_ids[0]] = "wrong"
    session.view_state = {"focus": session.target_ids[0], "scroll": 275}
    if submitted:
        grade_session(session)
        repository.save_submitted_session(session)
    else:
        repository.save_draft(session)
    repository.set_setting("zoom", 130)
    repository.set_setting("quick_write", {"text": "saved editor"})
    return repository, article, session


def fingerprint(repository):
    data = repository.database_path.parent
    return {str(file.relative_to(data)): hashlib.sha256(file.read_bytes()).hexdigest()
            for file in data.rglob("*") if file.is_file()}


def rewrite_zip(source, destination, change):
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    change(entries)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return destination


def test_roundtrip_replaces_library_restores_progress_settings_and_keeps_safety_copy(tmp_path):
    source, article, session = seed_library(tmp_path / "原电脑")
    submitted = new_full_session(article)
    grade_session(submitted)
    source.save_submitted_session(submitted)
    backup = backup_library(source, tmp_path / "旧学习备份.zip")
    article_path = Path(article.source_path)
    original = article_path.read_bytes()
    current, other, _ = seed_library(tmp_path / "新目录", "另一篇", submitted=True)
    before = fingerprint(current)
    prepared = prepare_library_restore(current, backup)
    assert fingerprint(current) == before
    safety = restore_library(current, prepared)
    reopened = Repository(current.database_path)
    assert reopened.list_article_paths() == [article.source_path]
    assert other.source_path not in reopened.list_article_paths()
    restored = reopened.load_drafts()[0]
    assert restored.answers == session.answers and restored.view_state == session.view_state
    assert len(reopened.list_history()) == len(reopened.list_wrong_words()) == 1
    assert reopened.get_setting("quick_write") == {"text": "saved editor"}
    assert reopened.get_setting("zoom") == 130
    bound = ArticleLibrary(reopened).bind_copy(reopened.load_articles()[0][0])
    assert Path(bound.metadata["document_path"]).is_file()
    assert Path(bound.metadata["document_path"]).parent == current.database_path.parent / "articles"
    assert article_path.read_bytes() == original
    saved = prepare_library_restore(current, safety)
    try:
        with closing(sqlite3.connect(saved.directory / "data" / "english_dictation.db")) as connection:
            assert connection.execute("SELECT source_path FROM article_library").fetchone()[0] == other.source_path
    finally:
        saved.close()
    assert not list((tmp_path / "新目录").glob(".ed_restore_*"))


def test_previous_backup_without_manifest_is_compatible(tmp_path):
    repo, _, _ = seed_library(tmp_path / "source")
    current = Repository(tmp_path / "current" / "data" / "english_dictation.db")
    backup = backup_library(repo, tmp_path / "new.zip")
    old = rewrite_zip(backup, tmp_path / "old.zip", lambda entries: entries.pop(backup_restore.MANIFEST))
    prepared = prepare_library_restore(current, old)
    restore_library(current, prepared)
    assert len(Repository(current.database_path).load_drafts()) == 1


def test_legacy_orphan_examples_are_preserved_with_a_compatibility_warning(tmp_path):
    source, _, _ = seed_library(tmp_path / "source", submitted=True)
    with source.connection() as connection:
        example = tuple(connection.execute("SELECT * FROM word_examples").fetchone())
    with closing(sqlite3.connect(source.database_path)) as connection:
        connection.execute("DELETE FROM wrong_words")  # Reproduce an early release's FK-disabled removal.
        connection.commit()
    backup = backup_library(source, tmp_path / "legacy.zip")
    current = Repository(tmp_path / "current" / "data" / "english_dictation.db")
    prepared = prepare_library_restore(current, backup)
    assert prepared.warnings and "1" in prepared.warnings[0]
    restore_library(current, prepared)
    with current.connection() as connection:
        assert tuple(connection.execute("SELECT * FROM word_examples").fetchone()) == example
    assert current.list_wrong_words() == []


def test_owner_old_backup_is_compatible_without_changing_learning_records(tmp_path):
    backup = Path("backups/before-unfinished-restore-20260930/learning-data.zip")
    if not backup.exists():
        pytest.skip("owner backup unavailable")
    current = Repository(tmp_path / "data" / "english_dictation.db")
    with zipfile.ZipFile(backup) as archive, closing(sqlite3.connect(":memory:")) as connection:
        connection.deserialize(archive.read("data/english_dictation.db"))
        tables = ("wrong_words", "word_examples", "practice_sessions", "practice_answers", "daily_review_days", "daily_review_items")
        original = {table: connection.execute(f"SELECT * FROM {table} ORDER BY rowid").fetchall() for table in tables}
    prepared = prepare_library_restore(current, backup)
    restore_library(current, prepared)
    reopened = Repository(current.database_path)
    with reopened.connection() as connection:
        for table in tables:
            width = len(original[table][0]) if original[table] else 0
            assert [tuple(row)[:width] for row in connection.execute(f"SELECT * FROM {table} ORDER BY rowid")] == original[table]


@pytest.mark.parametrize("mutation", ["bad_db", "missing_db", "missing_copy", "bad_copy", "future", "bad_draft", "schema", "foreign_keys"])
def test_invalid_backup_never_changes_current_data(tmp_path, mutation):
    source, _, _ = seed_library(tmp_path / "source", submitted=mutation == "foreign_keys")
    backup = backup_library(source, tmp_path / "good.zip")
    current, _, _ = seed_library(tmp_path / "current", "unchanged")
    before = fingerprint(current)

    def damage(entries):
        if mutation == "bad_db":
            entries["data/english_dictation.db"] = b"not sqlite"
        elif mutation == "missing_db":
            entries.pop("data/english_dictation.db")
        elif mutation in {"missing_copy", "bad_copy"}:
            name = next(name for name in entries if name.endswith(".docx"))
            if mutation == "missing_copy":
                entries.pop(name)
            else:
                entries[name] = b"not word"
        elif mutation == "future":
            entries[backup_restore.MANIFEST] = b'{"format":"EnglishDictation","version":999}'
        else:
            database = tmp_path / "damaged.db"
            database.write_bytes(entries["data/english_dictation.db"])
            with sqlite3.connect(database) as connection:
                if mutation == "bad_draft":
                    connection.execute("UPDATE practice_drafts SET payload='{broken'")
                elif mutation == "schema":
                    connection.execute("DROP TABLE app_settings")
                else:
                    connection.execute("UPDATE practice_answers SET session_id=9999")
            entries["data/english_dictation.db"] = database.read_bytes()

    damaged = rewrite_zip(backup, tmp_path / "damaged.zip", damage)
    with pytest.raises(ValueError):
        prepare_library_restore(current, damaged)
    assert fingerprint(current) == before
    assert not list((tmp_path / "current").glob(".ed_restore_*"))


@pytest.mark.parametrize("name", ["../evil", "data/../evil", "C:/evil", "/evil", "data\\evil", "data/articles/CON.docx", "data/articles/a?.docx"])
def test_unsafe_archive_paths_are_rejected(tmp_path, name):
    repo, _, _ = seed_library(tmp_path / "current")
    before = fingerprint(repo)
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as output:
        output.writestr(name, b"invalid")
    with pytest.raises(ValueError):
        prepare_library_restore(repo, archive)
    assert fingerprint(repo) == before
    assert not (tmp_path / "evil").exists()


@pytest.mark.parametrize("kind", ["duplicate", "symlink", "oversized"])
def test_duplicate_link_and_size_limits(tmp_path, kind, monkeypatch):
    repo, _, _ = seed_library(tmp_path / "current")
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as output:
        if kind == "duplicate":
            output.writestr("data/english_dictation.db", b"one")
            output.writestr("DATA/ENGLISH_DICTATION.DB", b"two")
        elif kind == "symlink":
            info = zipfile.ZipInfo("data/english_dictation.db")
            info.create_system = 3
            info.external_attr = 0o120777 << 16
            output.writestr(info, b"outside")
        else:
            monkeypatch.setattr(backup_restore, "MAX_FILE_BYTES", 2)
            output.writestr("data/english_dictation.db", b"1234")
    with pytest.raises(ValueError):
        prepare_library_restore(repo, archive)


@pytest.mark.parametrize("failure", ["backup", "backup_validation", "swap"])
def test_failure_preserves_current_database_and_copies(tmp_path, monkeypatch, failure):
    source, _, _ = seed_library(tmp_path / "source")
    backup = backup_library(source, tmp_path / "restore.zip")
    current, _, _ = seed_library(tmp_path / "current", "keep", submitted=True)
    before = fingerprint(current)
    prepared = prepare_library_restore(current, backup)
    if failure == "backup":
        monkeypatch.setattr(backup_restore, "backup_library", lambda *_: (_ for _ in ()).throw(OSError("disk full")))
    elif failure == "backup_validation":
        monkeypatch.setattr(backup_restore, "prepare_library_restore", lambda *_: (_ for _ in ()).throw(ValueError("invalid safety copy")))
    else:
        replace = os.replace
        failed = []

        def fail_once(source, destination):
            if Path(source) == prepared.directory / "data" / "english_dictation.db" and not failed:
                failed.append(True)
                raise PermissionError("locked database")
            return replace(source, destination)

        monkeypatch.setattr(backup_restore.os, "replace", fail_once)
    try:
        with pytest.raises((OSError, ValueError)):
            restore_library(current, prepared)
        assert fingerprint(current) == before
        assert len(Repository(current.database_path).list_history()) == 1
    finally:
        prepared.close()


def test_startup_rolls_back_a_process_interruption(tmp_path):
    source, _, _ = seed_library(tmp_path / "source")
    current, _, _ = seed_library(tmp_path / "current", "keep")
    before = fingerprint(current)
    backup = backup_library(source, tmp_path / "restore.zip")
    prepared = prepare_library_restore(current, backup)
    previous = prepared.directory / "previous"
    previous.mkdir()
    marker = current.database_path.parent / backup_restore.RECOVERY_MARKER
    marker.write_text(json.dumps({"stage": prepared.directory.name, "had_articles": True}), encoding="utf-8")
    os.replace(current.database_path, previous / current.database_path.name)
    os.replace(current.database_path.parent / "articles", previous / "articles")
    os.replace(prepared.directory / "data" / "articles", current.database_path.parent / "articles")
    prepared.close()  # Must not discard a pending recovery's previous files.
    assert prepared.directory.exists()
    assert recover_interrupted_restore(current.database_path)
    assert fingerprint(current) == before
    assert not recover_interrupted_restore(current.database_path)


def test_restore_empty_library_removes_only_cached_articles(tmp_path):
    source = Repository(tmp_path / "empty" / "data" / "english_dictation.db")
    current, article, _ = seed_library(tmp_path / "current")
    backup = backup_library(source, tmp_path / "empty.zip")
    lock = current.database_path.parent / "application.lock"
    lock.write_text("active lock", encoding="utf-8")
    restore_library(current, prepare_library_restore(current, backup))
    assert Repository(current.database_path).list_article_paths() == []
    assert list((current.database_path.parent / "articles").iterdir()) == []
    assert lock.read_text(encoding="utf-8") == "active lock"
    assert Path(article.source_path).exists()
