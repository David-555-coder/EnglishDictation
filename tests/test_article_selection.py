import hashlib
import sqlite3
from pathlib import Path

import pytest
from docx import Document
from PySide6.QtCore import QEvent, QPoint, QPointF, Qt, QCoreApplication, QItemSelectionModel, QTimer
from PySide6.QtGui import QMouseEvent, QTextDocument
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QLabel, QMenu, QMessageBox, QFrame

from src.database.repository import Repository
from src.services.grading import grade_session
from src.services.practice_service import new_full_session
from test_library_ui import ui


def library(tmp_path, ui, count=4):
    w = ui(Repository(tmp_path / "data" / "db"))
    paths = []
    for i in range(count):
        doc = Document(); doc.add_paragraph(f"Title {i}", "Title")
        for _ in range(8):
            p = doc.add_paragraph("We have "); p.add_run("ambition").bold = True
            p.add_run("（抱负）. " + "Enough context for reading. " * 15)
        path = tmp_path / f"article{i}.docx"; doc.save(path); paths.append(str(path))
    w._import_paths(paths)
    return w, paths


def select(w, *rows):
    w._article_switch_delay = 130
    w.article_list.clearSelection()
    for row in rows: w.article_list.item(row).setSelected(True)


def test_rapid_switch_only_renders_final_selection(tmp_path, ui, monkeypatch):
    w, paths = library(tmp_path, ui)
    w._article_switch_delay = 130
    previews = []
    original = w.show_preview
    def preview(): previews.append(w.current_article.source_path); original()
    monkeypatch.setattr(w, "show_preview", preview)
    for i in range(60):
        w.article_list.setCurrentRow((i % 3) + 1)
        QApplication.processEvents()
    assert not previews
    QTest.qWait(250)
    assert previews == [paths[3]]
    assert w.current_article.source_path == paths[3]
    assert len(w.content.findChildren(QLabel)) == 9


def test_old_paragraphs_detach_before_deferred_deletion(tmp_path, ui):
    w, _ = library(tmp_path, ui)
    old = w.content_layout.itemAt(0).widget()
    w.article_list.setCurrentRow(1)
    assert old.parent() is None and old.isHidden()
    assert len(w.content.findChildren(QLabel)) == 9
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


def test_preview_preserves_chinese_and_target_format(tmp_path, ui):
    w, _ = library(tmp_path, ui, 1)
    label = w.content.findChildren(QFrame, 'readingParagraph')[1].findChild(QLabel)
    document = QTextDocument(); document.setHtml(label.text())
    assert document.toPlainText() == w.current_article.paragraphs[1].text
    assert '<span' in label.text() and 'font-weight:700' in label.text() and 'ambition</span>' in label.text()
    assert label.textInteractionFlags() & Qt.TextSelectableByMouse


def test_batch_remove_keeps_unselected_draft_documents_and_history(tmp_path, ui):
    w, paths = library(tmp_path, ui)
    original_hashes = {p: hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in paths}
    record = new_full_session(w.current_article); grade_session(record)
    w.repository.save_submitted_session(record)
    w.start_full_practice(); edit = next(iter(w.answer_edits.values())); edit.setText("keep input")
    current_session = w.session
    select(w, 1, 2)
    w.remove_selected_article()
    assert w.session is current_session and edit.text() == "keep input"
    assert [a.source_path for a in w.articles] == [paths[0], paths[3]]
    assert len(w.repository.list_history()) == 1
    assert w.repository.list_wrong_words()
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest() == h for p, h in original_hashes.items())
    assert len(list(w.library.directory.glob('*.docx'))) == 4


def test_cancel_and_database_failure_preserve_batch(tmp_path, ui, monkeypatch):
    w, paths = library(tmp_path, ui)
    select(w, 1, 2)
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.No)
    w.remove_selected_article()
    assert len(w.articles) == 4 and len(w.article_list.selectedItems()) == 2
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.Yes)
    def fail(paths): raise sqlite3.OperationalError('disk error')
    monkeypatch.setattr(w.repository, 'delete_article_paths', fail)
    w.remove_selected_article()
    assert len(w.articles) == 4 and len(w.article_list.selectedItems()) == 2


def test_database_batch_delete_rolls_back_every_row_on_failure(tmp_path):
    repo = Repository(tmp_path / 'db')
    paths = [tmp_path / 'a.docx', tmp_path / 'b.docx']
    for p in paths: repo.save_article_path(p)
    with repo.connection() as db:
        db.execute("CREATE TRIGGER fail_second BEFORE DELETE ON article_library WHEN OLD.filename='b.docx' BEGIN SELECT RAISE(ABORT, 'failed'); END")
    with pytest.raises(sqlite3.IntegrityError): repo.delete_article_paths(paths)
    assert len(repo.list_article_paths()) == 2


def test_remove_all_cancels_pending_switch_and_does_not_resurrect_drafts(tmp_path, ui):
    w, _ = library(tmp_path, ui)
    w.start_full_practice(); next(iter(w.answer_edits.values())).setText('unfinished')
    w._save_active_draft()
    select(w, 0, 1, 2, 3)
    w.remove_selected_article(); QTest.qWait(200)
    assert not w.articles and w.session is None and not w._pending_article_path
    assert not w.repository.load_drafts() and not w.answer_edits
    w.close()
    reopened = ui(w.repository); reopened._restore_article_library()
    assert not reopened.articles


def test_right_click_preserves_multi_selection_and_menu_count(tmp_path, ui, monkeypatch):
    w, _ = library(tmp_path, ui)
    w.show(); QApplication.processEvents()
    select(w, 1, 2)
    point = w.article_list.visualItemRect(w.article_list.item(1)).center()
    seen = []
    event = QMouseEvent(QEvent.MouseButtonPress, QPointF(point), QPointF(w.article_list.viewport().mapToGlobal(point)),
                        Qt.RightButton, Qt.RightButton, Qt.NoModifier)
    QApplication.sendEvent(w.article_list.viewport(), event)
    assert len(w.article_list.selectedItems()) == 2
    timer = QTimer()
    def close_menu():
        menu = QApplication.activePopupWidget()
        if isinstance(menu, QMenu):
            seen.extend(a.text() for a in menu.actions())
            menu.close()
    timer.timeout.connect(close_menu)
    timer.start(50)
    w._article_context_menu(point)
    timer.stop()
    assert any('2 篇' in text for text in seen)
    assert len(w.article_list.selectedItems()) == 2


def test_drag_rubber_band_selects_multiple_articles_without_loading(tmp_path, ui):
    w, _ = library(tmp_path, ui, 3)
    w._article_switch_delay = 130
    w.show(); QApplication.processEvents()
    view = w.article_list.viewport()
    rect = w.article_list.visualItemRect(w.article_list.item(2))
    start = QPoint(view.width() - 8, rect.bottom() + 35)
    end = w.article_list.visualItemRect(w.article_list.item(0)).center()
    assert start.y() < view.height()
    QTest.mousePress(view, Qt.LeftButton, Qt.NoModifier, start)
    move = QMouseEvent(QEvent.MouseMove, QPointF(end), QPointF(view.mapToGlobal(end)),
                      Qt.NoButton, Qt.LeftButton, Qt.NoModifier)
    QApplication.sendEvent(view, move)
    QTest.mouseRelease(view, Qt.LeftButton, Qt.NoModifier, end)
    assert len(w.article_list.selectedItems()) >= 2
    QTest.qWait(200)
    assert not w._pending_article_path


def test_ctrl_shift_and_delete_keyboard_batch(tmp_path, ui):
    w, paths = library(tmp_path, ui)
    w._article_switch_delay = 130
    w.show(); QApplication.processEvents()
    view = w.article_list.viewport()
    point = lambda row: w.article_list.visualItemRect(w.article_list.item(row)).center()
    QTest.mouseClick(view, Qt.LeftButton, Qt.NoModifier, point(0))
    QTest.mouseClick(view, Qt.LeftButton, Qt.ControlModifier, point(2))
    assert {i.data(Qt.UserRole) for i in w.article_list.selectedItems()} == {paths[0], paths[2]}
    QTest.mouseClick(view, Qt.LeftButton, Qt.NoModifier, point(0))
    QTest.mouseClick(view, Qt.LeftButton, Qt.ShiftModifier, point(2))
    assert len(w.article_list.selectedItems()) == 3
    QTest.keyClick(w.article_list, Qt.Key_Delete)
    assert [a.source_path for a in w.articles] == [paths[3]]


def test_search_hidden_selection_is_not_removed(tmp_path, ui):
    w, paths = library(tmp_path, ui)
    select(w, 1, 2)
    w.article_search.setText('article1')
    w.remove_selected_article()
    assert paths[1] not in w.repository.list_article_paths()
    assert paths[2] in w.repository.list_article_paths()
