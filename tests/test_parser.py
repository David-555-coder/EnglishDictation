from docx import Document
from docx.shared import RGBColor

from src.models.rules import RecognitionRules
from src.parser.docx_parser import DocxParser


def save_doc(tmp_path, build):
    document = Document(); paragraph = document.add_paragraph(); build(paragraph)
    path = tmp_path / "测试 文章.docx"; document.save(path); return DocxParser().parse(path)


def test_bold_only_english_is_target_and_chinese_is_preserved(tmp_path):
    def build(p):
        p.add_run("I have "); run = p.add_run("ambition"); run.bold = True
        p.add_run("（抱负，雄心）.")
    article = save_doc(tmp_path, build)
    assert [x.answer for x in article.word_items] == ["ambition"]
    assert article.word_items[0].chinese == "抱负，雄心"
    assert article.paragraphs[0].text == "I have ambition（抱负，雄心）."


def test_red_only_english_is_target(tmp_path):
    def build(p):
        run = p.add_run("take advantage of"); run.font.color.rgb = RGBColor(255, 0, 0)
        p.add_run("（利用）")
    article = save_doc(tmp_path, build)
    assert [x.answer for x in article.word_items] == ["take advantage of"]


def test_adjacent_formatted_runs_merge_to_one_word(tmp_path):
    def build(p):
        first = p.add_run("amb"); first.bold = True
        second = p.add_run("ition"); second.font.color.rgb = RGBColor(255, 0, 0)
        p.add_run("（抱负）")
    article = save_doc(tmp_path, build)
    assert [x.answer for x in article.word_items] == ["ambition"]


def test_plain_english_is_not_target(tmp_path):
    article = save_doc(tmp_path, lambda p: p.add_run("ordinary text（普通文字）"))
    assert article.word_items == []


def test_formatted_heading_without_chinese_note_is_not_target(tmp_path):
    def build(p):
        heading = p.add_run("The Big Family Get-Together")
        heading.bold = True
    article = save_doc(tmp_path, build)
    assert article.word_items == []


def test_formatted_word_requires_its_own_adjacent_chinese_note(tmp_path):
    def build(p):
        heading = p.add_run("Exercise title ")
        heading.bold = True
        p.add_run("— ")
        word = p.add_run("ambition")
        word.bold = True
        p.add_run("（抱负）")
    article = save_doc(tmp_path, build)
    assert [item.answer for item in article.word_items] == ["ambition"]


def test_lenient_mode_uses_plain_run_immediately_before_chinese_note(tmp_path):
    document = Document(); paragraph = document.add_paragraph("I have ")
    paragraph.add_run("take advantage of")
    paragraph.add_run("（利用） every chance.")
    path = tmp_path / "plain phrase.docx"; document.save(path)
    article = DocxParser().parse(path, RecognitionRules("any", True))
    assert [item.answer for item in article.word_items] == ["take advantage of"]
    assert article.paragraphs[0].text == "I have take advantage of（利用） every chance."


def test_lenient_mode_keeps_context_when_everything_is_one_run(tmp_path):
    document = Document(); document.add_paragraph("I have ambition（抱负，雄心）.")
    path = tmp_path / "one run.docx"; document.save(path)
    article = DocxParser().parse(path, RecognitionRules("any", True))
    assert [item.answer for item in article.word_items] == ["ambition"]


def test_strict_mode_requires_red_and_bold_but_not_chinese(tmp_path):
    document = Document(); paragraph = document.add_paragraph()
    red_only = paragraph.add_run("redonly"); red_only.font.color.rgb = RGBColor(255, 0, 0)
    paragraph.add_run(" ")
    strict = paragraph.add_run("strictword"); strict.bold = True
    strict.font.color.rgb = RGBColor(255, 0, 0)
    path = tmp_path / "strict.docx"; document.save(path)
    article = DocxParser().parse(path, RecognitionRules("red_and_bold", False))
    assert [item.answer for item in article.word_items] == ["strictword"]


def test_lenient_mode_still_skips_word_heading(tmp_path):
    document = Document(); paragraph = document.add_heading(level=1)
    paragraph.add_run("Vocabulary（词汇）")
    path = tmp_path / "heading.docx"; document.save(path)
    article = DocxParser().parse(path, RecognitionRules("any", True, True))
    assert article.word_items == []
