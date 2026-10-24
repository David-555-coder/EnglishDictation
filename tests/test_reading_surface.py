"""Continuous reading, keyboard affordances, and menu-only library visibility."""
from copy import deepcopy

from docx import Document
from PySide6.QtCore import QPoint, Qt, QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFrame, QLabel, QPushButton

from test_library_ui import ui, sample_word
from test_focus_and_resume import practice
from src.database.repository import Repository


def story(tmp_path, count=24):
    path = tmp_path / 'story.docx'; doc = Document()
    doc.add_paragraph().add_run('The Quiet Town').bold = True
    for i in range(count):
        p = doc.add_paragraph(f'Paragraph {i}: Mia used ')
        p.add_run('ambition').bold = True
        p.add_run('（抱负） to describe her hopes & plans.\nShe kept reading with care.')
    doc.save(path)
    return path


def test_original_is_one_page_with_complete_unchanged_paragraphs(tmp_path, ui):
    repo = Repository(tmp_path / 'data/db'); window = ui(repo)
    window._import_paths([str(story(tmp_path))])
    original = deepcopy(window.current_article)
    page, = window.content.findChildren(QFrame, 'originalPage')
    paragraphs = page.findChildren(QFrame, 'readingParagraph')
    assert [p.property('paragraphText') for p in paragraphs] == [p.text for p in original.paragraphs if p.text.strip()]
    assert not window.content.findChildren(QFrame, 'paragraphCard')
    assert 'ambition' in paragraphs[1].findChild(QLabel).text()
    assert '（抱负）' in paragraphs[1].findChild(QLabel).text()
    assert '&amp;' in paragraphs[1].findChild(QLabel).text()
    assert '<br>' in paragraphs[1].findChild(QLabel).text()
    assert paragraphs[0].findChild(QLabel).property('articleHeading')
    assert window.current_article == original and repo.load_drafts() == [] and repo.list_history() == []
    window.start_full_practice()
    assert window.session.article == original and len(window.answer_edits) == original.target_count


def test_original_layout_reflows_at_reader_sizes_without_rebuilding_labels(tmp_path, ui):
    window = ui(Repository(tmp_path / 'data/db')); window._import_paths([str(story(tmp_path))]); window.show()
    labels = window.content.findChild(QFrame, 'originalPage').findChildren(QLabel)
    for width in (980, 1480, 2400):
        window.resize(width, 900)
        for zoom in (70, 100, 160, 180):
            window.set_zoom(zoom); QApplication.processEvents()
            assert window.scroll.horizontalScrollBar().maximum() == 0
            assert window.content.findChild(QFrame, 'originalPage').findChildren(QLabel) == labels
            assert labels[1].font().pixelSize() == round(20 * zoom / 100)
            assert labels[0].font().pixelSize() == round(24 * zoom / 100)
            for label in labels:
                assert label.height() >= label.heightForWidth(label.width())


def test_original_scroll_follows_visible_paragraph_during_zoom_and_focus(tmp_path, ui):
    window = ui(Repository(tmp_path / 'data/db')); window._import_paths([str(story(tmp_path, 40))])
    window.resize(1480, 900); window.show(); QApplication.processEvents()
    QTest.qWait(30); window._cancel_pending_follow()
    paragraph = window.content.findChild(QFrame, 'originalPage').findChildren(QFrame, 'readingParagraph')[20]
    bar = window.scroll.verticalScrollBar(); bar.setValue(paragraph.mapTo(window.content, QPoint()).y())
    QApplication.processEvents(); window._cancel_pending_follow()
    offset = paragraph.mapTo(window.content, QPoint()).y() - bar.value()
    for zoom in (130, 160, 100):
        window.set_zoom(zoom); QApplication.processEvents(); QTest.qWait(15)
        assert abs(paragraph.mapTo(window.content, QPoint()).y() - bar.value() - offset) <= 3
    for _ in range(2):
        window.actions['专注模式'].trigger(); QApplication.processEvents(); QTest.qWait(15)
        assert abs(paragraph.mapTo(window.content, QPoint()).y() - bar.value() - offset) <= 3


def test_import_default_focus_has_no_outline_but_tab_navigation_remains(tmp_path, ui):
    window = ui(Repository(tmp_path / 'db')); window.show(); window.activateWindow(); QApplication.processEvents()
    assert window.import_button.focusPolicy() & Qt.TabFocus
    assert not window.import_button.property('keyboardFocus')
    window.import_button.setFocus()
    QTest.keyClick(window.import_button, Qt.Key_Tab); QApplication.processEvents()
    assert QApplication.focusWidget() is window.export_button
    assert window.export_button.property('keyboardFocus') is True
    QTest.keyClick(window.export_button, Qt.Key_Tab, Qt.ShiftModifier); QApplication.processEvents()
    assert QApplication.focusWidget() is window.import_button
    assert window.import_button.property('keyboardFocus') is True


def test_mouse_clears_toolbar_outline_even_when_focus_stays_in_toolbar(tmp_path, ui):
    window = ui(Repository(tmp_path / 'db')); window._import_paths([str(sample_word(tmp_path))])
    window.show(); window.activateWindow(); QApplication.processEvents()
    window.import_button.setFocus()
    QTest.keyClick(window.import_button, Qt.Key_Tab)
    QTest.keyClick(window.export_button, Qt.Key_Tab, Qt.ShiftModifier); QApplication.processEvents()
    assert window.import_button.property('keyboardFocus') is True
    QTest.mouseClick(window.content, Qt.LeftButton, pos=QPoint(2, 2)); QApplication.processEvents()
    assert not window.import_button.property('keyboardFocus')
    window.import_button.setFocus(); QTest.keyClick(window.import_button, Qt.Key_Tab); QApplication.processEvents()
    assert window.export_button.property('keyboardFocus') is True


def test_keyboard_menu_keeps_focus_feedback_until_mouse_is_used(tmp_path, ui):
    window = ui(Repository(tmp_path / 'db')); window.show(); window.activateWindow(); QApplication.processEvents()
    window.import_button.setFocus(Qt.TabFocusReason)
    menu = window.import_button.menu(); opened = []
    timer = QTimer(); timer.setSingleShot(True)
    def close_menu():
        opened.append(menu.isVisible())
        QTest.keyClick(menu, Qt.Key_Escape)
    timer.timeout.connect(close_menu); timer.start(40)
    QTest.keyClick(window.import_button, Qt.Key_Space); QTest.qWait(65); QApplication.processEvents()
    assert opened == [True] and not menu.isVisible()
    assert window.import_button.property('keyboardFocus') is True
    QTest.mouseClick(window.status_panel, Qt.LeftButton, pos=QPoint(2, 2)); QApplication.processEvents()
    assert not window.import_button.property('keyboardFocus')


def test_menu_hides_entire_library_without_old_strip_and_f8_restores_it(tmp_path, ui):
    window = practice(tmp_path, ui, 24)
    item = next(iter(window.answer_edits.values())); item.setText('keep this answer')
    session, sizes = window.session, window.splitter.sizes()
    assert not any('收起文章列表' in b.text() for b in window.findChildren(QPushButton))
    window.actions['显示文章列表'].trigger(); QApplication.processEvents()
    assert window.library_shell.isHidden() and window.splitter.sizes()[0] == 0
    assert window.main_toolbar.isVisible() and window.status_panel.isVisible()
    window.actions['专注模式'].trigger(); QApplication.processEvents()
    window.actions['专注模式'].trigger(); QApplication.processEvents()
    assert window.library_shell.isHidden() and window.splitter.sizes()[0] == 0
    assert not window.actions['显示文章列表'].isChecked()
    window.actions['显示文章列表'].trigger(); QApplication.processEvents()
    assert window.library_shell.isVisible() and abs(window.splitter.sizes()[0] - sizes[0]) <= 5
    assert window.session is session and item.text() == 'keep this answer'


def test_preexisting_collapsed_setting_can_be_reopened_from_menu(tmp_path, ui):
    repo = Repository(tmp_path / 'db'); repo.set_setting('library_collapsed', True)
    window = ui(repo); window.show(); QApplication.processEvents()
    assert window.library_shell.isHidden() and window.splitter.sizes()[0] == 0
    assert not window.actions['显示文章列表'].isChecked()
    window.actions['显示文章列表'].trigger(); QApplication.processEvents()
    assert window.library_shell.isVisible() and window.splitter.sizes()[0] >= 200
