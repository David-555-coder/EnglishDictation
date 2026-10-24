from copy import deepcopy

import pytest
from docx import Document
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QInputDialog, QLabel, QMessageBox

from test_library_ui import ui
from test_daily_review import seed, daily, submit
from src.database.repository import Repository
from src.services.practice_service import new_wrong_session
from src.services.serialization import session_to_dict


def correction(repo, count=19):
    repo.enable_review_continuity()
    seed(repo, count, repo.today().isoformat())
    first = daily(repo)
    submit(repo, first, False)
    second = new_wrong_session(first)
    submit(repo, second, False)
    return second


def tables(repo, exclude=()):
    with repo.connection() as db:
        names = [r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")
                 if r[0] not in exclude]
        return {name: [tuple(r) for r in db.execute(f'SELECT * FROM "{name}" ORDER BY rowid')]
                for name in names}


def test_reader_preserves_result_answers_plan_models_and_drafts(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / 'db')
    second = correction(repo)
    window = ui(repo)
    window._activate_session(second)
    before = tables(repo, ('memory_exposures',))
    snapshot = deepcopy(session_to_dict(second))
    window.read_wrong_button.click()
    from PySide6.QtWidgets import QFrame
    shown = [c.property('paragraphText') for c in window.content.findChildren(QFrame) if c.property('paragraphText')]
    assert any('Learn word0 now.' in s for s in shown)
    assert window._wrong_reading and not window.answer_edits
    assert window.session is second and window.pending_review_session is None
    assert session_to_dict(second) == snapshot
    assert tables(repo, ('memory_exposures',)) == before
    with repo.connection() as db:
        assert db.execute("SELECT count(*) FROM memory_exposures WHERE event_key LIKE 'result-reading:%'").fetchone()[0] == 19
    assert window.progress.format() == '本轮已提交'
    window.reading_back_button.click()
    assert not window._wrong_reading and window.session is second


def test_reader_groups_correct_source_sentences_and_excludes_already_correct_sentence(tmp_path, ui):
    path = tmp_path / 'story.docx'
    doc = Document()
    p = doc.add_paragraph('The ')
    p.add_run('workforce').bold = True
    p.add_run('（劳动力） faces a ')
    p.add_run('forthcoming').bold = True
    p.add_run('（即将到来的） storm. ')
    p.add_run('Success').bold = True
    p.add_run('（成功） follows.')
    doc.save(path)
    window = ui(Repository(tmp_path / 'db'))
    window._import_paths([str(path)])
    window.start_full_practice()
    for item in window.session.article.word_items:
        window.answer_edits[item.id].setText(item.answer if item.answer == 'Success' else 'wrong-input')
    window._commit_submission()
    window.read_current_wrong()
    from PySide6.QtWidgets import QFrame
    text = '\n'.join(c.property('paragraphText') for c in window.content.findChildren(QFrame) if c.property('paragraphText'))
    assert text.count('storm.') == 1  # Two wrong words in one sentence, shown once.
    assert 'workforce（劳动力）' in text and 'forthcoming（即将到来的）' in text
    assert 'Success' not in text and 'wrong-input' not in text
    assert not window.answer_edits
    assert QApplication.activeModalWidget() is None


@pytest.mark.parametrize('all_correct', [False, True])
def test_group_returns_to_original_result_even_after_draft_restart(tmp_path, ui, monkeypatch, all_correct):
    monkeypatch.setattr(QMessageBox, 'information', lambda *a: QMessageBox.Ok)
    repo = Repository(tmp_path / 'db')
    parent = correction(repo)
    window = ui(repo)
    window._activate_session(parent)
    original = deepcopy(session_to_dict(parent))
    window.start_consolidation(6)
    group_key = window.session.session_key
    assert window.session.total == 6
    assert window.return_result_button.isHidden()
    assert window.read_wrong_button.isHidden()
    window._save_active_draft()
    restarted = Repository(repo.database_path)
    group = next(s for s in restarted.load_drafts() if s.session_key == group_key)
    assert group.view_state['consolidation_origin']['results'] == original['results']
    window._activate_session(group)
    if all_correct:
        for item_id, edit in window.answer_edits.items():
            edit.setText(group.article.item_map[item_id].answer)
    window._commit_submission()
    assert not window.return_result_button.isHidden()
    assert window.read_wrong_button.isHidden() is all_correct
    before = tables(repo)
    window.return_result_button.click()
    assert window.session.database_id == parent.database_id
    assert window.session.results == parent.results
    assert window.session.wrong_ids == parent.wrong_ids
    assert window.session.view_state['first_round_result'] == parent.view_state['first_round_result']
    assert window.primary_action_button.text() == '再练一组'
    assert not window.read_wrong_button.isHidden()
    assert len(repo.consolidation_remaining(window.session)) == 13
    assert tables(repo) == before  # Returning never submits/regrades the original.
    window.start_consolidation(10)
    window._commit_submission()
    window.return_to_consolidation_origin()
    window.start_consolidation(10)
    assert window.session.total == 3
    window._commit_submission()
    window.return_to_consolidation_origin()
    assert not window.primary_action_button.isEnabled()
    assert not window.read_wrong_button.isHidden()  # Reading remains after the daily extra limit.
    assert repo.consolidation_remaining(window.session) == []
    assert repo.review_status()['completed'] == 19
    assert len(repo.list_history()) == 5


def test_group_size_choice_cancel_and_fewer_words(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / 'db')
    parent = correction(repo)
    window = ui(repo)
    window._activate_session(parent)
    before = tables(repo)
    monkeypatch.setattr(QInputDialog, 'getInt', lambda *a: (10, False))
    window.primary_action_button.click()
    assert window.session is parent and tables(repo) == before
    choices = []

    def choose(*args):
        choices.append(args[3:])
        return 4, True

    monkeypatch.setattr(QInputDialog, 'getInt', choose)
    window.primary_action_button.click()
    assert choices == [(10, 1, 10)]
    assert window.session.total == 4
    assert len(repo.consolidation_remaining(parent)) == 15


def test_compact_submitted_sidebar_and_f8_keep_result_accessible(tmp_path, ui):
    repo = Repository(tmp_path / 'db')
    parent = correction(repo)
    window = ui(repo)
    window._activate_session(parent)
    window.resize(980, 650)
    window.show()
    QApplication.processEvents()
    assert window.status_panel.width() <= 210
    for button in (window.primary_action_button, window.read_wrong_button):
        assert button.isVisible()
        assert button.fontMetrics().horizontalAdvance(button.text()) + 30 <= button.width()
    for label in (window.progress_text, window.score_text):
        assert label.height() >= label.heightForWidth(label.width())
    hint = window.status_panel.findChild(QLabel, 'hintCard')
    assert 'F8' in hint.text()
    window.activateWindow()
    QApplication.processEvents()
    QTest.keyClick(window, Qt.Key_F8)
    QApplication.processEvents()
    assert window._focus_mode and window.status_panel.isVisible()
    assert window.read_wrong_button.isVisible()
    QTest.keyClick(window, Qt.Key_F8)
    QApplication.processEvents()
    assert not window._focus_mode
