"""Integrated reading-only translation: real menus, no production learning DB."""
from copy import deepcopy
from pathlib import Path
import os
import time

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import pytest
from docx import Document
from PySide6.QtCore import QPoint, Qt
from PySide6.QtGui import QFontDatabase
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from src.database.repository import Repository
from src.services.dictionary_lookup import Dictionary, DictionaryClient, Entry, Lookup
from src.services.reading_translation import Query, TranslationClient
from src.services.serialization import session_to_dict
from src.ui.main_window import MainWindow
from src.ui.reading_translation import ReadingTranslation
from src.ui.widgets.reading_selection import units
from src.ui.widgets.text_menus import TextLabel, plain_label_text, text_menu
from tools.translation_prototype.test_prototype import service, make_client


APP = QApplication.instance() or QApplication([])
for name in ('msyh.ttc', 'segoeui.ttf'):
    font = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / name
    if font.exists(): QFontDatabase.addApplicationFont(str(font))


def wait_until(test, timeout=3000):
    limit = time.monotonic() + timeout / 1000
    while not test() and time.monotonic() < limit:
        QTest.qWait(10)
    assert test()


@pytest.fixture
def reader(tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, 'question', lambda *a: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, 'information', lambda *a: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, 'warning', lambda *a: QMessageBox.Ok)
    monkeypatch.setattr(MainWindow, '_restore_article_library', lambda self: None)
    path = tmp_path / 'reading.docx'
    doc = Document()
    p = doc.add_paragraph('I was '); p.add_run('utterly').bold = True
    p.add_run('（完全地） exhausted after the long journey.')
    p = doc.add_paragraph('She wanted to '); p.add_run('outdo').bold = True
    p.add_run('（超过） her prior self. Never spill the beans about the surprise.')
    doc.save(path)
    window = MainWindow(Repository(tmp_path / 'learning.db'))
    window._article_switch_delay = 0
    window._import_paths([str(path)])
    window.resize(1100, 850); window.show(); APP.processEvents()
    yield window
    window.close(); window.deleteLater(); APP.processEvents()


def select(window, selected):
    APP.processEvents()
    window.reading_selection.clear()
    label = next(w for w in window.content.findChildren(TextLabel)
                 if w.isVisibleTo(window.content) and selected in plain_label_text(w))
    text = plain_label_text(label); start = text.index(selected)
    label.setSelection(units(text[:start]), units(selected))
    return label


def translate(window, selected):
    menu = text_menu(select(window, selected))
    try:
        action = next(a for a in menu.actions() if a.text() == '翻译')
        action.trigger()
    finally:
        menu.deleteLater()
    wait_until(lambda: window._reading_translation.popup.lookup_done)
    return window._reading_translation


def test_default_reading_menu_creates_only_lazy_offline_clients(reader):
    assert reader._reading_translation is None
    controller = translate(reader, 'utterly')
    assert '完全' in controller.popup.result.toPlainText()
    assert controller.network is None and controller.dictionary.lookups == 1
    assert controller.popup.width() <= 420 and controller.popup.height() < 280
    assert controller.popup.isModal() is False


def test_whole_phrase_and_missing_expression_do_not_send_network(reader):
    controller = translate(reader, 'spill the beans')
    assert '说漏嘴' in controller.popup.result.toPlainText()
    controller = translate(reader, 'about the surprise')
    assert '未收录' in controller.popup.result.toPlainText()
    assert controller.network is None


def test_independent_practice_has_no_translation_and_no_learning_side_effect(reader):
    reader.start_full_practice(); APP.processEvents()
    before = deepcopy(session_to_dict(reader.session))
    label = select(reader, 'journey')
    menu = text_menu(label)
    assert '翻译' not in [a.text() for a in menu.actions()]
    menu.deleteLater()
    reader._translate_reading_selection(label, QPoint(100, 100))
    assert reader._reading_translation is None
    assert session_to_dict(reader.session) == before
    assert not reader.session.viewed_answer


def test_stale_menu_cannot_open_translation_after_starting_practice(reader):
    menu = text_menu(select(reader, 'utterly'))
    action = next(a for a in menu.actions() if a.text() == '翻译')
    reader.start_full_practice(); action.trigger()
    assert reader._reading_translation is None
    menu.deleteLater()


def test_show_answers_during_active_round_still_disallows_translation(reader):
    reader.start_full_practice(); reader.show_answers(); APP.processEvents()
    assert not reader._can_translate_reading()
    assert '翻译' not in [a.text() for a in text_menu(select(reader, 'journey')).actions()]


def test_submitted_results_and_wrong_sentence_reading_allow_lookup_without_mutating_session(reader):
    reader.start_full_practice()
    reader._confirm_finished_submission = lambda: True
    reader.submit(); APP.processEvents()
    before = deepcopy(session_to_dict(reader.session))
    controller = translate(reader, 'journey')
    assert controller.popup.lookup_done and session_to_dict(reader.session) == before
    reader.read_current_wrong(); APP.processEvents()
    assert not controller.popup.isVisible()
    controller = translate(reader, 'journey')
    assert controller.network is None and reader.session.submitted


def test_switch_start_and_close_hide_popup(reader):
    controller = translate(reader, 'utterly')
    reader.show_preview()
    assert not controller.popup.isVisible()
    controller = translate(reader, 'utterly')
    reader.start_full_practice()
    assert not controller.popup.isVisible()
    reader.show_preview(); controller = translate(reader, 'utterly')
    reader.close(); assert not controller.popup.isVisible()


def test_escape_closes_non_modal_popup_from_reading_surface(reader):
    controller = translate(reader, 'utterly')
    reader.activateWindow(); label = select(reader, 'utterly'); label.setFocus()
    APP.processEvents(); QTest.keyClick(label, Qt.Key_Escape); APP.processEvents()
    assert not controller.popup.isVisible() and not controller.escape.isEnabled()


def test_online_first_click_only_shows_consent_and_actual_sentence(reader):
    controller = translate(reader, 'utterly')
    controller.online_action()
    assert controller.network is None and controller.popup.net_state == 'consent'
    assert 'long journey' in controller.popup.net_text.toPlainText()
    assert '（完全地）' not in controller.popup.net_text.toPlainText()
    assert 'outdo' not in controller.online_query.text


def test_selected_context_handles_unicode_without_hidden_storage(reader):
    selection = reader.reading_selection
    selection._text = '😀 She was utterly tired.'
    selection._anchor, selection._position = units('😀 She was '), units('😀 She was utterly')
    document, start, end = selection.selected_context()
    assert document[start:end] == 'utterly'


def test_latest_wins_queue_is_bounded_and_thread_is_lazy():
    class Slow:
        def lookup(self, text):
            QTest.qSleep(45)
            return Lookup(text, (Entry(text, '', '中文'),), 'exact')
    client = DictionaryClient(store=Slow()); results = []
    client.completed.connect(lambda *args: results.append(args))
    assert client._pool is None
    for index in range(80): client.request('word' + str(index))
    assert len(client._jobs) == 1 and client._pending[1] == 'word79'
    wait_until(lambda: bool(results))
    assert len(results) == 1 and results[0][1].selected == 'word79'
    assert client.lookups <= 2 and client._pool.maxThreadCount() == 1
    assert not client._jobs and client._pending is None
    client.deleteLater(); APP.processEvents()


def test_cancelled_running_read_and_cache_hit_stay_silent():
    client = DictionaryClient(); results = []
    client.completed.connect(lambda *args: results.append(args))
    client.request('utterly'); wait_until(lambda: len(results) == 1)
    client.request('utterly'); client.cancel(); QTest.qWait(30)
    assert len(results) == 1
    client.deleteLater(); APP.processEvents()


def test_dictionary_default_path_is_public_assets_only():
    store = Dictionary()
    assert store.path.name == 'dictionary.sqlite3' and store.path.parent.name == 'assets'
    assert 'learning' not in str(store.path)
    assert store.lookup('sustainable').entries[0].source == '补充释义'


def test_explicit_online_reference_and_return_restore_dictionary(reader, service):
    client, _ = make_client(service)
    controller = ReadingTranslation(reader, reader._can_translate_reading, network=client)
    reader._reading_translation = controller
    controller = translate(reader, 'utterly')
    controller.online_action()
    assert not service[1] and controller.popup.net_state == 'consent'
    controller.online_action()
    wait_until(lambda: controller.popup.net_state == 'done')
    assert service[1] == ['I was utterly exhausted after the long journey.']
    assert '译文' in controller.popup.net_text.toPlainText()
    assert '完全' in controller.popup.result.toPlainText()
    controller.online_action()
    assert not controller.popup.network.isVisible() and '完全' in controller.popup._copy_text
    client.deleteLater()


def test_late_network_reply_after_page_switch_cannot_reopen_popup(reader, service):
    client, _ = make_client(service)
    controller = ReadingTranslation(reader, reader._can_translate_reading, network=client)
    reader._reading_translation = controller
    controller = translate(reader, 'utterly')
    controller.online_query = Query('slow', 'slow response')
    controller.online_action(); controller.online_action()
    reader.show_preview(); QTest.qWait(260)
    assert not controller.popup.isVisible() and client._active is None
    assert controller._network_token is None
    client.deleteLater()


def test_repeated_page_lookup_reuses_one_popup_and_bounded_cache(reader):
    controller = translate(reader, 'utterly')
    for _ in range(12):
        reader.show_preview(); controller = translate(reader, 'utterly')
    assert controller.dictionary.lookups == 1
    assert len(controller.dictionary._cache) == 1
    assert len(reader.findChildren(type(controller.popup))) == 1


def test_pending_read_cannot_interrupt_native_input_after_start(reader):
    class Slow:
        def lookup(self, text):
            time.sleep(.15)
            return Lookup(text, (Entry(text, '', '迟到的释义'),), 'exact')
    local = DictionaryClient(store=Slow())
    controller = ReadingTranslation(reader, reader._can_translate_reading, dictionary=local)
    reader._reading_translation = controller
    menu = text_menu(select(reader, 'utterly'))
    next(a for a in menu.actions() if a.text() == '翻译').trigger(); menu.deleteLater()
    reader.start_full_practice(); APP.processEvents()
    edit = next(iter(reader.answer_edits.values()))
    edit.setFocus(); QTest.keyClicks(edit, 'typed answer', Qt.NoModifier, 1)
    QTest.qWait(180)
    assert edit.text() == 'typed answer'
    assert not controller.popup.isVisible() and not reader.session.viewed_answer
    local.deleteLater()


def test_missing_dictionary_is_a_local_error_not_a_learning_failure(reader, tmp_path):
    local = DictionaryClient(store=Dictionary(tmp_path / 'missing.sqlite3', supplements=False))
    reader._reading_translation = ReadingTranslation(reader, reader._can_translate_reading, dictionary=local)
    controller = translate(reader, 'utterly')
    assert '词库暂不可用' in controller.popup.result.toPlainText()
    assert not (tmp_path / 'missing.sqlite3').exists() and controller.network is None
    reader.start_full_practice(); APP.processEvents()
    assert reader.session is not None and reader.answer_edits
    assert not controller.popup.isVisible()
    local.deleteLater()
