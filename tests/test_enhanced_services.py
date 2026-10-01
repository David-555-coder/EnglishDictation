import hashlib
import os
import sqlite3
import zipfile
from copy import deepcopy

import pytest
from docx import Document
from docx.shared import RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from src.database.repository import Repository
from src.models.rules import RecognitionRules
from src.parser.docx_parser import DocxParser
from src.parser.word_structure import iter_paragraphs
from src.services.exporter import export_worksheet, export_answer
from src.services.grading import grade_session
from src.services.practice_service import new_full_session, new_saved_wrong_session, new_wrong_session, group_items_by_sentence, new_long_term_session
from src.services.library import ArticleLibrary, backup_library
from src.services.target_editor import set_target, remove_target


def document(tmp_path):
    path = tmp_path / "article.docx"
    doc = Document()
    doc.add_paragraph("The Story", "Title")
    p = doc.add_paragraph("The ")
    p.add_run("extensive").bold = True
    p.add_run("（广泛的） research gained ")
    p.add_run("momentum").bold = True
    p.add_run("（势头）. Another sentence stays visible.")
    doc.save(path)
    return path


def test_corrupt_draft_does_not_hide_healthy_drafts(tmp_path):
    repo = Repository(tmp_path / "db")
    session = new_full_session(DocxParser().parse(document(tmp_path)))
    repo.save_draft(session)
    with repo.connection() as connection:
        connection.execute("INSERT INTO practice_drafts VALUES ('bad', 'missing', '{broken', '2099')")
    assert [s.session_key for s in repo.load_drafts()] == [session.session_key]
    with repo.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM practice_drafts").fetchone()[0] == 2


def test_hyperlink_context_and_target_are_preserved_in_export(tmp_path):
    doc = Document(); p = doc.add_paragraph("Visit ")
    link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), "chapter")
    r = OxmlElement("w:r"); props = OxmlElement("w:rPr"); props.append(OxmlElement("w:b")); r.append(props)
    text = OxmlElement("w:t"); text.text = "library"; r.append(text); link.append(r); p._p.append(link)
    p.add_run("（图书馆） for research.")
    path = tmp_path / "link.docx"; doc.save(path)
    article = DocxParser().parse(path)
    assert article.paragraphs[0].text == "Visit library（图书馆） for research."
    assert article.word_items[0].answer == "library"
    output = tmp_path / "sheet.docx"; export_worksheet(article, output)
    result = Document(output)
    assert result.paragraphs[0].text.startswith("Visit ___")
    assert result.paragraphs[0].text.endswith("（图书馆） for research.")
    assert result.paragraphs[0]._p.find(qn("w:hyperlink")) is not None


def test_saved_wrong_answers_are_nonempty_and_same_sentence_grouped(tmp_path):
    article = DocxParser().parse(document(tmp_path))
    repository = Repository(tmp_path / "test.db")
    session = new_full_session(article)
    grade_session(session)
    session_id = repository.save_submitted_session(session)
    restored = new_saved_wrong_session(repository.get_practice_session(session_id), repository.get_wrong_answers(session_id))
    assert [i.answer for i in restored.article.word_items] == ["extensive", "momentum"]
    assert len(group_items_by_sentence(restored)) == 1
    assert "Another sentence" not in restored.article.word_items[0].sentence
    restored.answers[restored.target_ids[0]] = "Extensive"
    grade_session(restored)
    next_round = new_wrong_session(restored)
    assert [next_round.article.item_map[i].answer for i in next_round.target_ids] == ["momentum"]
    assert restored.correct_count == 1


def test_paragraph_zero_is_not_replaced_by_row_number(tmp_path):
    article = DocxParser().parse(document(tmp_path))
    session = new_full_session(article); grade_session(session)
    repo = Repository(tmp_path / "db")
    sid = repo.save_submitted_session(session)
    rows = repo.get_wrong_answers(sid)
    for row in rows: row["paragraph_index"] = 0
    restored = new_saved_wrong_session(repo.get_practice_session(sid), rows)
    assert len(group_items_by_sentence(restored)) == 1


def test_submission_retry_is_idempotent_and_deletes_draft(tmp_path):
    article = DocxParser().parse(document(tmp_path))
    repo = Repository(tmp_path / "db")
    session = new_full_session(article); repo.save_draft(session); grade_session(session)
    first = repo.save_submitted_session(session)
    assert repo.save_submitted_session(deepcopy(session)) == first
    assert len(repo.list_history()) == 1
    assert repo.load_drafts() == []
    assert all(w["error_count"] == 1 for w in repo.list_wrong_words())


def test_full_draft_round_trip_retains_input_snapshot_focus_and_answer_flag(tmp_path):
    repo = Repository(tmp_path / "db")
    session = new_full_session(DocxParser().parse(document(tmp_path)))
    session.answers[session.target_ids[0]] = "MY INPUT"
    session.viewed_answer = True
    session.view_state = {"scroll": 240, "focus": session.target_ids[1]}
    repo.save_draft(session)
    reopened = Repository(tmp_path / "db").load_drafts()[0]
    assert reopened.answers == session.answers
    assert reopened.view_state == session.view_state
    assert reopened.viewed_answer and not reopened.submitted
    assert reopened.results == []
    assert reopened.article == session.article


@pytest.mark.parametrize("export", [export_answer, export_worksheet])
def test_exports_refuse_original_and_hard_link(tmp_path, export):
    path = document(tmp_path)
    article = DocxParser().parse(path)
    before = path.read_bytes()
    with pytest.raises(ValueError): export(article, path)
    link = tmp_path / "alias.docx"
    os.link(path, link)
    with pytest.raises(ValueError): export(article, link)
    assert path.read_bytes() == before


def test_export_preserves_tables_annotations_split_runs_and_original(tmp_path):
    path = tmp_path / "table.docx"
    doc = Document(); doc.add_paragraph("Heading")
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    p.add_run("amb").bold = True
    r = p.add_run("ition（抱负） after."); r.bold = True
    doc.save(path)
    before = path.read_bytes()
    article = DocxParser().parse(path)
    export_worksheet(article, tmp_path / "worksheet.docx")
    out = Document(tmp_path / "worksheet.docx")
    assert len(out.tables) == 1
    assert "（抱负） after." in out.tables[0].cell(0, 0).text
    assert "ambition" not in out.tables[0].cell(0, 0).text
    export_answer(article, tmp_path / "answer.docx")
    assert "ambition（抱负） after." == Document(tmp_path / "answer.docx").tables[0].cell(0,0).text
    assert path.read_bytes() == before


def test_document_copy_and_snapshot_survive_missing_original_and_backup(tmp_path):
    path = document(tmp_path)
    repo = Repository(tmp_path / "data" / "english_dictation.db")
    library = ArticleLibrary(repo)
    article = library.import_file(path)
    path.unlink()
    articles, unresolved = Repository(repo.database_path).load_articles()
    assert not unresolved and articles[0].target_count == 2
    export_worksheet(library.bind_copy(articles[0]), tmp_path / "paper.docx")
    backup = backup_library(repo, tmp_path / "backup.zip")
    with zipfile.ZipFile(backup) as archive:
        assert "data/english_dictation.db" in archive.namelist()
        assert len([x for x in archive.namelist() if x.endswith(".docx")]) == 1


def test_manual_adjustments_merge_split_and_persist_without_modifying_word(tmp_path):
    path = document(tmp_path); before = path.read_bytes()
    repo = Repository(tmp_path / "data" / "db")
    article = ArticleLibrary(repo).import_file(path)
    p = article.paragraphs[1]
    start = p.text.index("research gained")
    set_target(article, p.index, start, start + len("research gained"))
    merged = next(i for i in article.word_items if i.answer == "research gained")
    remove_target(article, merged.id)
    set_target(article, p.index, start, start + 8)
    repo.save_article(article)
    loaded = repo.load_articles()[0][0]
    assert any(i.answer == "research" for i in loaded.word_items)
    assert "".join(s.text for s in loaded.paragraphs[1].segments) == p.text
    assert path.read_bytes() == before


def test_rules_heading_exclusion_and_inherited_style(tmp_path):
    path = tmp_path / "rules.docx"
    doc = Document()
    doc.add_paragraph("Title（题目）", "Heading 1")
    style = doc.styles.add_style("StudyWord", WD_STYLE_TYPE.CHARACTER); style.font.bold = True
    p = doc.add_paragraph(); p.add_run("ambition", style).italic = True; p.add_run("【抱负】")
    red = p.add_run(" plain"); red.font.color.rgb = RGBColor(255, 0, 0)
    doc.save(path)
    assert [i.answer for i in DocxParser().parse(path).word_items] == ["ambition"]
    rules = RecognitionRules(format_mode="red", require_chinese=False)
    assert [i.answer for i in DocxParser().parse(path, rules).word_items] == ["plain"]
    assert DocxParser().parse(path, RecognitionRules("red_and_bold")).target_count == 0


def test_theme_slot_is_resolved_not_assumed_red(tmp_path):
    from docx.enum.dml import MSO_THEME_COLOR_INDEX
    from lxml import etree
    doc = Document()
    p = doc.add_paragraph(); r = p.add_run("word"); r.font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_2
    p.add_run("（单词）")
    theme = next(rel.target_part for rel in doc.part.rels.values() if rel.reltype.endswith('/theme'))
    xml = etree.fromstring(theme.blob)
    node = xml.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}accent2')[0]
    node.set('val', '0000FF'); theme._blob = etree.tostring(xml)
    path = tmp_path / 'theme.docx'; doc.save(path)
    assert DocxParser().parse(path).target_count == 0
    node.set('val', 'FF0000'); theme._blob = etree.tostring(xml); doc.save(path)
    assert DocxParser().parse(path).target_count == 1


def test_long_term_review_survives_history_deletion(tmp_path):
    repo = Repository(tmp_path / "db")
    session = new_full_session(DocxParser().parse(document(tmp_path))); grade_session(session)
    sid = repo.save_submitted_session(session)
    repo.delete_practice_session(sid)
    records = repo.get_review_examples([w['id'] for w in repo.list_wrong_words()])
    review = new_long_term_session(records)
    assert review.total == 2
    for item in review.article.word_items: review.answers[item.id] = item.answer
    grade_session(review); repo.save_submitted_session(review)
    assert all(w['consecutive_correct'] == 1 for w in repo.list_wrong_words())
