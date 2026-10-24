from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PySide6.QtCore import QRect
from PySide6.QtWidgets import QApplication, QFrame, QLabel, QWidget

from test_library_ui import ui
from src.database.repository import Repository
from src.parser.docx_parser import DocxParser
from src.parser.word_structure import iter_paragraphs, iter_runs
from src.services.exporter import export_answer, export_worksheet
from src.ui.widgets.flow_layout import FlowLayout
from src.ui.typography import is_article_heading


def uneven_article(tmp_path):
    doc = Document()
    doc.add_paragraph('The School Festival', 'Title')
    doc.add_paragraph('   ')
    doc.add_paragraph()
    p = doc.add_paragraph('Mia used a ')
    p.add_run('paintbrush').bold = True
    p.add_run('（画笔） to make a colourful picture. ' * 3)
    p.paragraph_format.line_spacing = Pt(65)
    p.paragraph_format.space_after = Pt(80)
    p.paragraph_format.first_line_indent = Pt(24)
    p._p.get_or_add_pPr().find(qn('w:ind')).set(qn('w:firstLineChars'), '200')
    mark = OxmlElement('w:rPr'); size = OxmlElement('w:sz'); size.set(qn('w:val'), '96')
    mark.append(size); p._p.get_or_add_pPr().append(mark)
    for run in p.runs: run.font.size = Pt(32)
    doc.add_paragraph()
    p = doc.add_paragraph('She decided to ')
    p.add_run('take advantage of').bold = True
    p.add_run('（利用） the afternoon.')
    path = tmp_path / 'uneven.docx'
    doc.save(path)
    return path


@pytest.mark.parametrize('zoom', [100, 120, 130, 180])
def test_display_skips_blank_cards_but_preserves_article_and_answers(tmp_path, ui, zoom):
    repo = Repository(tmp_path / 'db')
    w = ui(repo)
    w._import_paths([str(uneven_article(tmp_path))])
    snapshot = deepcopy(w.current_article)
    assert len(w.content.findChildren(QFrame, 'originalPage')) == 1
    cards = w.content.findChildren(QFrame, 'readingParagraph')
    assert len(cards) == 3
    assert all(c.property('paragraphText').strip() for c in cards)
    assert 'line-height:150%' in cards[1].findChild(QLabel).text()
    w.start_full_practice()
    w.resize(1000, 700); w.show(); w.set_zoom(zoom)
    QApplication.processEvents()
    assert len(w.answer_edits) == snapshot.target_count == 2
    first = next(iter(w.answer_edits.values()))
    first.setText('my spelling')
    w.show_answers(); w.show_answers()
    assert first.text() == 'my spelling'
    assert not w.session.viewed_answer
    assert w.session.article == snapshot
    assert w.scroll.horizontalScrollBar().maximum() == 0
    w._save_active_draft()
    restored = Repository(tmp_path / 'db').load_drafts()[0]
    assert restored.article == snapshot
    assert restored.answers[first.item_id] == 'my spelling'


@pytest.mark.parametrize('export', [export_answer, export_worksheet])
def test_export_normalizes_print_copy_only(tmp_path, export):
    source = uneven_article(tmp_path)
    before = source.read_bytes()
    article = DocxParser().parse(source)
    snapshot = deepcopy(article)
    output = export(article, tmp_path / 'out.docx')
    doc = Document(output)
    assert len(doc.paragraphs) == 3
    assert round(doc.sections[0].page_width.mm) == 210
    assert round(doc.sections[0].page_height.mm) == 297
    assert doc.paragraphs[0].paragraph_format.keep_with_next
    for p in doc.paragraphs[1:]:
        assert p.paragraph_format.line_spacing == 1.6
        assert p.paragraph_format.space_after.pt == 14
        assert all(r.font.size.pt == 13 for r in iter_runs(p))
        assert p._p.get_or_add_pPr().find(qn('w:ind')).get(qn('w:firstLineChars')) is None
        assert '（' in p.text
    assert doc.paragraphs[1]._p.find('.//' + qn('w:rPr')).find(qn('w:sz')).get(qn('w:val')) == '26'
    if export is export_answer:
        assert [p.text for p in doc.paragraphs] == [p.text for p in article.paragraphs if p.text.strip()]
        assert any(r.text == 'paintbrush' and r.bold for r in doc.paragraphs[1].runs)
    else:
        assert 'paintbrush' not in doc.paragraphs[1].text
        assert '________' in doc.paragraphs[1].text
    assert source.read_bytes() == before
    assert article == snapshot


def test_export_preserves_objects_and_structural_empty_paragraphs(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p._p.append(OxmlElement('w:bookmarkStart'))
    doc.add_paragraph().add_run().add_break()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_paragraph('table text')
    path = tmp_path / 'objects.docx'; doc.save(path)
    article = DocxParser().parse(path)
    result = Document(export_answer(article, tmp_path / 'out.docx'))
    assert result.element.body.find('.//' + qn('w:bookmarkStart')) is not None
    assert result.element.body.find('.//' + qn('w:br')) is not None
    assert len(result.tables[0].cell(0, 0).paragraphs) == 2
    assert len(list(iter_paragraphs(result))) == 4


def test_flow_centers_text_and_keeps_short_hint_with_input(tmp_path, ui):
    ui(Repository(tmp_path / 'db'))
    parent = QWidget(); flow = FlowLayout(parent, margin=0, h_spacing=2)
    prefix = QLabel('prefix'); prefix.setFixedSize(80, 20)
    field = QLabel('field'); field.setFixedSize(90, 40)
    hint = QLabel('（提示）'); hint.setFixedSize(60, 20); hint.setProperty('keepWithPrevious', True)
    for widget in (prefix, field, hint): flow.addWidget(widget)
    flow._do_layout(QRect(0, 0, 220, 0), False)
    assert field.y() > prefix.y()
    assert abs(field.geometry().center().y() - hint.geometry().center().y()) <= 1
    assert hint.x() > field.x()
    parent.deleteLater()


def test_title_requires_explicit_style():
    assert is_article_heading('Title')
    assert is_article_heading('Heading 1')
    assert is_article_heading('标题 2')
    assert not is_article_heading('Normal')
    assert not is_article_heading(None)
