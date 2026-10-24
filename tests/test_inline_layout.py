"""Actual flow-layout regressions for compact fields and Chinese annotations."""
import os
from pathlib import Path

import pytest
from PySide6.QtCore import QRect
from PySide6.QtGui import QFontDatabase
from PySide6.QtWidgets import QApplication, QLabel, QWidget
from docx import Document

from test_library_ui import ui
from src.ui.main_window import MainWindow
from src.ui.widgets.answer_line_edit import AnswerLineEdit
from src.ui.widgets.flow_layout import FlowLayout
from src.database.repository import Repository


@pytest.fixture
def fonts(ui):
    app = QApplication.instance()
    for name in ('msyh.ttc', 'segoeui.ttf'):
        path = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / name
        if path.exists():
            QFontDatabase.addApplicationFont(str(path))
    return app


def text_flow(app, text, width, size=40):
    parent = QWidget()
    parent.setStyleSheet(f"QWidget {{font-family:'Microsoft YaHei UI';font-size:{size}px;}}")
    flow = FlowLayout(parent, margin=16, h_spacing=1, v_spacing=12)
    class Factory:
        _text_label = staticmethod(MainWindow._text_label)
        _add_text = MainWindow._add_text
    Factory()._add_text(flow, text)
    parent.resize(width, 900); parent.show(); app.processEvents()
    flow.setGeometry(QRect(0, 0, width, 900))
    return parent, flow


def test_chinese_annotation_stays_one_line_when_it_fits(fonts):
    parent, flow = text_flow(fonts, 'most thrilling （令人激动的、惊险的） find ever,', 1200)
    cn = next(flow.itemAt(i).widget() for i in range(flow.count())
              if '令人' in flow.itemAt(i).widget().text())
    assert cn.width() >= cn.fontMetrics().horizontalAdvance(cn.text())
    assert cn.height() <= cn.fontMetrics().height() + 2
    parent.close(); parent.deleteLater()


@pytest.mark.parametrize('scale', [0.7, 1., 1.3, 1.8])
@pytest.mark.parametrize('answer,low,high', [('cat',120,170), ('take advantage of',180,250),
                                            ("As far as I'm concerned",220,320)])
def test_field_width_buckets_follow_zoom_and_keep_long_input(fonts, scale, answer, low, high):
    field = AnswerLineEdit('id', len(answer))
    field.setProperty('answerText', answer)
    field.setProperty('flowScale', scale)
    field.setStyleSheet(f'font-size:{round(20*scale)}px;')
    field.setText('an exceptionally long user answer ' * 6)
    field.update_preferred_width()
    width = field.property('preferredFlowWidth')
    assert round(low*scale) <= width <= round(high*scale)
    before = width
    field.setText('even longer input ' * 30)
    assert field.property('preferredFlowWidth') == before
    assert field.text() == 'even longer input ' * 30
    field.deleteLater()


def test_long_chinese_hint_wraps_as_one_annotation_without_clipping(fonts):
    text = '（一种需要充分说明而不能被截断或拆散到其他英文单词之间的中文释义）'
    parent, flow = text_flow(fonts, text + ' next word.', 250)
    cn = [flow.itemAt(i).widget() for i in range(flow.count())
          if '中文' in flow.itemAt(i).widget().text()]
    assert len(cn) == 1 and cn[0].text().strip() == text
    label = cn[0]
    assert label.height() >= label.heightForWidth(label.width())
    assert label.geometry().right() < parent.width()
    parent.close(); parent.deleteLater()


@pytest.mark.parametrize('width', [980, 1480, 2400])
@pytest.mark.parametrize('zoom', [70, 100, 130, 180])
def test_actual_practice_layout_preserves_fields_hints_and_chinese(tmp_path, ui, fonts, width, zoom):
    document = Document()
    document.add_paragraph('Layout Verification', 'Title')
    paragraph = document.add_paragraph("This was the most ")
    paragraph.add_run('thrilling').bold = True
    paragraph.add_run('（令人激动的、惊险的） find ever. ')
    paragraph.add_run("As far as I'm concerned").bold = True
    paragraph.add_run('（就我而言）, it was excellent.')
    path = tmp_path / 'layout.docx'; document.save(path)
    window = ui(Repository(tmp_path / 'data' / 'learning.db'))
    window._import_paths([str(path)]); window.start_full_practice()
    window.show(); window.resize(width, 1000); window.set_zoom(zoom)
    fonts.processEvents()
    answers = [edit.property('answerText') for edit in window.answer_edits.values()]
    assert answers == ['thrilling', "As far as I'm concerned"]
    first = next(iter(window.answer_edits.values()))
    first.setText('long user input ' * 25)
    first.setFocus(); fonts.processEvents()
    preferred = first.property('preferredFlowWidth')
    first.setText('long user input ' * 50)
    assert first.property('preferredFlowWidth') == preferred

    def check_layout():
        fonts.processEvents()
        assert window.scroll.horizontalScrollBar().maximum() == 0
        for flow in window.content.findChildren(FlowLayout):
            left, top, right, bottom = flow.getContentsMargins()
            available = flow.parentWidget().width() - left - right
            visible = [flow.itemAt(i).widget() for i in range(flow.count())
                       if not flow.itemAt(i).widget().isHidden()]
            for widget in visible:
                rect = widget.geometry()
                assert rect.left() >= left and rect.right() < flow.parentWidget().width() - right
                assert rect.bottom() < flow.parentWidget().height() - bottom
                if widget.hasHeightForWidth():
                    assert widget.height() >= widget.heightForWidth(widget.width())
                if isinstance(widget, AnswerLineEdit):
                    cap = 170 if len(widget.property('answerText').split()) == 1 else 320
                    assert widget.width() <= round(cap * zoom / 100)
                elif widget.text().startswith('（') and widget.fontMetrics().horizontalAdvance(widget.text()) + 2 <= available:
                    assert widget.height() <= widget.fontMetrics().height() + 2
                elif widget.property('flowAnswerHint'):
                    assert widget.width() <= round(280 * zoom / 100)
                    margins = widget.contentsMargins()
                    natural = widget.fontMetrics().horizontalAdvance(widget.text()) + margins.left() + margins.right() + 2
                    if natural <= widget.width():
                        assert widget.height() <= widget.fontMetrics().height() + margins.top() + margins.bottom() + 2
            for i, widget in enumerate(visible):
                for other in visible[i+1:]:
                    assert not widget.geometry().intersects(other.geometry())

    check_layout()
    window.show_answers(); check_layout()
    assert window.answer_hints[next(iter(window.answer_edits))].text() == '答案：thrilling'
    window.show_answers(); check_layout()
    assert first.text() == 'long user input ' * 50
    assert first.hasFocus()
    assert window._save_active_draft()
    assert window.repository.load_drafts()[0].answers[first.item_id] == first.text()
