import hashlib
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from docx import Document
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication, QLabel, QMessageBox, QFrame

from src.database.repository import Repository
from src.services.library import ArticleLibrary
from src.ui.main_window import MainWindow


@pytest.fixture
def ui(monkeypatch):
    app = QApplication.instance() or QApplication([])
    monkeypatch.setattr(QTimer, "singleShot", lambda *args: None)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.Yes)
    windows = []

    def create(repository):
        window = MainWindow(repository)
        # Most functional tests drive selection synchronously; stress tests enable the real timer.
        window._article_switch_delay = 0
        monkeypatch.setattr(window, "_confirm_finished_submission", lambda: True)
        windows.append(window)
        return window

    yield create
    for window in windows:
        window.close()
        window.deleteLater()
    app.processEvents()


def sample_word(tmp_path):
    path = tmp_path / "article 中文.docx"
    doc = Document()
    doc.add_paragraph().add_run("The Big Family Get-Together").bold = True
    p = doc.add_paragraph("I have ")
    p.add_run("ambition").bold = True
    p.add_run("（抱负，雄心）.")
    doc.save(path)
    return path


def test_reopen_restores_article_title_and_removal_persists(tmp_path, ui):
    path = sample_word(tmp_path)
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    db = tmp_path / "learning.db"
    first = ui(Repository(db))
    first._import_paths([str(path)])
    first.close()
    second = ui(Repository(db))
    second._restore_article_library()
    assert second.article_list.count() == 1
    assert second.articles[0].target_count == 1
    second.start_full_practice()
    assert len(second.answer_edits) == 1
    assert not second.session.submitted
    heading = next(card for card in second.content.findChildren(QFrame)
                   if card.property("paragraphText") == "The Big Family Get-Together")
    assert "".join(label.text() for label in heading.findChildren(QLabel)) == "The Big Family Get-Together"
    second.remove_selected_article()
    third = ui(Repository(db))
    third._restore_article_library()
    assert third.article_list.count() == 0
    assert hashlib.sha256(path.read_bytes()).hexdigest() == original_hash


def test_missing_file_is_not_forgotten(tmp_path, ui):
    repository = Repository(tmp_path / "learning.db")
    missing = tmp_path / "temporarily missing.docx"
    repository.save_article_path(missing)
    window = ui(repository)
    window._restore_article_library()
    assert window.articles == []
    assert repository.list_article_paths() == [str(missing.resolve())]


def test_failed_save_does_not_claim_article_was_imported(tmp_path, ui, monkeypatch):
    repository = Repository(tmp_path / "learning.db")
    window = ui(repository)
    def fail(*args):
        raise OSError("database unavailable")
    monkeypatch.setattr(repository, "save_article", fail)
    window._import_paths([str(sample_word(tmp_path))])
    assert window.articles == []
    assert repository.list_article_paths() == []


def test_legacy_doc_is_converted_to_internal_copy_without_touching_original(tmp_path, monkeypatch):
    source = tmp_path / "旧式 文章.doc"
    source.write_bytes(b"legacy-word-content")
    original = source.read_bytes()

    def fake_convert(input_path, destination):
        assert input_path == source.resolve()
        document = Document(); paragraph = document.add_paragraph("I have ")
        paragraph.add_run("ambition").bold = True
        paragraph.add_run("（抱负）.")
        document.save(destination)

    monkeypatch.setattr("src.services.legacy_doc.convert_doc_to_docx", fake_convert)
    repository = Repository(tmp_path / "data" / "learning.db")
    library = ArticleLibrary(repository)
    article = library.import_file(source)

    assert article.filename == source.name
    assert article.source_path == str(source.resolve())
    assert [item.answer for item in article.word_items] == ["ambition"]
    assert article.metadata["source_format"] == ".doc"
    assert Path(article.metadata["document_path"]).suffix == ".docx"
    assert source.read_bytes() == original
