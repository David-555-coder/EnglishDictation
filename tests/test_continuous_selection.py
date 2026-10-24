"""Real mouse/clipboard regression for sentences split into flow widgets."""
from copy import deepcopy
import os
from pathlib import Path

import pytest
from docx import Document
from PySide6.QtCore import QEvent, QPoint, Qt
from PySide6.QtGui import QFontDatabase
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFrame, QLabel, QLineEdit, QWidget, QVBoxLayout, QScrollArea

from test_library_ui import ui
from src.database.repository import Repository
from src.services.practice_service import new_full_session
from src.services.serialization import session_to_dict
from src.ui.widgets.text_menus import TextLabel, text_menu
from src.ui.widgets.reading_selection import ReadingSelection, units

SENTENCES = ('On the whole（总体而言）, it was a stress-free（无压力的） night.',
             'Her goal was to outdo（超过） her prior self & keep reading.')


@pytest.fixture
def reader(tmp_path, ui):
    for name in ('msyh.ttc','segoeui.ttf'):
        QFontDatabase.addApplicationFont(str(Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/name))
    path=tmp_path/'story.docx'; doc=Document()
    p=doc.add_paragraph('On the whole（总体而言）, it was a '); p.add_run('stress-free').bold=True
    p.add_run('（无压力的） night.')
    p=doc.add_paragraph('Her goal was to '); p.add_run('outdo').bold=True
    p.add_run('（超过） her prior self & keep reading.')
    doc.save(path)
    window=ui(Repository(tmp_path/'db')); window._import_paths([str(path)])
    window.resize(1100,900); window.show(); window.activateWindow(); QApplication.processEvents()
    return window


def labels(window):
    return [w for w in window.content.findChildren(TextLabel) if w.isVisibleTo(window.content)
            and w.property('flowText') and not w.property('flowAnswerHint') and w.objectName()!='learningNote']


def drag(first,last,start=None):
    if start is None: start=QPoint(1,first.height()//2)
    end=first.mapFromGlobal(last.mapToGlobal(QPoint(max(1,last.width()-1),last.height()//2)))
    QTest.mousePress(first,Qt.LeftButton,pos=start)
    QTest.mouseMove(first,end,10)
    QTest.mouseRelease(first,Qt.LeftButton,pos=end)
    QApplication.processEvents()


def copy_key(label):
    QApplication.clipboard().setText('sentinel')
    QTest.keyClick(label,Qt.Key_C,Qt.ControlModifier)
    return QApplication.clipboard().text()


def test_preview_drag_copies_whole_sentence_across_words(reader):
    window=reader; window._show_wrong_preview(new_full_session(window.current_article)); QApplication.processEvents()
    card=window.content.findChildren(QFrame,'paragraphCard')[0]
    words=[w for w in card.findChildren(TextLabel) if w.property('flowText')]
    drag(words[0],words[-1])
    assert copy_key(words[0]).strip()==SENTENCES[0]


def test_preview_drag_copies_across_sentence_cards(reader):
    window=reader; window._show_wrong_preview(new_full_session(window.current_article)); QApplication.processEvents()
    words=labels(window); drag(words[0],words[-1])
    assert copy_key(words[0]).strip()=='\n\n'.join(SENTENCES)


@pytest.mark.parametrize('zoom',[70,100,180])
def test_partial_sentence_and_partial_word_drag_positions(reader,zoom):
    window=reader;window._show_wrong_preview(new_full_session(window.current_article));window.set_zoom(zoom)
    QApplication.processEvents();words=labels(window)
    first=next(w for w in words if w.text()=='the ');last=next(w for w in words if w.text()=='was ')
    drag(first,last);assert copy_key(first).strip()=='the whole（总体而言）, it was'
    target=next(w for w in words if w.text()=='stress-free')
    drag(target,target,start=QPoint(target.fontMetrics().horizontalAdvance('stress-'),target.height()//2))
    assert copy_key(target)=='free'


@pytest.mark.parametrize('width,zoom',[(980,70),(980,180),(1480,100),(2400,130)])
def test_reverse_cross_card_selection_respects_zoom_and_line_wrapping(reader,width,zoom):
    window=reader; window._show_wrong_preview(new_full_session(window.current_article))
    window.resize(width,900);window.set_zoom(zoom);QApplication.processEvents()
    words=labels(window);first,last=words[0],words[-1]
    start=QPoint(last.width()-1,last.height()//2)
    end=last.mapFromGlobal(first.mapToGlobal(QPoint(1,first.height()//2)))
    QTest.mousePress(last,Qt.LeftButton,pos=start);QTest.mouseMove(last,end,10)
    QTest.mouseRelease(last,Qt.LeftButton,pos=end)
    assert copy_key(last).strip()=='\n\n'.join(SENTENCES)
    assert not window.reading_selection._scroll_timer.isActive()
    assert window.scroll.horizontalScrollBar().maximum()==0


def test_right_click_copy_and_select_all_cover_visible_sentence_page(reader):
    window=reader;window._show_wrong_preview(new_full_session(window.current_article));QApplication.processEvents()
    words=labels(window);drag(words[0],words[-1])
    menu=text_menu(words[-1]);next(a for a in menu.actions() if a.text()=='复制').trigger()
    assert QApplication.clipboard().text().strip()=='\n\n'.join(SENTENCES)
    menu.deleteLater();window.reading_selection.clear()
    menu=text_menu(words[0]);next(a for a in menu.actions() if a.text()=='全选').trigger()
    assert window.reading_selection.selected_text().strip()=='\n\n'.join(SENTENCES)
    menu.deleteLater()


def test_active_practice_copy_never_reads_hidden_answers_and_keeps_native_edits(reader):
    window=reader;window.start_full_practice();QApplication.processEvents()
    session=window.session; words=labels(window)
    before=deepcopy(session_to_dict(session))
    drag(words[0],words[-1]);text=copy_key(words[0])
    assert 'stress-free' not in text and 'outdo' not in text and '正确答案' not in text
    assert '____' in text and 'On the whole' in text and 'keep reading.' in text
    assert session_to_dict(session)==before
    edit=next(iter(window.answer_edits.values()));edit.setText('typed answer');edit.setFocus();edit.selectAll()
    QTest.keyClick(edit,Qt.Key_C,Qt.ControlModifier)
    assert QApplication.clipboard().text()=='typed answer'
    QTest.keyClicks(edit,'replacement');assert edit.text()=='replacement'
    QTest.keyClick(edit,Qt.Key_Return)
    assert QApplication.focusWidget() is list(window.answer_edits.values())[1]
    assert not session.view_state['memory']['exposed']


def test_wrong_reader_selection_does_not_change_results_or_model(reader):
    window=reader;window.start_full_practice();window._commit_submission();window.read_current_wrong()
    QApplication.processEvents();session=window.session;before=deepcopy(session_to_dict(session))
    with window.repository.connection() as db:
        tables={t:[tuple(r) for r in db.execute('SELECT * FROM '+t)] for t in
                ('wrong_words','practice_sessions','practice_answers','memory_events','memory_exposures','daily_review_items')}
    words=labels(window);drag(words[0],words[-1])
    assert copy_key(words[0]).strip()=='\n\n'.join(SENTENCES)
    assert session_to_dict(session)==before
    with window.repository.connection() as db:
        assert all([tuple(r) for r in db.execute('SELECT * FROM '+t)]==rows for t,rows in tables.items())
    window.return_from_wrong_reading();assert not window._wrong_reading and window.session is session


def test_submitted_answers_copy_only_visible_result_text(reader):
    window=reader;window.start_full_practice();window._commit_submission();QApplication.processEvents()
    first=labels(window)[0];QTest.keyClick(first,Qt.Key_A,Qt.ControlModifier)
    copied=copy_key(first)
    assert '正确答案：stress-free' in copied and '正确答案：outdo' in copied
    assert 'On the whole' in copied and 'keep reading.' in copied
    assert all(edit.isReadOnly() for edit in window.answer_edits.values())


def test_drag_can_start_on_readonly_result_answer_and_context_menu_copies_the_range(reader):
    window=reader;window.start_full_practice()
    edit=next(iter(window.answer_edits.values()));edit.setText('typed wrong')
    window._commit_submission();QApplication.processEvents()
    last=labels(window)[-1]
    drag(edit,last,start=QPoint(0,edit.height()//2))
    copied=copy_key(edit)
    assert copied.startswith('typed wrong') and '正确答案：stress-free' in copied and 'keep reading.' in copied
    menu=text_menu(edit);next(a for a in menu.actions() if a.text()=='复制').trigger()
    assert QApplication.clipboard().text()==copied
    assert all(not a.isEnabled() for a in menu.actions() if a.text() in ('剪切','粘贴','删除'))
    menu.deleteLater()


def test_original_rich_paragraphs_can_also_select_across_cards(reader):
    window=reader;window.resize(980,900);window.set_zoom(180);QApplication.processEvents()
    page=window.content.findChild(QFrame,'originalPage');paragraphs=page.findChildren(TextLabel)
    first,last=paragraphs[0],paragraphs[-1]
    start=QPoint(1,1);end=first.mapFromGlobal(last.mapToGlobal(QPoint(last.width()-1,last.height()-1)))
    QTest.mousePress(first,Qt.LeftButton,pos=start);QTest.mouseMove(first,end,10)
    QTest.mouseRelease(first,Qt.LeftButton,pos=end)
    assert copy_key(first).strip()=='\n\n'.join(SENTENCES)
    assert '<span' not in QApplication.clipboard().text()


def test_unicode_keyboard_selection_copy_and_escape(reader):
    window=reader;window._show_wrong_preview(new_full_session(window.current_article));QApplication.processEvents()
    first=labels(window)[0];first.setFocus()
    QTest.keyClick(first,Qt.Key_A,Qt.ControlModifier)
    QTest.keyClick(first,Qt.Key_Left,Qt.ShiftModifier)
    assert copy_key(first)=='\n\n'.join(SENTENCES)[:-1]
    QTest.keyClick(first,Qt.Key_Escape)
    assert window.reading_selection.selected_text()==''
    assert all(not w.hasSelectedText() for w in labels(window))


def test_double_click_word_and_page_rebuild_leave_no_stale_selection(reader):
    window=reader;window._show_wrong_preview(new_full_session(window.current_article));QApplication.processEvents()
    word=next(w for w in labels(window) if w.text().strip()=='goal')
    QTest.mouseDClick(word,Qt.LeftButton,pos=QPoint(4,word.height()//2))
    assert copy_key(word)=='goal'
    old=window.reading_selection._bound.copy()
    window.start_full_practice();QApplication.processEvents()
    assert window.reading_selection.selected_text()==''
    assert not any(w in window.reading_selection._bound for w in old)


def test_zoom_preserves_selected_sentence_range(reader):
    window=reader;window._show_wrong_preview(new_full_session(window.current_article));QApplication.processEvents()
    words=labels(window);drag(words[0],words[-1]);before=copy_key(words[0])
    for zoom in (180,70,130):
        window.set_zoom(zoom);QApplication.processEvents()
        assert copy_key(words[0])==before
        assert all(w.hasSelectedText() for w in labels(window))
        assert not window.reading_selection._scroll_timer.isActive()


def test_rich_unicode_multiline_label_uses_utf16_positions():
    app=QApplication.instance() or QApplication([])
    scroll=QScrollArea();root=QWidget();layout=QVBoxLayout(root);scroll.setWidget(root);scroll.setWidgetResizable(True)
    label=TextLabel('<p>Keep <b>专注😀</b> &amp; read.<br>Another line.</p>')
    label.setTextFormat(Qt.RichText);label.setWordWrap(True);label.setTextInteractionFlags(Qt.TextSelectableByMouse)
    layout.addWidget(label);selection=ReadingSelection(root,scroll);selection.bind();scroll.resize(300,150);scroll.show();app.processEvents()
    QTest.mousePress(label,Qt.LeftButton,pos=QPoint(1,1));QTest.mouseMove(label,QPoint(label.width()-1,label.height()-1),5)
    QTest.mouseRelease(label,Qt.LeftButton,pos=QPoint(label.width()-1,label.height()-1))
    assert copy_key(label)=='Keep 专注😀 & read.\nAnother line.'
    assert label.selectedText().replace('\u2028','\n')==selection.selected_text()
    selection.reset();scroll.close();scroll.deleteLater()


def test_drag_timer_stops_on_deactivation_and_no_idle_polling(reader):
    window=reader;window._show_wrong_preview(new_full_session(window.current_article));QApplication.processEvents()
    first=labels(window)[0];QTest.mousePress(first,Qt.LeftButton,pos=QPoint(1,1))
    assert window.reading_selection._scroll_timer.isActive()
    QApplication.sendEvent(window,QEvent(QEvent.WindowDeactivate))
    assert not window.reading_selection._dragging
    assert not window.reading_selection._scroll_timer.isActive()
    QTest.mouseRelease(first,Qt.LeftButton,pos=QPoint(1,1))


def test_drag_near_page_bottom_scrolls_only_while_pressed(reader):
    window=reader
    for _ in range(14):
        p=deepcopy(window.current_article.paragraphs[0]);p.index=len(window.current_article.paragraphs)
        window.current_article.paragraphs.append(p)
    window.show_preview();window.resize(980,450);window.set_zoom(130);QApplication.processEvents()
    paragraphs=window.content.findChild(QFrame,'originalPage').findChildren(TextLabel)
    first=paragraphs[0];bar=window.scroll.verticalScrollBar();bar.setValue(0)
    start=QPoint(1,1);end=first.mapFromGlobal(window.scroll.viewport().mapToGlobal(
        QPoint(window.scroll.viewport().width()-8,window.scroll.viewport().height()-2)))
    QTest.mousePress(first,Qt.LeftButton,pos=start);QTest.mouseMove(first,end,10)
    before=bar.value();QTest.qWait(150);QApplication.processEvents()
    assert bar.value()>before
    QTest.mouseRelease(first,Qt.LeftButton,pos=end)
    assert not window.reading_selection._scroll_timer.isActive()
    settled=bar.value();QTest.qWait(100);assert bar.value()==settled
