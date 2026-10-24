import json
import sqlite3
import time
from pathlib import Path

from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtCore import QPoint, QPointF, Qt
from PySide6.QtGui import QWheelEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QDialog, QFrame, QLabel
from docx import Document

from src.database.repository import Repository
from src.models.article import ArticleSummary
from src.ui.dialogs import WrongWordsDialog, WrongPracticeHistoryDialog
from test_library_ui import ui, sample_word


def test_unchanged_save_skips_database_and_keeps_failure_retry(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / "db")
    w = ui(repo); w._import_paths([str(sample_word(tmp_path))]); w.start_full_practice()
    w._save_active_draft()
    original = repo.save_draft
    calls = []
    def save(session):
        calls.append(session.session_key)
        original(session)
    monkeypatch.setattr(repo, "save_draft", save)
    assert w._save_active_draft() and w._save_active_draft()
    assert calls == []
    edit = next(iter(w.answer_edits.values())); edit.setText("unsaved")
    def fail(session): raise sqlite3.OperationalError("disk unavailable")
    monkeypatch.setattr(repo, "save_draft", fail)
    assert not w._save_active_draft()
    monkeypatch.setattr(repo, "save_draft", save)
    assert w._save_active_draft() and len(calls) == 1
    assert repo.load_drafts()[0].answers[edit.item_id] == "unsaved"
    w.session.viewed_answer = True
    assert w._save_active_draft() and len(calls) == 2
    assert repo.load_drafts()[0].viewed_answer


def test_view_save_never_postpones_answer_save_or_queries_history(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / "db")
    w = ui(repo); w._import_paths([str(sample_word(tmp_path))]); w.start_full_practice()
    calls = []
    monkeypatch.setattr(repo, "has_wrong_history", lambda: calls.append(1) or False)
    edit = next(iter(w.answer_edits.values())); edit.setText("my input")
    assert w._autosave.isActive()
    w._queue_view_save()
    assert w._autosave.isActive() and w._view_save.isActive()
    assert w._view_save.interval() > w._autosave.interval()
    assert w._save_active_draft()
    assert not w._view_save.isActive() and not w._autosave.isActive()
    assert calls == []


def test_catalog_materializes_only_selected_article_and_survives_missing_word(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / "data" / "db")
    w = ui(repo)
    paths = []
    for i in range(4):
        folder = tmp_path / str(i); folder.mkdir()
        paths.append(str(sample_word(folder)))
    w._import_paths(paths); w.close()
    for path in paths: Path(path).unlink()
    calls = []
    original = repo.load_article
    def load(path): calls.append(path); return original(path)
    monkeypatch.setattr(repo, "load_article", load)
    reopened = ui(repo); reopened._restore_article_library()
    assert len(calls) == 1
    assert sum(isinstance(a, ArticleSummary) for a in reopened.articles) == 3
    reopened.article_list.setCurrentRow(2)
    assert len(calls) == 2
    reopened.start_full_practice()
    assert len(reopened.answer_edits) == 1
    reopened.article_list.setCurrentRow(0)
    assert len(calls) == 2


def test_startup_previews_article_and_leaves_draft_unchanged(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / "db")
    w = ui(repo); w._import_paths([str(sample_word(tmp_path))]); w.start_full_practice()
    next(iter(w.answer_edits.values())).setText("restored")
    w.close()
    reopened = ui(repo)
    before = repo.load_drafts()[0].answers.copy()
    reopened._restore_article_library()
    assert reopened.session is None and not reopened.answer_edits
    assert reopened.mode_title.text().startswith('原文模式')
    assert repo.load_drafts()[0].answers == before


def test_closed_history_and_wordbook_are_destroyed(tmp_path, ui, monkeypatch):
    w = ui(Repository(tmp_path / "db"))
    monkeypatch.setattr(QDialog, "exec", lambda self: QDialog.Rejected)
    for _ in range(12):
        w.show_history(); w.show_wrong_words()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    assert not w.findChildren(WrongWordsDialog)
    assert not w.findChildren(WrongPracticeHistoryDialog)


def test_invalid_settings_fall_back_without_deleting_data(tmp_path, ui):
    repo = Repository(tmp_path / "db")
    with repo.connection() as db:
        db.execute("INSERT INTO app_settings VALUES ('zoom', '{broken')")
    assert ui(repo).zoom_percent == 100
    repo.set_setting("zoom", [120])
    repo.set_setting("rules", {"format_mode": "unsupported"})
    w = ui(repo)
    assert w.zoom_percent == 100 and w._default_rules().require_chinese
    with repo.connection() as db:
        assert json.loads(db.execute("SELECT value FROM app_settings WHERE key='zoom'").fetchone()[0]) == [120]


def test_catalog_keeps_corrupt_and_legacy_records(tmp_path):
    repo = Repository(tmp_path / "db")
    with repo.connection() as db:
        db.execute("INSERT INTO article_library(source_path, filename, added_at, snapshot_json) VALUES ('bad','bad.docx','now','{bad')")
    rows, missing = repo.load_article_catalog()
    assert rows == [] and missing == ['bad']
    assert repo.list_article_paths() == ['bad']


def test_real_autosave_timer_persists_without_grading(tmp_path, ui):
    repo = Repository(tmp_path / "db")
    w = ui(repo); w._import_paths([str(sample_word(tmp_path))]); w.start_full_practice()
    edit = next(iter(w.answer_edits.values())); edit.setText("my answer")
    w._queue_view_save()
    QTest.qWait(750)
    assert repo.load_drafts()[0].answers[edit.item_id] == "my answer"
    assert not w.session.submitted and not w.session.results


def test_wheel_zoom_batches_relayout_without_replacing_inputs(tmp_path, ui, monkeypatch):
    w = ui(Repository(tmp_path / "db")); w._import_paths([str(sample_word(tmp_path))]); w.start_full_practice()
    edit = next(iter(w.answer_edits.values())); edit.setText("keep this")
    count = []
    w._zoom_timer.timeout.disconnect()
    w._zoom_timer.timeout.connect(lambda: count.append(w.zoom_percent))
    for _ in range(4):
        event = QWheelEvent(QPointF(10, 10), QPointF(10, 10), QPoint(), QPoint(0, 120),
                            Qt.NoButton, Qt.ControlModifier, Qt.NoScrollPhase, False)
        assert w.eventFilter(w.scroll.viewport(), event)
    assert count == []
    QTest.qWait(90)
    assert count == [140]
    w.set_zoom(140)
    assert next(iter(w.answer_edits.values())) is edit and edit.text() == "keep this"


def test_large_submission_batches_result_widgets(tmp_path, ui):
    path = tmp_path / "many words.docx"
    document = Document()
    for index in range(100):
        paragraph = document.add_paragraph(f"Sentence {index}: ")
        paragraph.add_run("vocabulary").bold = True
        paragraph.add_run(f"（词汇{index}）.")
    document.save(path)

    window = ui(Repository(tmp_path / "db"))
    window._import_paths([str(path)])
    window.start_full_practice()
    started = time.perf_counter()
    window._commit_submission()
    elapsed = time.perf_counter() - started

    panels = [panel for panel in window.content.findChildren(QFrame)
              if panel.objectName() == "wrongSummary"]
    assert elapsed < 3.0
    assert len(panels) == 1
    assert len(panels[0].findChildren(QLabel)) == 2  # 标题 + 一个批量摘要文本
