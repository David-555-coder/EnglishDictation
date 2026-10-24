from copy import deepcopy
import hashlib
import json

import pytest
from docx import Document
from PySide6.QtWidgets import QDialog, QLabel, QMessageBox

from src.database.repository import Repository
from src.services.library import ArticleLibrary
from src.services.practice_service import new_full_session, new_wrong_session, new_saved_wrong_session
from src.services.grading import grade_session
from src.services.learning_content import (prepare_article, edit_hint, edit_paragraph, cancel_blank,
                                          add_blank, note_for, context_for)
from src.services.serialization import session_to_dict
from src.ui.editing_dialogs import TargetEditorDialog
from test_library_ui import ui


def story(tmp_path):
    path = tmp_path / 'story.docx'
    doc = Document(); doc.add_paragraph('The Story', 'Title')
    p = doc.add_paragraph('We '); p.add_run('hold').bold = True
    p.add_run("（容纳） everyone's attention, then ")
    p.add_run('hold').bold = True; p.add_run('（握住） the rope. ')
    p.add_run('internal').bold = True; p.add_run('（内部的） rules matter.')
    doc.save(path)
    repo = Repository(tmp_path / 'data' / 'db')
    article = ArticleLibrary(repo).import_file(path)
    return repo, article, path


def submitted(repo, article):
    session = new_full_session(article); grade_session(session); repo.save_submitted_session(session)
    return session


def update(repo, article, note='internal / inherent'):
    before = prepare_article(article); after = deepcopy(before)
    edit_hint(after, after.word_items[0].id, '保持大家的注意力', note)
    assert repo.save_learning_content(before, after) == 1
    return after


def test_sentence_specific_hint_updates_body_and_preserves_other_occurrence(tmp_path):
    repo, article, path = story(tmp_path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    after = update(repo, article)
    first, second, third = after.word_items
    assert first.chinese == '保持大家的注意力' and second.chinese == '握住'
    assert 'hold（保持大家的注意力）' in first.paragraph_text
    assert all(i.paragraph_text[i.start:i.end] == i.answer for i in after.word_items)
    assert context_for(first)['key'] == context_for(article.word_items[0])['key']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    assert repo.load_article(article.source_path).word_items[0].chinese == first.chinese


def test_history_memory_and_old_draft_unchanged_but_future_examples_corrected(tmp_path):
    repo, article, _ = story(tmp_path)
    parent = submitted(repo, article)
    draft = new_wrong_session(parent); repo.save_draft(draft)
    frozen = session_to_dict(draft)
    with repo.connection() as db:
        old = {t: [tuple(r) for r in db.execute(f'SELECT * FROM {t}')] for t in ('practice_sessions','practice_answers','wrong_words','memory_events','practice_drafts')}
        ids = [r[0] for r in db.execute('SELECT id FROM wrong_words')]
    update(repo, article)
    with repo.connection() as db:
        assert {t: [tuple(r) for r in db.execute(f'SELECT * FROM {t}')] for t in old} == old
    assert session_to_dict(repo.load_drafts()[0]) == frozen
    correction = repo.prepare_learning_session(new_wrong_session(parent))
    assert correction.article.word_items[0].chinese == '保持大家的注意力'
    assert parent.article.word_items[0].chinese == '容纳'
    examples = repo.get_review_examples(ids)
    # The most recent hold occurrence is the second sense, which stays unchanged.
    assert next(r for r in examples if r['answer']=='hold')['chinese'] == '握住'
    original = repo.get_wrong_answers(parent.database_id)
    renewed = repo.prepare_learning_session(new_saved_wrong_session({'id':parent.database_id},original))
    assert renewed.article.word_items[0].chinese == '保持大家的注意力'


def test_origin_survives_reconstructed_history_and_second_correction(tmp_path):
    repo, article, _ = story(tmp_path)
    parent = submitted(repo, article)
    after = update(repo, article)
    next_round = repo.prepare_learning_session(new_wrong_session(parent))
    grade_session(next_round); repo.save_submitted_session(next_round)
    before = repo.load_article(article.source_path); latest = prepare_article(before)
    edit_hint(latest, latest.word_items[0].id, '吸引注意力', '新的备注')
    repo.save_learning_content(before, latest)
    rows = repo.get_wrong_answers(next_round.database_id)
    restored = repo.prepare_learning_session(new_saved_wrong_session({'id':next_round.database_id},rows))
    assert restored.article.word_items[0].chinese == '吸引注意力'
    assert note_for(restored.article.word_items[0]) == '新的备注'


def test_paragraph_edit_maps_unicode_offsets_and_changes_target_identity(tmp_path):
    repo, article, _ = story(tmp_path)
    before = prepare_article(article); after = deepcopy(before)
    p = after.paragraphs[1]
    keys = [context_for(i)['key'] for i in after.word_items]
    edit_paragraph(after, p.index, '😀 Today, '+p.text.replace('internal','inherent'))
    assert all(i.paragraph_text[i.start:i.end] == i.answer for i in after.word_items)
    assert [context_for(i)['key'] for i in after.word_items[:2]] == keys[:2]
    assert context_for(after.word_items[-1])['key'] != keys[-1]
    submitted(repo, article); repo.save_learning_content(before, after)
    with repo.connection() as db:
        word_id = db.execute("SELECT id FROM wrong_words WHERE answer='internal'").fetchone()[0]
        assert db.execute("SELECT count(*) FROM wrong_words WHERE answer='inherent'").fetchone()[0] == 0
    assert repo.get_review_examples([word_id]) == []
    assert new_full_session(repo.load_article(article.source_path)).article.word_items[-1].answer == 'inherent'


def test_reparse_and_reimport_keep_local_revisions(tmp_path):
    repo, article, path = story(tmp_path); update(repo, article)
    from src.models.rules import RecognitionRules
    library = ArticleLibrary(repo)
    assert library.reparse(repo.load_article(article.source_path), RecognitionRules()).word_items[0].chinese == '保持大家的注意力'
    assert library.import_file(path).word_items[0].chinese == '保持大家的注意力'


def test_cancel_blank_does_not_retarget_old_word_to_another_answer(tmp_path):
    repo, article, _ = story(tmp_path); parent = submitted(repo, article)
    before = prepare_article(article); after = deepcopy(before)
    cancel_blank(after, after.word_items[-1].id); repo.save_learning_content(before, after)
    revised = repo.prepare_learning_session(new_wrong_session(parent))
    assert all(i.answer != 'internal' for i in revised.article.word_items)
    assert len(revised.target_ids) == 2
    assert len(parent.target_ids) == 3


def test_editor_starts_at_target_and_save_applies_fields_without_extra_click(tmp_path, ui):
    repo, article, _ = story(tmp_path); window = ui(repo)
    editor = TargetEditorDialog(article, window)
    assert editor._paragraph_index == article.word_items[0].paragraph_index
    assert not hasattr(editor,'paragraphs') and not hasattr(editor,'targets')
    editor.chinese_edit.setText('吸引注意力'); editor.note_edit.setPlainText('保留具体原句')
    editor.accept()
    assert editor.result() == QDialog.Accepted
    assert editor.article.word_items[0].chinese == '吸引注意力'
    assert note_for(editor.article.word_items[0]) == '保留具体原句'
    assert article.word_items[0].chinese == '容纳'


def test_cancelled_editor_still_records_revelation_and_protects_existing_input(tmp_path, ui, monkeypatch):
    repo, article, _ = story(tmp_path); window = ui(repo)
    window._restore_article_library(); window.start_full_practice()
    first, second, third = window.session.target_ids
    window.answer_edits[first].setText('hold')
    snapshot = window.session
    monkeypatch.setattr(TargetEditorDialog,'exec',lambda self: QDialog.Rejected)
    window.edit_targets()
    assert window.session is snapshot and snapshot.answers[first] == 'hold'
    from src.services.learning_evidence import kind_for_item
    assert kind_for_item(snapshot,first) in {'first','study'}
    assert kind_for_item(snapshot,second) == 'exposed'
    assert repo.load_drafts()[0].view_state['memory']['exposed']
    with repo.connection() as db:
        assert db.execute('SELECT count(*) FROM learning_content_edits').fetchone()[0] == 0
        assert db.execute("SELECT count(*) FROM memory_exposures WHERE event_key LIKE 'content-editor:%'").fetchone()[0] > 0


def test_edit_during_round_preserves_inputs_then_new_round_uses_revision(tmp_path, ui, monkeypatch):
    repo, article, _ = story(tmp_path); window = ui(repo); window._restore_article_library(); window.start_full_practice()
    session = window.session; first = session.target_ids[0]
    window.answer_edits[first].setText('attempt')
    def accepted(dialog):
        dialog.chinese_edit.setText('吸引注意力'); dialog.accept(); return dialog.result()
    monkeypatch.setattr(TargetEditorDialog,'exec',accepted)
    window.edit_targets()
    assert window.session is session and window.answer_edits[first].text() == 'attempt'
    assert session.article.item_map[first].chinese == '容纳'
    assert repo.load_article(article.source_path).word_items[0].chinese == '吸引注意力'
    window.start_full_practice()
    assert window.session is not session and window.session.article.word_items[0].chinese == '吸引注意力'


def test_notes_hidden_during_dictation_visible_on_feedback_and_reading(tmp_path, ui):
    repo, article, _ = story(tmp_path); update(repo,article)
    window = ui(repo); window._restore_article_library(); window.start_full_practice()
    notes = window.content.findChildren(QLabel,'learningNote')
    assert notes and all(n.isHidden() for n in notes)
    window.show_answers(); assert all(not n.isHidden() for n in notes)
    window.show_answers(); assert all(n.isHidden() for n in notes)
    window._commit_submission(); assert all(not n.isHidden() for n in notes)
    assert '原词回忆结果' in window.score_text.text()
    window.read_current_wrong()
    assert window.content.findChildren(QLabel,'learningNote')


def test_add_cancel_and_undo_preserve_preview_integrity(tmp_path, ui):
    repo, article, _ = story(tmp_path); window=ui(repo)
    editor=TargetEditorDialog(article,window)
    editor._remove(); assert editor.article.target_count == 2
    editor._undo_change(); assert editor.article.target_count == 3
    p=editor.article.paragraphs[1]; start=p.text.index('rules')
    add_blank(editor.article,p.index,start,start+5)
    assert editor.article.target_count==4
    assert all(i.paragraph_text[i.start:i.end]==i.answer for i in editor.article.word_items)
    editor.article=deepcopy(editor.original_article);editor._undo=[];editor.reject()


@pytest.mark.parametrize('old,new,answer,expected', [
    ('I see cat.','I see nothing.','cat','nothing'),
    ('A cat runs.','A cats runs.','cat','cats'),
    ('Be kind.','Be unkind.','kind','unkind'),
])
def test_body_replacements_never_create_word_fragments(old,new,answer,expected):
    from src.models.article import Article,Paragraph
    article=prepare_article(Article('source','story',[Paragraph(0,old,[])],[]))
    add_blank(article,0,old.index(answer),old.index(answer)+len(answer))
    edit_paragraph(article,0,new)
    assert article.word_items[0].answer==expected


def test_large_rewrite_requires_reselection_instead_of_guessing():
    from src.models.article import Article,Paragraph
    article=prepare_article(Article('source','story',[Paragraph(0,'The quick cat ran.',[])],[]))
    add_blank(article,0,10,13)
    with pytest.raises(ValueError,match='无法可靠对应'):
        edit_paragraph(article,0,'The slow dog walked.')
    assert article.word_items[0].answer=='cat'


def test_partial_history_correction_does_not_remove_unseen_targets_on_reimport(tmp_path):
    from src.models.article import Article,Paragraph
    repo,article,path=story(tmp_path)
    item=deepcopy(article.word_items[0]); item.learning_context=context_for(item);item.paragraph_index=0
    partial=Article(article.source_path,article.filename,[Paragraph(0,item.paragraph_text,[])],[item],{'learning_partial':True})
    before=prepare_article(partial);after=deepcopy(before)
    edit_hint(after,item.id,'吸引注意力','仅本句')
    repo.delete_article_paths([article.source_path])
    repo.save_learning_content(before,after)
    restored=ArticleLibrary(repo).import_file(path)
    assert restored.target_count==3
    assert [i.chinese for i in restored.word_items]==['吸引注意力','握住','内部的']


def test_pending_preview_and_daily_examples_use_latest_revision(tmp_path,ui,monkeypatch):
    repo,article,_=story(tmp_path);parent=submitted(repo,article)
    window=ui(repo);window._restore_article_library()
    window._show_wrong_preview(new_wrong_session(parent))
    def accepted(dialog):
        dialog.chinese_edit.setText('吸引注意力');dialog.accept();return dialog.result()
    monkeypatch.setattr(TargetEditorDialog,'exec',accepted)
    window.edit_targets()
    assert window.pending_review_session.article.word_items[0].chinese=='吸引注意力'
    window.start_full_practice()
    assert window.session.article.word_items[0].chinese=='吸引注意力'
    current=repo.load_article(article.source_path);after=prepare_article(current)
    edit_hint(after,after.word_items[-1].id,'里面的','internal / inherent')
    repo.save_learning_content(current,after)
    with repo.connection() as db: word_id=db.execute("SELECT id FROM wrong_words WHERE answer='internal'").fetchone()[0]
    example=repo.get_review_examples([word_id])[0]
    assert example['chinese']=='里面的'
    assert 'internal / inherent' in example['learning_context']


def test_merge_annotation_ranges_and_split_preview(tmp_path,ui,monkeypatch):
    from src.models.article import Article,Paragraph
    from src.services.learning_content import merge_blank
    from PySide6.QtWidgets import QInputDialog
    text='We hold（保持） attention（注意力） today.'
    article=prepare_article(Article('source','story',[Paragraph(0,text,[])],[]))
    add_blank(article,0,text.index('hold'),text.index('hold')+4)
    add_blank(article,0,text.index('attention'),text.index('attention')+9)
    merge_blank(article,0,text.index('hold'),text.index('attention')+9,'保持注意力')
    assert article.target_count==1 and article.word_items[0].answer=='hold attention'
    assert article.paragraphs[0].text=='We hold attention（保持注意力） today.'
    repo=Repository(tmp_path/'db');window=ui(repo);editor=TargetEditorDialog(article,window)
    monkeypatch.setattr(QInputDialog,'getMultiLineText',lambda *a:('hold | 保持\nattention | 注意力',True))
    editor._split()
    assert [i.answer for i in editor.article.word_items]==['hold','attention']
    assert [i.chinese for i in editor.article.word_items]==['保持','注意力']
    assert all(i.paragraph_text[i.start:i.end]==i.answer for i in editor.article.word_items)
    editor._undo_change();assert editor.article.target_count==1
    editor.article=deepcopy(editor.original_article);editor._undo=[];editor.reject()


@pytest.mark.parametrize('suffix', ['（容纳）','[容纳]','【容纳】',': 容纳'])
def test_supported_gloss_styles_replaced_without_leaving_old_chinese(suffix):
    from src.models.article import Article,Paragraph
    text='We hold'+suffix+' attention.'
    article=prepare_article(Article('source','story',[Paragraph(0,text,[])],[]))
    add_blank(article,0,3,7)
    edit_hint(article,article.word_items[0].id,'保持','备注')
    assert article.paragraphs[0].text=='We hold（保持） attention.'


def test_note_edit_keeps_english_parenthetical_context():
    from src.models.article import Article,Paragraph
    text='We hold (as expected) attention.'
    article=prepare_article(Article('source','story',[Paragraph(0,text,[])],[]))
    add_blank(article,0,3,7)
    edit_hint(article,article.word_items[0].id,'','some note')
    assert article.paragraphs[0].text==text


def test_legacy_sentence_only_example_can_match_unique_source_sentence(tmp_path):
    repo,article,_=story(tmp_path)
    before=prepare_article(article);after=deepcopy(before)
    item=after.word_items[-1];old=deepcopy(item)
    edit_hint(after,item.id,'里面的','内部 / 固有')
    repo.save_learning_content(before,after)
    sentence=old.sentence.strip();start=sentence.index('internal')
    row={'source_path':article.source_path,'source_article':article.filename,'paragraph_index':-1,
         'paragraph_text':sentence,'sentence':sentence,'item_start':start,'answer':'internal','chinese':'内部的'}
    from src.services.learning_content import ContentResolver
    with repo.connection() as db: revised=ContentResolver(db).resolve(row)
    assert revised['chinese']=='里面的'
