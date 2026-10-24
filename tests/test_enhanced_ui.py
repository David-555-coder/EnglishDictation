import time

from PySide6.QtCore import QItemSelectionModel, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFrame, QMessageBox, QLabel
from docx import Document

from test_library_ui import ui, sample_word
from test_enhanced_services import document
from src.database.repository import Repository
from src.models.practice import PracticeMode
from src.ui.dialogs import WrongPracticeHistoryDialog, WrongWordsDialog


def test_close_reopen_restores_input_manually_even_if_original_moved(tmp_path, ui, monkeypatch):
    path = sample_word(tmp_path)
    repo = Repository(tmp_path / "data" / "db")
    first = ui(repo); first._import_paths([str(path)]); first.start_full_practice()
    item_id = first.session.target_ids[0]
    first.answer_edits[item_id].setText("Ambition")
    assert first.answer_edits[item_id].property("answerState") is None
    assert not first.session.submitted and not first.session.results
    first.close(); path.unlink()
    reopened = ui(Repository(repo.database_path)); reopened._restore_article_library()
    assert reopened.session is None and not reopened.answer_edits
    from src.ui.unfinished_dialog import UnfinishedPracticeDialog
    from PySide6.QtWidgets import QDialog
    monkeypatch.setattr(UnfinishedPracticeDialog, 'exec', lambda dialog: QDialog.Accepted)
    reopened.resume_draft()
    assert reopened.answer_edits[item_id].text() == "Ambition"
    assert not reopened.session.submitted
    assert reopened.repository.list_history() == []


def test_autosave_is_silent_but_still_persists(tmp_path, ui):
    repo = Repository(tmp_path / "data" / "db")
    window = ui(repo); window._import_paths([str(sample_word(tmp_path))]); window.start_full_practice()
    item_id = window.session.target_ids[0]
    window.statusBar().clearMessage()
    window.answer_edits[item_id].setText("unfinished")
    window._save_active_draft()
    assert "草稿已自动保存" not in window.statusBar().currentMessage()
    assert repo.load_drafts()[0].answers[item_id] == "unfinished"


def test_switching_articles_preserves_inputs_and_zoom_does_not_rebuild(tmp_path, ui, monkeypatch):
    a = sample_word(tmp_path)
    second_dir = tmp_path / 'second'; second_dir.mkdir()
    b = sample_word(second_dir)
    window = ui(Repository(tmp_path / "data" / "db"))
    window._import_paths([str(a), str(b)]); window.start_full_practice()
    first_id = window.session.target_ids[0]
    edit = window.answer_edits[first_id]; edit.setText("ambitio")
    window.set_zoom(130)
    assert window.answer_edits[first_id] is edit
    assert edit.text() == "ambitio"
    window.article_list.setCurrentRow(1)
    window.article_list.setCurrentRow(0)
    assert window.session is None and not window.answer_edits
    from src.ui.unfinished_dialog import UnfinishedPracticeDialog
    from PySide6.QtWidgets import QDialog
    monkeypatch.setattr(UnfinishedPracticeDialog, 'exec', lambda dialog: QDialog.Accepted)
    window.resume_draft()
    assert window.answer_edits[first_id].text() == "ambitio"


def test_actual_submit_history_restore_retry_and_all_correct(tmp_path, ui, monkeypatch):
    monkeypatch.setattr(QMessageBox, "information", lambda *a: QMessageBox.Ok)
    window = ui(Repository(tmp_path / "data" / "db"))
    window._import_paths([str(document(tmp_path))]); window.start_full_practice(); window.submit()
    sid = window.session.database_id
    assert window.session.wrong_count == 2
    window.start_saved_wrong_practice(sid)
    assert window.session is None and window.pending_review_session.total == 2
    assert len(window.answer_edits) == 0
    preview_text = "".join(label.text() for label in window.content.findChildren(QLabel))
    assert "extensive" in preview_text and "momentum" in preview_text
    window.start_full_practice()
    assert len(window.answer_edits) == 2
    cards = [c for c in window.content.findChildren(QFrame) if c.property("paragraphText")]
    # Old widgets are deleteLater; inspect current layout cards only.
    texts = [window.content_layout.itemAt(i).widget().property('paragraphText')
             for i in range(window.content_layout.count()) if window.content_layout.itemAt(i).widget()]
    assert len([text for text in texts if text]) == 1
    assert "Another sentence" not in texts[0]
    first = window.session.article.word_items[0]
    window.answer_edits[first.id].setText(first.answer); window.submit()
    assert window.session.wrong_count == 1
    window.retry_current_wrong()
    assert window.pending_review_session is None and window.session.total == 1
    assert window.session.view_state['consolidation']
    assert not window._answers_visible  # Extra group begins closed-book, not another preview.
    assert len(window.answer_edits) == 1
    last_id = window.session.target_ids[0]
    window.answer_edits[last_id].setText(window.session.article.item_map[last_id].answer)
    window.submit()
    assert window.session.correct_count == 1 and not window.session.wrong_ids
    assert not window.primary_action_button.isEnabled()
    assert window.primary_action_button.text() == "本轮已全部正确"


def test_failed_result_save_is_recoverable_after_reopen(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / 'data' / 'db')
    window = ui(repo); window._import_paths([str(sample_word(tmp_path))]); window.start_full_practice()
    def fail(*args): raise OSError("simulated unavailable database")
    monkeypatch.setattr(repo, 'save_submitted_session', fail)
    window.submit()
    assert window.session.submitted and window.session.database_id is None
    window.close()
    reopened = ui(Repository(repo.database_path)); reopened._restore_article_library()
    assert reopened.session is None
    from src.ui.unfinished_dialog import UnfinishedPracticeDialog
    from PySide6.QtWidgets import QDialog
    monkeypatch.setattr(UnfinishedPracticeDialog, 'exec', lambda dialog: QDialog.Accepted)
    reopened.resume_draft()
    assert reopened.session.submitted
    assert all(edit.isReadOnly() for edit in reopened.answer_edits.values())
    reopened.retry_save_results(); reopened.retry_save_results()
    assert len(reopened.repository.list_history()) == 1
    assert reopened.repository.list_wrong_words()[0]['error_count'] == 1


def test_preview_does_not_mark_current_draft(tmp_path, ui, monkeypatch):
    from src.ui.unfinished_dialog import UnfinishedPracticeDialog
    from PySide6.QtWidgets import QDialog
    monkeypatch.setattr(UnfinishedPracticeDialog, "exec", lambda _: QDialog.Accepted)
    window = ui(Repository(tmp_path / 'data' / 'db'))
    window._import_paths([str(sample_word(tmp_path))]); window.start_full_practice()
    window.show_preview()
    assert window.session is None
    window.resume_draft()
    assert not window.session.viewed_answer
    assert not window.repository.load_drafts()[0].viewed_answer


def test_long_term_dialog_launches_selected_word(tmp_path, ui, monkeypatch):
    window = ui(Repository(tmp_path / 'data' / 'db'))
    window._import_paths([str(sample_word(tmp_path))]); window.start_full_practice(); window.submit()
    dialog = WrongWordsDialog(window.repository, window)
    dialog.practiceRequested.connect(window.start_long_term_practice)
    dialog.table.selectRow(0); dialog._practice_selected()
    assert window.session is None
    assert window.pending_review_session.mode == PracticeMode.LONG_TERM_WRONG
    window.start_full_practice()
    assert window.session.mode == PracticeMode.LONG_TERM_WRONG
    assert len(window.answer_edits) == 1
    dialog.deleteLater()


def test_background_import_continues_past_bad_file(tmp_path, ui):
    good = sample_word(tmp_path)
    bad = tmp_path / "broken.docx"; bad.write_bytes(b"not a word file")
    window = ui(Repository(tmp_path / 'data' / 'db'))
    window._start_import([str(bad), str(good)])
    deadline = time.monotonic() + 15
    while window._import_worker is not None and time.monotonic() < deadline:
        QApplication.processEvents(); QTest.qWait(10)
    QApplication.processEvents()
    assert window._import_worker is None
    assert len(window.articles) == 1
    assert window.actions['导入文章'].isEnabled()
def test_keyboard_navigation_does_not_grade_or_lose_inputs(tmp_path, ui):
    window = ui(Repository(tmp_path / "db"))
    window._import_paths([str(document(tmp_path))]); window.start_full_practice()
    window.show(); window.activateWindow(); QApplication.processEvents()
    first, second = list(window.answer_edits.values())
    first.setFocus(); QTest.keyClicks(first, "wrong spelling")
    QTest.keyClick(first, Qt.Key_Return); QApplication.processEvents()
    assert QApplication.focusWidget() is second
    QTest.keyClicks(second, "momentum")
    QTest.keyClick(second, Qt.Key_Return, Qt.ShiftModifier); QApplication.processEvents()
    assert QApplication.focusWidget() is first
    assert first.text() == "wrong spelling"
    assert not window.session.submitted and not window.session.results
    assert first.property("answerState") is None
    assert second.property("answerState") is None
    assert "%" not in window.score_text.text()


def test_focus_scrolls_down_and_text_can_be_selected(tmp_path, ui):
    path = tmp_path / "long.docx"
    doc = Document()
    for index in range(28):
        p = doc.add_paragraph(f"Line {index} has ")
        p.add_run("example").bold = True
        p.add_run("（单词） and enough context to wrap naturally across the page.")
    doc.save(path)
    window = ui(Repository(tmp_path / "db")); window._import_paths([str(path)])
    window.start_full_practice(); window.resize(1000, 650); window.show(); QApplication.processEvents()
    last_id = window.session.target_ids[-1]
    assert window.scroll.verticalScrollBar().maximum() > 0
    window._center_answer(last_id)
    assert window.scroll.verticalScrollBar().value() > 0
    labels = [label for label in window.content.findChildren(QLabel) if label.text().strip() == "enough"]
    assert labels and labels[0].textInteractionFlags() & Qt.TextSelectableByMouse


def test_answers_are_shown_in_full_inline_only_after_request_or_submit(tmp_path, ui, monkeypatch):
    confirmations = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: confirmations.append(args) or QMessageBox.Yes)
    window = ui(Repository(tmp_path / "db")); window._import_paths([str(document(tmp_path))])
    window.start_full_practice()
    item_id = window.session.target_ids[0]; hint = window.answer_hints[item_id]
    assert hint.isHidden() and "正确答案" not in window.progress_text.text()
    window.show_answers()
    assert not hint.isHidden() and hint.text() == "答案：extensive"
    assert window.actions["显示答案"].text() == "隐藏答案"
    window.show_answers()
    assert hint.isHidden() and window.actions["显示答案"].text() == "显示答案"
    window.show_answers()
    assert not hint.isHidden() and len(confirmations) == 0
    assert not window.session.viewed_answer
    assert window.answer_edits[item_id].placeholderText() == ""
    window.submit()
    assert window.answer_hints[item_id].text() == "正确答案：extensive"
    assert not window.answer_hints[item_id].isHidden()


def test_resumed_draft_hides_answers_without_viewing_flag(tmp_path, ui, monkeypatch):
    from src.ui.unfinished_dialog import UnfinishedPracticeDialog
    from PySide6.QtWidgets import QDialog
    monkeypatch.setattr(UnfinishedPracticeDialog, "exec", lambda _: QDialog.Accepted)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: QMessageBox.Yes)
    repository = Repository(tmp_path / "db")
    window = ui(repository); window._import_paths([str(document(tmp_path))])
    window.start_full_practice(); window.show_answers()
    assert not window.session.viewed_answer and window._answers_visible
    session_key = window.session.session_key
    window._save_active_draft()
    window.session = None
    window.resume_draft()
    assert window.session.session_key == session_key
    assert not window.session.viewed_answer and not window._answers_visible
    assert all(hint.isHidden() for hint in window.answer_hints.values())
    assert window.actions["显示答案"].text() == "显示答案"


def test_answer_toggle_preserves_inputs_and_does_not_save_or_grade(tmp_path, ui, monkeypatch):
    window = ui(Repository(tmp_path / "db"))
    window._import_paths([str(document(tmp_path))]); window.start_full_practice()
    item_id = window.session.target_ids[0]
    window.answer_edits[item_id].setText("my input")
    answers = dict(window.session.answers)
    def unexpected_save():
        raise AssertionError("Visibility toggle must not save a draft")
    with monkeypatch.context() as patcher:
        patcher.setattr(window, "_save_active_draft", unexpected_save)
        for _ in range(4):
            window.show_answers()
    assert window.answer_edits[item_id].text() == "my input"
    assert window.session.answers == answers
    assert not window.session.submitted and not window.session.results
    assert not window.session.viewed_answer


def test_history_does_not_display_viewing_classification(tmp_path, ui):
    repository = Repository(tmp_path / "db")
    window = ui(repository)
    window._import_paths([str(document(tmp_path))]); window.start_full_practice(); window.submit()
    with repository.connection() as conn:
        conn.execute("UPDATE practice_sessions SET viewed_answer=1")
    dialog = WrongPracticeHistoryDialog(repository, window)
    assert dialog.table.columnCount() == 6
    assert dialog.table.rowCount() == 1
    assert repository.list_history()[0]["viewed_answer"] == 1
    dialog.close()


def test_wrong_history_dialog_multiselect_delete_and_natural_sort(tmp_path, ui):
    repository = Repository(tmp_path / "history.db")
    with repository.connection() as connection:
        for index, name in enumerate(("djci10.docx", "djci2.docx", "djci1.docx")):
            connection.execute("""INSERT INTO practice_sessions
                (submitted_at,article_name,source_path,practice_type,target_count,
                 correct_count,wrong_count,accuracy,viewed_answer)
                VALUES (?, ?, ?, '整篇文章', 2, 1, 1, 50, 0)""",
                (f"2026-01-01T00:00:0{index}", name, name))
    parent = ui(repository)
    dialog = WrongPracticeHistoryDialog(repository, parent)
    dialog.sort.setCurrentIndex(1)
    assert [dialog.table.item(row, 1).text() for row in range(3)] == [
        "djci1.docx", "djci2.docx", "djci10.docx"]
    for row in (0, 2):
        dialog.table.selectionModel().select(
            dialog.table.model().index(row, 0),
            QItemSelectionModel.Select | QItemSelectionModel.Rows)
    assert len(dialog._selected_ids()) == 2
    assert not dialog.start_button.isEnabled()
    assert "2 条" in dialog.delete_button.text()
    dialog._delete_selected()
    assert [record["article_name"] for record in repository.list_history()] == ["djci2.docx"]
    dialog.deleteLater()


def test_primary_practice_action_is_in_status_panel_and_resume_removed(tmp_path, ui):
    window = ui(Repository(tmp_path / "layout.db"))
    toolbar_texts = [action.text() for action in window.main_toolbar.actions()]
    assert "继续练习" not in toolbar_texts
    assert "开始默写" not in toolbar_texts and "提交批改" not in toolbar_texts

    window._import_paths([str(sample_word(tmp_path))])
    assert window.primary_action_button.text() == "开始默写"
    assert window.primary_action_button.isEnabled()
    window.primary_action_button.click()
    assert window.session is not None and not window.session.submitted
    assert window.primary_action_button.text() == "提交批改"
    assert len(window._drafts) == 1

    window.primary_action_button.click()
    assert window.session.submitted and window.session.wrong_ids
    assert window.primary_action_button.text() == "只练本轮错词"
    window.primary_action_button.click()
    assert window.session is None and window.pending_review_session is not None
    assert window.primary_action_button.text() == "开始错词默写"


def test_failed_draft_save_keeps_previous_article_selected(tmp_path, ui, monkeypatch):
    repo = Repository(tmp_path / "db"); window = ui(repo)
    one = document(tmp_path); two = tmp_path / "z.docx"; two.write_bytes(one.read_bytes())
    window._import_paths([str(one), str(two)]); window.start_full_practice()
    next(iter(window.answer_edits.values())).setText("unsaved change")
    def fail(*args): raise OSError("unavailable")
    with monkeypatch.context() as patch:
        patch.setattr(repo, "save_draft", fail)
        window.article_list.setCurrentRow(1)
        assert window.article_list.currentRow() == 0
        assert window.session.article.source_path == str(one)


def test_full_practice_from_history_uses_complete_library_article(tmp_path, ui):
    window = ui(Repository(tmp_path / "db"))
    window._import_paths([str(document(tmp_path))]); window.start_full_practice()
    first_id = window.session.target_ids[0]
    window.answer_edits[first_id].setText("extensive"); window.submit()
    window.start_saved_wrong_practice(window.session.database_id)
    assert window.pending_review_session.total == 1
    window.start_full_practice()
    assert window.session.total == 1 and window.session.mode == PracticeMode.CURRENT_WRONG
    window.start_full_practice()
    assert window.session.total == 2
    assert window.session.mode == PracticeMode.FULL_ARTICLE
